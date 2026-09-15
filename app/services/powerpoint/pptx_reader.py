from collections.abc import Callable
from io import BytesIO

from pptx import Presentation

from app.services.document.document_operation_registry import DocumentType
from app.services.document.extraction_models import (
    DocumentBlock,
    DocumentBlockType,
    DocumentLocation,
    DocumentMetadata,
    ExtractedDocument,
    ExtractedTable,
)


class PptxReader:
    def __init__(self, presentation_factory: Callable = Presentation) -> None:
        self.presentation_factory = presentation_factory

    def read(self, file_data: bytes) -> ExtractedDocument:
        presentation = self.presentation_factory(BytesIO(file_data))
        parts: list[str] = []
        blocks: list[DocumentBlock] = []
        tables: list[ExtractedTable] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    rows = tuple(
                        tuple(cell.text.strip() for cell in row.cells) for row in shape.table.rows
                    )
                    if rows:
                        location = DocumentLocation(slide=slide_index)
                        tables.append(
                            ExtractedTable(
                                name=f"Slide {slide_index} table",
                                rows=rows,
                                columns=rows[0],
                                header_detected=len(rows) > 1,
                                location=location,
                            )
                        )
                        blocks.append(
                            DocumentBlock(
                                type=DocumentBlockType.TABLE,
                                cells=rows,
                                location=location,
                                style={"shape_name": str(getattr(shape, "name", ""))},
                            )
                        )
                        slide_parts.append("\n".join("\t".join(row) for row in rows))
                elif getattr(shape, "has_text_frame", False):
                    value = str(getattr(shape, "text", "")).strip()
                    if value:
                        slide_parts.append(value)
                        blocks.append(
                            DocumentBlock(
                                type=DocumentBlockType.SLIDE,
                                text=value,
                                location=DocumentLocation(slide=slide_index),
                                style={
                                    "shape_name": str(getattr(shape, "name", "")),
                                    "shape_type": str(getattr(shape, "shape_type", "")),
                                },
                            )
                        )
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"Notes: {notes}")
                    blocks.append(
                        DocumentBlock(
                            type=DocumentBlockType.PARAGRAPH,
                            text=notes,
                            location=DocumentLocation(slide=slide_index),
                            style={"notes": True},
                        )
                    )
            if slide_parts:
                parts.append(f"[Slide {slide_index}]\n" + "\n".join(slide_parts))
        return ExtractedDocument(
            document_type=DocumentType.PPTX,
            text="\n".join(parts).strip(),
            tables=tuple(tables),
            blocks=tuple(blocks),
            metadata=DocumentMetadata(
                slide_count=len(presentation.slides),
                author=presentation.core_properties.author or None,
                title=presentation.core_properties.title or None,
            ),
            coverage=f"{len(presentation.slides)} of {len(presentation.slides)} slides",
        )
