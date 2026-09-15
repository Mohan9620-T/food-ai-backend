from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

from app.services.document.content_parser import document_title
from app.services.document.extraction_models import (
    GeneratedTableContent,
    StructuredDocumentContent,
)


class PptxGenerator:
    MAX_LINES_PER_SLIDE = 7

    def generate(self, content: str) -> bytes:
        presentation = Presentation()
        self._configure(presentation)
        title = document_title(content)
        self._add_title_slide(presentation, title)

        lines = [line.strip().lstrip("#").strip() for line in content.splitlines() if line.strip()]
        if lines and lines[0] == title:
            lines = lines[1:]
        for start in range(0, len(lines), self.MAX_LINES_PER_SLIDE):
            chunk = lines[start : start + self.MAX_LINES_PER_SLIDE]
            self._add_text_slide(
                presentation,
                "Details" if start == 0 else "Details (continued)",
                chunk,
            )
        return self._save(presentation)

    def generate_structured(self, content: StructuredDocumentContent) -> bytes:
        presentation = Presentation()
        self._configure(presentation)
        self._add_title_slide(presentation, content.title)
        self._add_text_slides(
            presentation,
            "Overview",
            self._block_lines(content.paragraphs, content.bullet_lists),
        )
        for table in content.tables:
            self._add_table_slide(presentation, table)
        for section in content.sections:
            self._add_text_slides(
                presentation,
                section.heading,
                self._block_lines(section.paragraphs, section.bullet_lists),
            )
            for table in section.tables:
                self._add_table_slide(presentation, table, prefix=section.heading)
        return self._save(presentation)

    @staticmethod
    def _configure(presentation) -> None:
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

    @staticmethod
    def _add_title_slide(presentation, title: str) -> None:
        title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        title_slide.shapes.title.text = title
        if len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = "Generated from the conversation"

    def _add_text_slides(self, presentation, title: str, lines: list[str]) -> None:
        for start in range(0, len(lines), self.MAX_LINES_PER_SLIDE):
            suffix = "" if start == 0 else " (continued)"
            self._add_text_slide(
                presentation,
                f"{title}{suffix}",
                lines[start : start + self.MAX_LINES_PER_SLIDE],
            )

    @staticmethod
    def _add_text_slide(presentation, title: str, lines: list[str]) -> None:
        if not lines:
            return
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line.lstrip("-*+ ")
            paragraph.font.size = Pt(20)

    @staticmethod
    def _block_lines(paragraphs: list[str], bullet_lists: list[list[str]]) -> list[str]:
        return [*paragraphs, *(f"• {item}" for items in bullet_lists for item in items)]

    @staticmethod
    def _add_table_slide(
        presentation,
        table_content: GeneratedTableContent,
        *,
        prefix: str | None = None,
    ) -> None:
        rows = [*([table_content.headers] if table_content.headers else []), *table_content.rows]
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = (
            f"{prefix}: {table_content.title}" if prefix else table_content.title
        )[:120]
        shape = slide.shapes.add_table(
            len(rows),
            len(rows[0]),
            Inches(0.7),
            Inches(1.5),
            Inches(11.9),
            Inches(5.2),
        )
        for row_index, values in enumerate(rows):
            for column_index, value in enumerate(values):
                cell = shape.table.cell(row_index, column_index)
                cell.text = value
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(12)
                    if table_content.headers and row_index == 0:
                        paragraph.font.bold = True

    @staticmethod
    def _save(presentation) -> bytes:
        buffer = BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()
