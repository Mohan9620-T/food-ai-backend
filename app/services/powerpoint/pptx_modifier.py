from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from app.services.document.exceptions import (
    DocumentLocatorError,
    UnsupportedDocumentModificationError,
)
from app.services.document.extraction_models import DocumentEdit, DocumentEditAction


class PptxModifier:
    """Apply text, table, and basic paragraph-format edits to PowerPoint files."""

    def modify(self, file_data: bytes, edits: tuple[DocumentEdit, ...]) -> bytes:
        presentation = Presentation(BytesIO(file_data))
        for edit in edits:
            self._apply(presentation, edit)
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    def _apply(self, presentation, edit: DocumentEdit) -> None:
        if edit.action == DocumentEditAction.REPLACE_TEXT:
            paragraph, old = self._resolve_text_paragraph(presentation, edit)
            new = "" if edit.value is None else str(edit.value)
            self._replace_across_runs(paragraph, old, new)
        elif edit.action == DocumentEditAction.UPDATE_TABLE_CELL:
            cell = self._resolve_table_cell(presentation, edit)
            cell.text = "" if edit.value is None else str(edit.value)
        elif edit.action == DocumentEditAction.FORMAT_PARAGRAPH:
            paragraph, _ = self._resolve_text_paragraph(presentation, edit)
            self._format_paragraph(paragraph, edit.options)
        else:
            raise UnsupportedDocumentModificationError(
                f"The PowerPoint edit '{edit.action.value}' is not supported."
            )

    def _resolve_text_paragraph(self, presentation, edit):
        text = (edit.locator.text or edit.locator.heading or "").strip()
        if not text:
            raise DocumentLocatorError("A PowerPoint text edit requires a text locator.")
        matches: list[Any] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            if edit.locator.slide is not None and edit.locator.slide != slide_number:
                continue
            for shape in slide.shapes:
                if (
                    edit.locator.shape_name
                    and shape.name.casefold() != edit.locator.shape_name.casefold()
                ):
                    continue
                if not getattr(shape, "has_text_frame", False):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    matches.extend(paragraph for _ in range(paragraph.text.count(text)))
        if not matches:
            location = f" on slide {edit.locator.slide}" if edit.locator.slide else ""
            raise DocumentLocatorError(
                f"I couldn't find text matching '{text}'{location} in this presentation."
            )
        if len(matches) > 1:
            raise DocumentLocatorError(
                f"More than one presentation location matches '{text}'. "
                "Please identify the slide and shape."
            )
        return matches[0], text

    @staticmethod
    def _replace_across_runs(paragraph, old: str, new: str) -> None:
        runs = list(paragraph.runs)
        combined = "".join(run.text for run in runs)
        start = combined.find(old)
        end = start + len(old)
        offset = 0
        for run in runs:
            run_start = offset
            run_end = offset + len(run.text)
            offset = run_end
            if run_end <= start or run_start >= end:
                continue
            local_start = max(start - run_start, 0)
            local_end = min(end - run_start, len(run.text))
            prefix = run.text[:local_start]
            suffix = run.text[local_end:]
            run.text = prefix + new + suffix if run_start <= start < run_end else suffix

    @staticmethod
    def _resolve_table_cell(presentation, edit):
        if edit.locator.slide is None or edit.locator.table_index is None:
            raise DocumentLocatorError("A PowerPoint table edit requires slide and table indexes.")
        if edit.locator.slide > len(presentation.slides):
            raise DocumentLocatorError(f"Slide {edit.locator.slide} does not exist.")
        slide = presentation.slides[edit.locator.slide - 1]
        tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
        if edit.locator.table_index > len(tables):
            raise DocumentLocatorError(
                f"Table {edit.locator.table_index} does not exist on slide {edit.locator.slide}."
            )
        if edit.locator.row is None or edit.locator.column is None:
            raise DocumentLocatorError("A PowerPoint table edit requires row and column numbers.")
        table = tables[edit.locator.table_index - 1]
        if edit.locator.row > len(table.rows) or edit.locator.column > len(table.columns):
            raise DocumentLocatorError("The requested PowerPoint table cell does not exist.")
        return table.cell(edit.locator.row - 1, edit.locator.column - 1)

    @staticmethod
    def _format_paragraph(paragraph, options: dict[str, object]) -> None:
        if "alignment" in options:
            alignment = str(options["alignment"]).casefold()
            values = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "centre": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
                "justify": PP_ALIGN.JUSTIFY,
            }
            if alignment not in values:
                raise UnsupportedDocumentModificationError(
                    "Paragraph alignment must be left, center, right, or justify."
                )
            paragraph.alignment = values[alignment]
        for run in paragraph.runs:
            if "bold" in options:
                run.font.bold = bool(options["bold"])
            if "italic" in options:
                run.font.italic = bool(options["italic"])
            if "font_size_pt" in options:
                run.font.size = Pt(float(str(options["font_size_pt"])))
