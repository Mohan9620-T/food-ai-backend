"""Append user-supplied records without regenerating an existing workbook."""

import csv
import json
import math
import re
from copy import copy
from dataclasses import dataclass
from io import BytesIO, StringIO

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from pydantic import BaseModel, ConfigDict, Field

from app.services.document.exceptions import DocumentLocatorError, InvalidDocumentError
from app.services.spreadsheet.workbook_append_patch import preserve_original_package


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


class AppendRowsRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    data: str = Field(min_length=1, max_length=100000)
    sheet_name: str | None = Field(default=None, min_length=1, max_length=128)

    @staticmethod
    def _is_action_line(line: str) -> bool:
        return bool(
            re.search(
                r"\b(?:add|append|insert)\b.{0,70}\b(?:new\s+)?(?:rows?|records?|items?|data(?:['’]?s)?)\b",
                line,
                re.I,
            )
        ) and not bool(
            re.search(
                r"\b(?:add|append|insert)\s+(?:(?:a|an|the|new|these|following|extra)\s+)*"
                r"(?:columns?|headers?|formulas?|filters?)\b",
                line,
                re.I,
            )
        )

    @classmethod
    def is_request(cls, instruction: str) -> bool:
        # People paste the instruction before or after the data, so inspect
        # whichever end actually carries it, not action words inside the records.
        lines = [line for line in instruction.splitlines() if line.strip()]
        if not lines:
            return False
        return cls._is_action_line(lines[0]) or cls._is_action_line(lines[-1])

    @classmethod
    def from_instruction(cls, instruction: str) -> "AppendRowsRequest | None":
        stripped = instruction.strip(" \r\n")
        lines = stripped.splitlines()
        non_blank = [index for index, line in enumerate(lines) if line.strip()]
        if not non_blank:
            return None
        first_index, last_index = non_blank[0], non_blank[-1]
        if cls._is_action_line(lines[first_index]):
            lead = lines[first_index]
            body = "\n".join(lines[first_index + 1 :])
        elif cls._is_action_line(lines[last_index]):
            lead = lines[last_index]
            body = "\n".join(lines[:last_index])
        else:
            return None
        if not body.strip():
            _, _, body = lead.partition(":")
        if not body.strip():
            return None
        sheet = re.search(r"\b(?:sheet|worksheet)\s+[`\"']([^`\"']+)[`\"']", lead, re.I)
        return cls(data=body.strip(" \r\n"), sheet_name=sheet[1] if sheet else None)


@dataclass(frozen=True)
class ResolvedAppend:
    sheet_name: str
    header_row: int
    columns: tuple[int, ...]
    rows: tuple[tuple[object, ...], ...]
    start_row: int

    @property
    def summary(self) -> str:
        return f"Added {len(self.rows)} new rows to '{self.sheet_name}' (rows {self.start_row}–{self.start_row + len(self.rows) - 1})."


class WorkbookRowAppender:
    @staticmethod
    def _parse(data: str) -> tuple[list[str] | None, list[list[object]]]:
        text = re.sub(r"^```(?:json|csv|tsv)?\s*\n|\n```$", "", data.strip(" \r\n"), flags=re.I)
        if text.startswith("[") or text.startswith("{"):
            try:
                payload = json.loads(text)
            except ValueError as error:
                raise DocumentLocatorError(
                    "The new rows contain invalid JSON. Paste a JSON array of rows."
                ) from error
            if isinstance(payload, dict):
                payload = [payload]
            if not isinstance(payload, list) or not payload or len(payload) > 500:
                raise DocumentLocatorError("Provide between 1 and 500 new rows.")
            if all(isinstance(row, dict) for row in payload):
                headers = list(dict.fromkeys(key for row in payload for key in row))
                return headers, [[row.get(key) for key in headers] for row in payload]
            if all(isinstance(row, list) for row in payload):
                return None, payload
            raise DocumentLocatorError("Use either named JSON records or arrays of cell values.")
        lines = [line for line in text.splitlines() if line.strip()]
        if "\t" in text:
            return None, [list(row) for row in csv.reader(StringIO(text), delimiter="\t")]
        if lines and all("|" in line for line in lines):
            rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
            return None, [
                list(row)
                for row in rows
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)
            ]
        if "," in text:
            return None, [list(row) for row in csv.reader(StringIO(text))]
        return None, [[line] for line in lines]

    @staticmethod
    def _catalog_row(line: str, sheet, header_row: int, headers: list[str]) -> list[object] | None:
        # Spaces inside names and categories are data, not column separators.
        # Resolve only a known item-master layout, using its existing type/category values.
        if [header_key(h) for h in headers[:7]] != [
            "itemcode",
            "itemname",
            "itemtype",
            "category",
            "uom",
            "baseuom",
            "uomfactor",
        ]:
            return None
        types = {
            str(row[0])
            for row in sheet.iter_rows(
                min_row=header_row + 1, min_col=3, max_col=3, values_only=True
            )
            if row[0]
        }
        categories = {
            str(row[0])
            for row in sheet.iter_rows(
                min_row=header_row + 1, min_col=4, max_col=4, values_only=True
            )
            if row[0]
        }
        matches = []
        for item_type in types:
            for category in categories:
                pattern = rf"^(\S+)\s+(.+?)\s+({re.escape(item_type)})\s+({re.escape(category)})\s+(.+?)\s+([+-]?\d+(?:\.\d+)?)$"
                match = re.fullmatch(pattern, line.strip(), re.I)
                if match:
                    tail = match[5]
                    unit = re.match(r"(?:\d+(?:\.\d+)?\s+)?\S+", tail)
                    assert unit is not None
                    base_unit = tail[unit.end() :].strip()
                    if base_unit:
                        matches.append(
                            [match[1], match[2], match[3], match[4], unit[0], base_unit, match[6]]
                        )
        return matches[0] if len(matches) == 1 else None

    def resolve(self, file_data: bytes, request: AppendRowsRequest) -> ResolvedAppend:
        workbook = load_workbook(BytesIO(file_data))
        try:
            return self._resolve(workbook, request)
        finally:
            workbook.close()

    def _resolve(self, workbook, request: AppendRowsRequest) -> ResolvedAppend:
        requested_headers, parsed = self._parse(request.data)
        if not parsed or len(parsed) > 500:
            raise DocumentLocatorError(
                "Paste between 1 and 500 rows to add, copied from Excel with the column separators."
            )
        candidates = []
        for sheet in workbook:
            if request.sheet_name and sheet.title.casefold() != request.sheet_name.casefold():
                continue
            if sheet.max_row * sheet.max_column > 2_000_000:
                raise InvalidDocumentError("This worksheet is too large to update safely.")
            # Headers may follow title/notes rows. Only use a contiguous, unique header row.
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 20)):
                headers = [str(cell.value or "").strip() for cell in row]
                while headers and not headers[-1]:
                    headers.pop()
                keys = [header_key(value) for value in headers]
                if not keys or any(not key for key in keys) or len(set(keys)) != len(keys):
                    continue
                rows = [list(item) for item in parsed]
                supplied_headers = requested_headers
                if (
                    supplied_headers is None
                    and all(isinstance(cell, str) for cell in rows[0])
                    and len(rows[0]) > 1
                    and all(header_key(cell) in keys for cell in rows[0])
                ):
                    supplied_headers = [str(cell) for cell in rows.pop(0)]
                if supplied_headers:
                    normalized = [header_key(h) for h in supplied_headers]
                    if len(set(normalized)) != len(normalized) or any(
                        key not in keys for key in normalized
                    ):
                        continue
                    columns = tuple(keys.index(key) + 1 for key in normalized)
                else:
                    if all(len(item) == 1 for item in rows) and len(headers) > 1:
                        expanded = [
                            self._catalog_row(str(item[0]), sheet, row[0].row, headers)
                            for item in rows
                        ]
                        if not all(expanded):
                            continue
                        rows = [item for item in expanded if item is not None]
                    width = len(rows[0]) if rows else 0
                    if width < 1 or width > len(headers):
                        continue
                    columns = tuple(range(1, width + 1))
                    if (
                        len(workbook.worksheets) > 1
                        and not request.sheet_name
                        and not re.search(r"(?:code|id)$", keys[0])
                    ):
                        continue
                    existing_first = [
                        item[0]
                        for item in sheet.iter_rows(
                            min_row=row[0].row + 1,
                            max_row=min(sheet.max_row, row[0].row + 20),
                            max_col=1,
                            values_only=True,
                        )
                        if item[0] is not None
                    ]
                    if (
                        existing_first
                        and all(isinstance(value, (int, float)) for value in existing_first)
                        and any(not str(item[0]).isdigit() for item in rows)
                    ):
                        continue
                    # An audit/reference sheet can share the leading code/name columns.
                    # Reject positional mappings that put text into consistently numeric columns.
                    incompatible = False
                    for column in columns:
                        samples = [
                            cells[0]
                            for cells in sheet.iter_rows(
                                min_row=row[0].row + 1,
                                max_row=min(sheet.max_row, row[0].row + 20),
                                min_col=column,
                                max_col=column,
                                values_only=True,
                            )
                            if cells[0] is not None
                        ]
                        if samples and all(
                            isinstance(value, (int, float)) and not isinstance(value, bool)
                            for value in samples
                        ):
                            incompatible = any(
                                item[column - 1] not in (None, "")
                                and not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", str(item[column - 1]))
                                for item in rows
                                if len(item) >= column
                            )
                            if incompatible:
                                break
                    if incompatible:
                        continue
                if not rows or any(len(item) != len(columns) for item in rows):
                    continue
                if any(
                    not isinstance(value, (str, int, float, bool, type(None)))
                    or (isinstance(value, float) and not math.isfinite(value))
                    or (isinstance(value, str) and len(value) > 32767)
                    for item in rows
                    for value in item
                ):
                    raise DocumentLocatorError(
                        "Each new cell must contain text up to 32,767 characters, a finite number, a boolean, or a blank."
                    )
                last_data = max(
                    (
                        cell.row
                        for cells in sheet
                        for cell in cells
                        if cell.value is not None and cell.data_type != "f"
                    ),
                    default=row[0].row,
                )
                start = last_data + 1
                # Formula-only template rows can be filled without replacing their formulas.
                while any(
                    sheet.cell(index, column).value is not None
                    for index in range(start, start + len(rows))
                    for column in columns
                ):
                    start += 1
                if start + len(rows) - 1 > 1_048_576:
                    raise InvalidDocumentError(
                        "There is no space for these rows in this worksheet."
                    )
                if any(table.totalsRowCount for table in sheet.tables.values()):
                    raise DocumentLocatorError(
                        "Name a worksheet without a totals row before appending records."
                    )
                if any(
                    merged.min_row <= start + len(rows) - 1 and merged.max_row >= start
                    for merged in sheet.merged_cells.ranges
                ):
                    raise DocumentLocatorError(
                        "The destination rows contain merged cells. Choose a worksheet with unmerged data rows."
                    )
                candidates.append(
                    ResolvedAppend(
                        sheet.title, row[0].row, columns, tuple(tuple(item) for item in rows), start
                    )
                )
                break
        if len(candidates) != 1:
            names = ", ".join(item.sheet_name for item in candidates) or ", ".join(
                workbook.sheetnames
            )
            raise DocumentLocatorError(
                f"Specify the worksheet and paste the new rows with column headers (tabs, CSV, a table, or JSON). Available sheets: {names}."
            )
        return candidates[0]

    def append(self, file_data: bytes, request: AppendRowsRequest) -> tuple[bytes, ResolvedAppend]:
        workbook = load_workbook(BytesIO(file_data))
        try:
            resolved = self._resolve(workbook, request)
            sheet = workbook[resolved.sheet_name]
            template_row = resolved.start_row - 1
            end_row = resolved.start_row + len(resolved.rows) - 1
            numeric_columns = set()
            for column in resolved.columns:
                samples = [
                    cells[0]
                    for cells in sheet.iter_rows(
                        min_row=resolved.header_row + 1,
                        max_row=min(template_row, resolved.header_row + 20),
                        min_col=column,
                        max_col=column,
                        values_only=True,
                    )
                    if cells[0] is not None
                ]
                key = header_key(sheet.cell(resolved.header_row, column).value)
                if (
                    samples
                    and not re.search(r"(?:code|id)$", key)
                    and all(
                        isinstance(value, (int, float)) and not isinstance(value, bool)
                        for value in samples
                    )
                ):
                    numeric_columns.add(column)
            for offset, values in enumerate(resolved.rows):
                index = resolved.start_row + offset
                for column in range(1, sheet.max_column + 1):
                    target = sheet.cell(index, column)
                    template = sheet.cell(template_row, column)
                    if not target.has_style:
                        target._style = copy(template._style)
                    if (
                        target.value is None
                        and column not in resolved.columns
                        and template.data_type == "f"
                    ):
                        target.value = Translator(
                            template.value, origin=template.coordinate
                        ).translate_formula(target.coordinate)
                for column, value in zip(resolved.columns, values):
                    target = sheet.cell(index, column)
                    # Keep codes/leading zeroes as text, but numeric data columns as numbers.
                    if (
                        isinstance(value, str)
                        and column in numeric_columns
                        and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", value)
                    ):
                        value = float(value) if "." in value else int(value)
                    target.value = value if value != "" else None
                    if isinstance(value, str):
                        target.data_type = "s"
                if sheet.row_dimensions[index].height is None:
                    sheet.row_dimensions[index].height = sheet.row_dimensions[template_row].height
            for table in sheet.tables.values():
                if table.totalsRowCount:
                    raise DocumentLocatorError(
                        "Name a worksheet without a totals row before appending records."
                    )
                min_col, min_row, max_col, max_row = range_boundaries(table.ref)
                if min_row == resolved.header_row and max_row + 1 >= resolved.start_row:
                    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max(max_row, end_row)}"
                    if table.autoFilter:
                        table.autoFilter.ref = table.ref
            if sheet.auto_filter.ref:
                left, top, right, bottom = range_boundaries(sheet.auto_filter.ref)
                if bottom + 1 >= resolved.start_row:
                    sheet.auto_filter.ref = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{max(bottom, end_row)}"
            for validation in sheet.data_validations.dataValidation:
                for cell_range in list(validation.ranges):
                    if (
                        cell_range.min_row <= template_row <= cell_range.max_row
                        and end_row > cell_range.max_row
                    ):
                        validation.add(
                            f"{get_column_letter(cell_range.min_col)}{cell_range.max_row + 1}:{get_column_letter(cell_range.max_col)}{end_row}"
                        )
            if workbook.calculation is None:
                workbook.calculation = CalcProperties()
            workbook.calculation.fullCalcOnLoad = True
            output = BytesIO()
            workbook.save(output)
            file_data = preserve_original_package(
                file_data, output.getvalue(), resolved.sheet_name, resolved.start_row, end_row
            )
            # Reopen the actual output before accepting it as a generated file.
            check = load_workbook(BytesIO(file_data))
            try:
                for index in range(resolved.start_row, end_row + 1):
                    for column in resolved.columns:
                        actual = check[resolved.sheet_name].cell(index, column).value
                        expected = sheet.cell(index, column).value
                        if actual == expected or (
                            isinstance(actual, float)
                            and isinstance(expected, float)
                            and math.isclose(actual, expected, rel_tol=1e-15)
                        ):
                            continue
                        raise InvalidDocumentError(
                            "The appended rows could not be verified. The original workbook is unchanged."
                        )
            finally:
                check.close()
            return file_data, resolved
        finally:
            workbook.close()
