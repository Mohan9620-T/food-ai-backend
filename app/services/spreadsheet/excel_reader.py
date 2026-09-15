from collections.abc import Callable
from contextlib import closing
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.document.document_operation_registry import DocumentType
from app.services.document.extraction_models import (
    DocumentBlock,
    DocumentBlockType,
    DocumentLocation,
    DocumentMetadata,
    ExtractedDocument,
    ExtractedTable,
)


class ExcelReader:
    def __init__(self, workbook_loader: Callable = load_workbook) -> None:
        self.workbook_loader = workbook_loader

    def read(self, file_data: bytes) -> ExtractedDocument:
        text_parts: list[str] = []
        tables: list[ExtractedTable] = []
        blocks: list[DocumentBlock] = []
        with (
            closing(
                self.workbook_loader(BytesIO(file_data), read_only=True, data_only=True)
            ) as values_workbook,
            closing(
                self.workbook_loader(BytesIO(file_data), read_only=True, data_only=False)
            ) as formulas_workbook,
        ):
            for formula_sheet in formulas_workbook.worksheets:
                value_sheet = values_workbook[formula_sheet.title]
                rows_list: list[tuple[str, ...]] = []
                formula_cells: list[dict[str, object]] = []
                first_row: int | None = None
                last_row = 0
                first_column: int | None = None
                max_column = 0
                # Read the formula and cached-value worksheets in lockstep. Restarting a
                # read-only worksheet iterator for every row makes large workbooks O(n^2)
                # and caused simple questions such as "how many rows?" to take minutes.
                for row_number, (formula_row, value_row) in enumerate(
                    zip(
                        formula_sheet.iter_rows(),
                        value_sheet.iter_rows(),
                        strict=False,
                    ),
                    start=1,
                ):
                    if not any(cell.value is not None for cell in formula_row):
                        continue
                    first_row = first_row or row_number
                    last_row = row_number
                    populated_columns = [
                        cell.column for cell in formula_row if cell.value is not None
                    ]
                    first_column = min(first_column or populated_columns[0], *populated_columns)
                    max_column = max(max_column, *populated_columns)
                    rendered: list[str] = []
                    for formula_cell, value_cell in zip(formula_row, value_row, strict=False):
                        formula = formula_cell.value
                        cached = value_cell.value
                        display = cached if cached is not None else formula
                        rendered.append("" if display is None else str(display))
                        if isinstance(formula, str) and formula.startswith("="):
                            formula_cells.append(
                                {
                                    "cell": formula_cell.coordinate,
                                    "formula": formula,
                                    "cached_value": cached,
                                }
                            )
                    rows_list.append(tuple(rendered))
                rows = tuple(rows_list)
                if not rows:
                    continue
                cell_range = (
                    f"{get_column_letter(first_column)}{first_row}:"
                    f"{get_column_letter(max_column)}{last_row}"
                    if first_row is not None and first_column is not None
                    else None
                )
                location = DocumentLocation(sheet=formula_sheet.title, cell_range=cell_range)
                tables.append(
                    ExtractedTable(
                        name=formula_sheet.title,
                        rows=rows,
                        columns=rows[0],
                        header_detected=len(rows) > 1,
                        location=location,
                    )
                )
                blocks.append(
                    DocumentBlock(
                        type=DocumentBlockType.SHEET_RANGE,
                        text=f"Worksheet {formula_sheet.title}",
                        cells=rows,
                        location=location,
                        style={"formulas": formula_cells},
                    )
                )
                text_parts.append(
                    f"[{formula_sheet.title}]\n" + "\n".join("\t".join(row) for row in rows)
                )
        return ExtractedDocument(
            document_type=DocumentType.XLSX,
            text="\n".join(text_parts).strip(),
            tables=tuple(tables),
            blocks=tuple(blocks),
            metadata=DocumentMetadata(sheet_names=tuple(table.name for table in tables)),
            coverage=f"{len(tables)} of {len(formulas_workbook.sheetnames)} populated sheets",
        )
