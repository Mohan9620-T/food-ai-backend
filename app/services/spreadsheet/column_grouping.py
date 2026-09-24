"""Split actual workbook records into worksheets without regenerating their data."""

import re
from collections import Counter
from copy import copy
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, quote_sheetname
from pydantic import BaseModel, ConfigDict, Field

from app.services.document.exceptions import DocumentLocatorError, InvalidDocumentError


def header_key(value: object) -> str:
    return re.sub(r"[\W_]+", "", str(value or "").casefold())


def positive_instruction(instruction: str) -> str:
    return re.sub(r"\b(?:do not|don't|never)\b[^.\n]*", " ", instruction, flags=re.I)


class ColumnGroupingRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column: str = Field(min_length=1, max_length=128)
    sheet_name: str | None = Field(default=None, min_length=1, max_length=128)
    instruction: str = Field(default="", max_length=50000)

    @classmethod
    def from_instruction(cls, instruction: str) -> "ColumnGroupingRequest | None":
        text = positive_instruction(instruction)
        sheets = re.search(r"\b(?:(?:work)?sheets?|tabs?)\b|ஷீட்|தாள்", text, re.I)
        separate = re.search(
            r"\b(?:split|separate|different|thani|thaniya|thanithani)\b|தனி|பிரித்"
            r"|\b(?:one|a)\s+(?:work)?sheet\s+(?:for\s+)?(?:each|every|per)\b"
            r"|\b(?:each|every)\s+(?:price|rate|cost)[^.\n]{0,40}\b(?:work)?sheet\b",
            text,
            re.I,
        )
        if not sheets or not separate:
            return None
        if re.search(r"\b(?:prices?|rates?|cost)\b|விலை|விலையில்|விலைக்கு", text, re.I):
            column = "Price"
        else:
            match = re.search(
                r'\b(?:by|per|each|every)\s+(?:the\s+|unique\s+)?(?:column\s+)?["`]([^"`\n]+)["`]',
                text,
                re.I,
            )
            # Keep compound legacy category-formatting requests with the full planner.
            if not match:
                return None
            column = match.group(1)
        sheet_match = re.search(
            r'\bsource\s+worksheet\s+is\s+named\s*:\s*([^\n]+)|\b(?:from|use)\s+(?:only\s+)?(?:the\s+)?["`]([^"`\n]+)["`]\s+(?:work)?sheet',
            text,
            re.I,
        )
        return cls(
            column=column,
            sheet_name=(
                next(v for v in sheet_match.groups() if v).strip() if sheet_match else None
            ),
            instruction=instruction,
        )


@dataclass(frozen=True)
class GroupedWorkbook:
    file_data: bytes
    summary: str


class WorkbookColumnGrouper:
    MAX_GROUPS = 200

    @staticmethod
    def _key(value: object) -> tuple[int, object]:
        if value is None or isinstance(value, str) and not value.strip():
            return (2, "")
        if not isinstance(value, bool):
            try:
                number = Decimal(str(value).strip())
                if number.is_finite():
                    return (0, number)
            except InvalidOperation:
                pass
        return (1, str(value))

    @staticmethod
    def _sheet_name(key: tuple[int, object], used: set[str], *, money: bool) -> str:
        kind, value = key
        if kind == 2:
            name = "Blank"
        elif kind == 0:
            number = Decimal(str(value))
            places = max(2 if money else 0, -int(number.normalize().as_tuple().exponent))
            name = f"{number:.{min(places, 20)}f}"
        else:
            name = str(value)
        name = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "-", name).strip(" '")[:31] or "Group"
        candidate, suffix = name, 2
        while candidate.casefold() in used or candidate.casefold() == "history":
            tail = f" ({suffix})"
            candidate = name[: 31 - len(tail)] + tail
            suffix += 1
        used.add(candidate.casefold())
        return candidate

    @staticmethod
    def _source(workbook, request: ColumnGroupingRequest, filename: str):
        aliases = {header_key(request.column)}
        if aliases & {"price", "rate", "cost", "unitprice", "unitrate"}:
            aliases |= {"price", "rate", "cost", "unitprice", "unitrate", "விலை"}
        candidates = []
        for sheet in workbook:
            for row in sheet.iter_rows(max_row=min(sheet.max_row, 25)):
                columns = [c.column for c in row if header_key(c.value) in aliases]
                if len(columns) == 1:
                    candidates.append((sheet, row[0].row, columns[0]))
                    break
        if request.sheet_name:
            named = [
                item
                for item in candidates
                if header_key(item[0].title) == header_key(request.sheet_name)
            ]
            if named:
                return named[0]
            # A filename is commonly mistaken for the tab name. Recover only when
            # the named source is this file and exactly one tab has the real column.
            if not (
                len(candidates) == 1
                and header_key(request.sheet_name) == header_key(Path(filename).stem)
            ):
                raise DocumentLocatorError(
                    f"Which source sheet should I use? '{request.sheet_name}' was not found with a {request.column} column."
                )
        text = positive_instruction(request.instruction)
        named = [
            item
            for item in candidates
            if re.search(r"(?<!\w)" + re.escape(item[0].title) + r"(?!\w)", text, re.I)
        ]
        if len(named) == 1:
            return named[0]
        if len(candidates) == 1:
            return candidates[0]
        choices = ", ".join(item[0].title for item in candidates)
        raise DocumentLocatorError(
            f"Which sheet should I split by {request.column}? Matching sheets: {choices}."
            if candidates
            else f"I couldn't find a column named {request.column} in the workbook."
        )

    def split(self, data: bytes, request: ColumnGroupingRequest, filename: str) -> GroupedWorkbook:
        workbook = load_workbook(BytesIO(data), data_only=False)
        cached = load_workbook(BytesIO(data), data_only=True)
        try:
            source, header_row, column = self._source(workbook, request, filename)
            if any(r.max_row > header_row for r in source.merged_cells.ranges):
                raise InvalidDocumentError(
                    "Unmerge the item data cells before splitting rows into sheets."
                )
            groups: dict[tuple[int, object], list[int]] = {}
            for row in source.iter_rows(min_row=header_row + 1):
                if not any(cell.value is not None for cell in row):
                    continue
                cell = row[column - 1]
                value = cell.value
                if cell.data_type == "f":
                    value = cached[source.title][cell.coordinate].value
                    if value is None:
                        raise InvalidDocumentError(
                            "Recalculate and save the workbook first; the grouping column contains formulas without saved results."
                        )
                if cell.data_type == "e":
                    raise InvalidDocumentError(
                        "Correct the Excel errors in the grouping column before splitting."
                    )
                groups.setdefault(self._key(value), []).append(row[0].row)
            if not groups or len(groups) > self.MAX_GROUPS:
                raise InvalidDocumentError(
                    f"Splitting requires item rows and at most {self.MAX_GROUPS} distinct values."
                )
            originals = {s.title: list(s.values) for s in workbook}
            used = {name.casefold() for name in workbook.sheetnames}
            expected: dict[str, list[tuple[object, ...]]] = {}
            money = header_key(request.column) in {"price", "rate", "cost", "unitprice", "unitrate"}
            for key, row_indexes in sorted(groups.items()):
                title = self._sheet_name(key, used, money=money)
                target = workbook.create_sheet(title)
                for col, dimension in source.column_dimensions.items():
                    target.column_dimensions[col] = copy(dimension)
                target.freeze_panes = "A2"
                expected[title] = []
                for destination, original in enumerate([header_row, *row_indexes], start=1):
                    target.row_dimensions[destination] = copy(source.row_dimensions[original])
                    target.row_dimensions[destination].index = destination
                    values = []
                    for cell in source[original]:
                        value = cell.value
                        if cell.data_type == "f":
                            value = f"={quote_sheetname(source.title)}!{cell.coordinate}"
                        # Explicit blank cells retain trailing unnamed columns even
                        # when this entire price group has no notes in that column.
                        new = target.cell(destination, cell.column, "" if value is None else value)
                        if cell.data_type != "f":
                            new.data_type = cell.data_type
                        new._style = copy(cell._style)
                        new.number_format = cell.number_format
                        if cell.comment:
                            new.comment = copy(cell.comment)
                        if cell.hyperlink:
                            new.hyperlink = copy(cell.hyperlink)
                        values.append(value)
                    expected[title].append(tuple(values))
                for cells in target.columns:
                    letter = get_column_letter(cells[0].column)
                    width = min(60, max(len(str(c.value or "")) for c in cells) + 2)
                    target.column_dimensions[letter].width = max(
                        target.column_dimensions[letter].width or 0, width
                    )
                target.auto_filter.ref = (
                    f"A1:{get_column_letter(source.max_column)}{len(row_indexes) + 1}"
                )
            output = BytesIO()
            workbook.save(output)
            result = output.getvalue()
            with BytesIO(result) as buffer:
                reopened = load_workbook(buffer, data_only=False)
                try:
                    for name, rows in {**originals, **expected}.items():
                        if list(reopened[name].values) != rows:
                            raise InvalidDocumentError(
                                "Workbook row preservation check failed; no output was returned."
                            )
                    actual_indexes = Counter(i for indexes in groups.values() for i in indexes)
                    if any(count != 1 for count in actual_indexes.values()):
                        raise InvalidDocumentError(
                            "A source row appeared more than once in the grouped output."
                        )
                finally:
                    reopened.close()
            count = sum(len(rows) for rows in groups.values())
            summary = f"Split all {count} item rows from '{source.title}' into {len(groups)} sheets by {source.cell(header_row, column).value}. Every item appears exactly once across the new sheets; all original sheets and columns are retained."
            return GroupedWorkbook(result, summary)
        finally:
            workbook.close()
            cached.close()
