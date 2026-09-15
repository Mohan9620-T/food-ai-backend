import hashlib
import math
from contextlib import closing
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook


def cell_value(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


class SpreadsheetPreviewService:
    """Page through the exact stored workbook, retaining sheet/row/column positions."""

    def preview(
        self,
        data: bytes,
        *,
        sheet: str | None = None,
        row_offset: int = 0,
        column_offset: int = 0,
        row_limit: int = 100,
        column_limit: int = 50,
    ) -> dict:
        with (
            closing(
                load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            ) as formulas,
            closing(
                load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
            ) as values,
        ):
            sheets = [
                {"name": ws.title, "rows": ws.max_row or 0, "columns": ws.max_column or 0}
                for ws in formulas.worksheets
            ]
            active_name = sheet if sheet is not None else sheets[0]["name"]
            if active_name not in formulas.sheetnames:
                raise ValueError("Worksheet not found")
            source = formulas[active_name]
            cached = values[active_name]
            end_row = min(source.max_row or 0, row_offset + row_limit)
            end_column = min(source.max_column or 0, column_offset + column_limit)
            rows, formula_rows = [], []
            if row_offset < end_row and column_offset < end_column:
                bounds = dict(
                    min_row=row_offset + 1,
                    max_row=end_row,
                    min_col=column_offset + 1,
                    max_col=end_column,
                )
                for original, stored in zip(
                    source.iter_rows(**bounds), cached.iter_rows(**bounds), strict=True
                ):
                    rendered, row_formulas = [], []
                    for cell, cached_cell in zip(original, stored, strict=True):
                        formula = cell.value if cell.data_type == "f" else None
                        rendered.append(
                            cell_value(
                                cached_cell.value
                                if formula and cached_cell.value is not None
                                else cell.value
                            )
                        )
                        row_formulas.append(formula)
                    rows.append(rendered)
                    formula_rows.append(row_formulas)
            return {
                "sha256": hashlib.sha256(data).hexdigest(),
                "sheets": sheets,
                "sheet": active_name,
                "row_offset": row_offset,
                "column_offset": column_offset,
                "rows": rows,
                "formulas": formula_rows,
            }
