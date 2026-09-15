import re
from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.document.content_parser import tabular_rows
from app.services.document.extraction_models import (
    ExtractedTable,
    GeneratedTableContent,
    StructuredDocumentContent,
)


class ExcelGenerator:
    def generate(self, content: str) -> bytes:
        rows = tabular_rows(content)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Document"
        self._write_rows(sheet, rows)
        return self._save(workbook)

    def generate_tables(self, tables: Sequence[ExtractedTable]) -> bytes:
        populated = [table for table in tables if table.rows]
        if not populated:
            raise ValueError("The document does not contain a table to export.")
        workbook = Workbook()
        existing_titles: set[str] = set()
        for index, table in enumerate(populated):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = self._safe_sheet_title(table.name, existing_titles)
            existing_titles.add(sheet.title.casefold())
            self._write_rows(sheet, [list(row) for row in table.rows])
        return self._save(workbook)

    def generate_structured(self, content: StructuredDocumentContent) -> bytes:
        workbook = Workbook()
        overview = workbook.active
        overview.title = "Overview"
        overview_rows: list[list[str]] = [[content.title]]
        overview_rows.extend([paragraph] for paragraph in content.paragraphs)
        for items in content.bullet_lists:
            overview_rows.extend([[f"• {item}"] for item in items])
        for section in content.sections:
            overview_rows.append([section.heading])
            overview_rows.extend([paragraph] for paragraph in section.paragraphs)
            for items in section.bullet_lists:
                overview_rows.extend([[f"• {item}"] for item in items])
        self._write_rows(overview, overview_rows)

        existing_titles = {overview.title.casefold()}
        for content_table in content.all_tables():
            sheet = workbook.create_sheet(
                self._safe_sheet_title(content_table.title, existing_titles)
            )
            existing_titles.add(sheet.title.casefold())
            self._write_rows(sheet, self._content_table_rows(content_table))
        return self._save(workbook)

    @staticmethod
    def _content_table_rows(table: GeneratedTableContent) -> list[list[str]]:
        rows = [list(row) for row in table.rows]
        if table.headers:
            rows.insert(0, list(table.headers))
        return rows

    @staticmethod
    def _write_rows(sheet, rows: list[list[str]]) -> None:
        width = max(len(row) for row in rows)
        for row in rows:
            sheet.append([*row, *([""] * (width - len(row)))])
        sheet.freeze_panes = "A2" if len(rows) > 1 else None
        sheet.auto_filter.ref = sheet.dimensions if len(rows) > 1 else None
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="left", vertical="center")
        for column in range(1, width + 1):
            values = [
                str(sheet.cell(row=row, column=column).value or "")
                for row in range(1, len(rows) + 1)
            ]
            sheet.column_dimensions[get_column_letter(column)].width = min(
                60, max(10, max(map(len, values), default=0) + 2)
            )
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _safe_sheet_title(title: str, existing_titles: set[str]) -> str:
        base = re.sub(r"[\\/*?:\[\]]", " ", title).strip() or "Table"
        base = " ".join(base.split())[:31]
        candidate = base
        suffix = 2
        while candidate.casefold() in existing_titles:
            marker = f" {suffix}"
            candidate = f"{base[: 31 - len(marker)]}{marker}"
            suffix += 1
        return candidate

    @staticmethod
    def _save(workbook: Workbook) -> bytes:
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()
