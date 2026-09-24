from __future__ import annotations

from copy import copy
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from app.services.document.exceptions import (
    DocumentLocatorError,
    UnsupportedDocumentModificationError,
)
from app.services.document.extraction_models import DocumentEdit, DocumentEditAction


class ExcelModifier:
    """Apply validated, deterministic cell and worksheet edits to a copied workbook."""

    def modify(self, file_data: bytes, edits: tuple[DocumentEdit, ...]) -> bytes:
        workbook = load_workbook(BytesIO(file_data), data_only=False)
        try:
            for edit in edits:
                sheet = self._resolve_sheet(workbook, edit.locator.sheet, edit)
                self._apply(sheet, edit)
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    def _apply(self, sheet: Worksheet, edit: DocumentEdit) -> None:
        action = edit.action
        if action == DocumentEditAction.UPDATE_CELL:
            cell = self._resolve_cell(sheet, edit.locator.cell)
            cell.value = edit.value
        elif action in {DocumentEditAction.INSERT_ROWS, DocumentEditAction.DELETE_ROWS}:
            row = edit.locator.row
            if row is None:
                raise DocumentLocatorError("A row operation requires a row number.")
            amount = self._positive_amount(edit.options)
            if action == DocumentEditAction.INSERT_ROWS:
                sheet.insert_rows(row, amount)
            else:
                self._require_existing_row(sheet, row, amount)
                sheet.delete_rows(row, amount)
        elif action in {DocumentEditAction.INSERT_COLUMNS, DocumentEditAction.DELETE_COLUMNS}:
            column = edit.locator.column
            if column is None:
                raise DocumentLocatorError("A column operation requires a column number.")
            amount = self._positive_amount(edit.options)
            if action == DocumentEditAction.INSERT_COLUMNS:
                sheet.insert_cols(column, amount)
            else:
                if column + amount - 1 > sheet.max_column:
                    raise DocumentLocatorError(
                        f"Columns {column} through {column + amount - 1} do not exist."
                    )
                sheet.delete_cols(column, amount)
        elif action == DocumentEditAction.SORT_RANGE:
            self._sort_range(sheet, edit)
        elif action == DocumentEditAction.FILTER_ROWS:
            self._filter_rows(sheet, edit)
        elif action == DocumentEditAction.FORMAT_RANGE:
            self._format_range(sheet, edit)
        elif action == DocumentEditAction.FREEZE_PANES:
            coordinate = edit.locator.cell
            if not coordinate:
                raise DocumentLocatorError("Freeze panes requires a cell locator such as A2.")
            self._resolve_cell(sheet, coordinate)
            sheet.freeze_panes = coordinate.upper()
        elif action == DocumentEditAction.UPDATE_TABLE_CELL:
            if edit.locator.row is None or edit.locator.column is None:
                raise DocumentLocatorError("A cell edit requires row and column numbers.")
            if edit.locator.row > sheet.max_row or edit.locator.column > sheet.max_column:
                raise DocumentLocatorError("The requested worksheet cell does not exist.")
            sheet.cell(edit.locator.row, edit.locator.column).value = edit.value
        else:
            raise UnsupportedDocumentModificationError(
                f"The Excel edit '{action.value}' is not supported."
            )

    @staticmethod
    def _resolve_sheet(
        workbook, requested: str | None, edit: DocumentEdit | None = None
    ) -> Worksheet:
        if requested:
            matches = [
                sheet
                for sheet in workbook.worksheets
                if sheet.title.casefold() == requested.casefold()
            ]
            if not matches:
                raise DocumentLocatorError(f"I couldn't find a worksheet named '{requested}'.")
            return matches[0]
        if len(workbook.worksheets) == 1:
            return workbook.active

        candidates = ExcelModifier._candidate_sheets(workbook.worksheets, edit)
        if len(candidates) == 1:
            return candidates[0]
        if len(workbook.worksheets) != 1:
            names = ", ".join(sheet.title for sheet in workbook.worksheets[:8])
            raise DocumentLocatorError(
                f"This workbook has multiple worksheets ({names}). Please name the worksheet."
            )
        return workbook.active

    @staticmethod
    def _candidate_sheets(
        worksheets: list[Worksheet], edit: DocumentEdit | None
    ) -> list[Worksheet]:
        if edit is None:
            return []
        if edit.action == DocumentEditAction.FILTER_ROWS:
            expected = edit.options.get("equals", "active")
            return [
                sheet
                for sheet in worksheets
                if any(
                    ExcelModifier._values_equal(cell.value, expected)
                    for row in sheet.iter_rows(min_row=2)
                    for cell in row
                )
            ]
        if edit.action == DocumentEditAction.SORT_RANGE:
            requested_column = edit.options.get("column")
            if requested_column is None:
                return []
            expected = str(requested_column).strip().casefold()
            return [
                sheet
                for sheet in worksheets
                if any(
                    str(cell.value or "").strip().casefold() == expected
                    for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())
                )
            ]
        return []

    @staticmethod
    def _resolve_cell(sheet: Worksheet, coordinate: str | None):
        if not coordinate:
            raise DocumentLocatorError("A cell edit requires a cell locator such as B4.")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(coordinate.upper())
        except ValueError as error:
            raise DocumentLocatorError(f"'{coordinate}' is not a valid cell locator.") from error
        if min_col != max_col or min_row != max_row:
            raise DocumentLocatorError("This operation requires exactly one cell, not a range.")
        return sheet.cell(min_row, min_col)

    @staticmethod
    def _positive_amount(options: dict[str, object]) -> int:
        value = options.get("amount", 1)
        if not isinstance(value, int) or isinstance(value, bool) or value < 1:
            raise UnsupportedDocumentModificationError(
                "Row or column amount must be a positive integer."
            )
        return value

    @staticmethod
    def _require_existing_row(sheet: Worksheet, row: int, amount: int) -> None:
        if row + amount - 1 > sheet.max_row:
            raise DocumentLocatorError(f"Rows {row} through {row + amount - 1} do not exist.")

    def _sort_range(self, sheet: Worksheet, edit: DocumentEdit) -> None:
        bounds = self._range_bounds(sheet, edit.locator.cell_range)
        min_col, min_row, max_col, max_row = bounds
        header = bool(edit.options.get("header", True))
        sort_column = edit.options.get("column")
        column = self._column_in_range(sheet, min_row, min_col, max_col, sort_column)
        data_start = min_row + 1 if header else min_row
        if data_start > max_row:
            raise DocumentLocatorError("The requested range has no rows to sort.")
        self._reject_merged_range(sheet, bounds)
        snapshots = [
            [self._snapshot_cell(sheet.cell(row, col)) for col in range(min_col, max_col + 1)]
            for row in range(data_start, max_row + 1)
        ]
        key_index = column - min_col
        reverse = bool(edit.options.get("descending", False))
        snapshots.sort(key=lambda row: self._sort_key(row[key_index]["value"]), reverse=reverse)
        for row_offset, snapshot_row in enumerate(snapshots, start=data_start):
            for col_offset, snapshot in enumerate(snapshot_row, start=min_col):
                self._restore_cell(sheet.cell(row_offset, col_offset), snapshot)

    def _filter_rows(self, sheet: Worksheet, edit: DocumentEdit) -> None:
        bounds = self._range_bounds(sheet, edit.locator.cell_range)
        min_col, min_row, max_col, max_row = bounds
        header = bool(edit.options.get("header", True))
        filter_column = edit.options.get("column")
        expected = edit.options.get("equals", "active")
        if filter_column is None:
            candidates = [
                column
                for column in range(min_col, max_col + 1)
                if any(
                    self._values_equal(sheet.cell(row, column).value, expected)
                    for row in range(min_row + (1 if header else 0), max_row + 1)
                )
            ]
            if not candidates:
                raise DocumentLocatorError(f"I couldn't find a column containing '{expected}'.")
            if len(candidates) > 1:
                names = ", ".join(str(sheet.cell(min_row, column).value) for column in candidates)
                raise DocumentLocatorError(
                    f"More than one column contains '{expected}' ({names}). Please name the column."
                )
            column = candidates[0]
        else:
            column = self._column_in_range(sheet, min_row, min_col, max_col, filter_column)
        data_start = min_row + 1 if header else min_row
        matching_rows = [
            row
            for row in range(data_start, max_row + 1)
            if self._values_equal(sheet.cell(row, column).value, expected)
        ]
        if not matching_rows:
            raise DocumentLocatorError(
                f"I couldn't find records matching '{expected}' in the requested column."
            )
        self._reject_merged_range(sheet, bounds)
        snapshots = [
            [self._snapshot_cell(sheet.cell(row, col)) for col in range(min_col, max_col + 1)]
            for row in matching_rows
        ]
        for target_row, snapshot_row in enumerate(snapshots, start=data_start):
            for col, snapshot in zip(range(min_col, max_col + 1), snapshot_row, strict=True):
                self._restore_cell(sheet.cell(target_row, col), snapshot)
        delete_start = data_start + len(snapshots)
        delete_count = max_row - delete_start + 1
        if delete_count:
            sheet.delete_rows(delete_start, delete_count)
        sheet.auto_filter.ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{data_start + len(snapshots) - 1}"
        )

    def _format_range(self, sheet: Worksheet, edit: DocumentEdit) -> None:
        min_col, min_row, max_col, max_row = self._range_bounds(sheet, edit.locator.cell_range)
        horizontal = edit.options.get("alignment")
        vertical = edit.options.get("vertical_alignment")
        wrap = edit.options.get("wrap_text")
        border_requested = bool(edit.options.get("border", False))
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        for row in sheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        ):
            for cell in row:
                cell.alignment = copy(cell.alignment)
                if horizontal is not None:
                    cell.alignment = Alignment(
                        horizontal=str(horizontal).casefold(),
                        vertical=cell.alignment.vertical,
                        wrap_text=cell.alignment.wrap_text,
                    )
                if vertical is not None or wrap is not None:
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=(
                            str(vertical).casefold()
                            if vertical is not None
                            else cell.alignment.vertical
                        ),
                        wrap_text=(bool(wrap) if wrap is not None else cell.alignment.wrap_text),
                    )
                if "bold" in edit.options:
                    cell.font = copy(cell.font)
                    cell.font = Font(
                        name=cell.font.name,
                        size=cell.font.sz,
                        bold=bool(edit.options["bold"]),
                        italic=cell.font.italic,
                        color=cell.font.color,
                    )
                if "fill" in edit.options:
                    color = str(edit.options["fill"]).lstrip("#")
                    cell.fill = PatternFill("solid", fgColor=color)
                if border_requested:
                    cell.border = border
        if "column_width" in edit.options:
            width = float(str(edit.options["column_width"]))
            for column in range(min_col, max_col + 1):
                sheet.column_dimensions[get_column_letter(column)].width = width

    @staticmethod
    def _range_bounds(sheet: Worksheet, cell_range: str | None) -> tuple[int, int, int, int]:
        if not cell_range:
            if sheet.max_row < 1 or sheet.max_column < 1:
                raise DocumentLocatorError("The selected worksheet is empty.")
            return 1, 1, sheet.max_column, sheet.max_row
        try:
            bounds = range_boundaries(cell_range.upper())
        except ValueError as error:
            raise DocumentLocatorError(f"'{cell_range}' is not a valid cell range.") from error
        min_col, min_row, max_col, max_row = bounds
        if max_row > sheet.max_row or max_col > sheet.max_column:
            raise DocumentLocatorError(
                f"Range {cell_range} extends beyond worksheet '{sheet.title}'."
            )
        return bounds

    @staticmethod
    def _column_in_range(
        sheet: Worksheet,
        header_row: int,
        min_col: int,
        max_col: int,
        requested: object,
    ) -> int:
        if isinstance(requested, int) and not isinstance(requested, bool):
            column = requested
        elif isinstance(requested, str) and requested.strip():
            value = requested.strip()
            if value.isalpha() and len(value) <= 3:
                column = column_index_from_string(value.upper())
            else:
                matches = [
                    cell.column
                    for cell in sheet[header_row][min_col - 1 : max_col]
                    if str(cell.value or "").strip().casefold() == value.casefold()
                ]
                if not matches:
                    raise DocumentLocatorError(f"I couldn't find a column named '{value}'.")
                if len(matches) > 1:
                    raise DocumentLocatorError(
                        f"More than one column is named '{value}'. Please use a column letter."
                    )
                column = matches[0]
        else:
            raise DocumentLocatorError(
                "Sort and filter operations require a column name or number."
            )
        if not min_col <= column <= max_col:
            raise DocumentLocatorError("The requested column is outside the selected range.")
        return column

    @staticmethod
    def _reject_merged_range(sheet: Worksheet, bounds: tuple[int, int, int, int]) -> None:
        min_col, min_row, max_col, max_row = bounds
        if any(
            merged.min_col <= max_col
            and merged.max_col >= min_col
            and merged.min_row <= max_row
            and merged.max_row >= min_row
            for merged in sheet.merged_cells.ranges
        ):
            raise UnsupportedDocumentModificationError(
                "Sorting or filtering a range containing merged cells is not supported safely."
            )

    @staticmethod
    def _snapshot_cell(cell) -> dict[str, Any]:
        return {
            "value": cell.value,
            "style": copy(cell._style),
            "number_format": cell.number_format,
            "protection": copy(cell.protection),
            "alignment": copy(cell.alignment),
            "comment": copy(cell.comment),
            "hyperlink": copy(cell.hyperlink),
        }

    @staticmethod
    def _restore_cell(cell, snapshot: dict[str, Any]) -> None:
        cell.value = snapshot["value"]
        cell._style = copy(snapshot["style"])
        cell.number_format = snapshot["number_format"]
        cell.protection = copy(snapshot["protection"])
        cell.alignment = copy(snapshot["alignment"])
        cell.comment = copy(snapshot["comment"])
        cell._hyperlink = copy(snapshot["hyperlink"])

    @staticmethod
    def _sort_key(value: object) -> tuple[int, Decimal | str]:
        if value is None or (isinstance(value, str) and not value.strip()):
            return (2, "")
        # Numeric cells and numeric text must sort by value: 2.00 before 10.00.
        if not isinstance(value, bool):
            try:
                number = Decimal(str(value).strip())
                if number.is_finite():
                    return (0, number)
            except InvalidOperation:
                pass
        return (1, str(value).casefold())

    @staticmethod
    def _values_equal(actual: object, expected: object) -> bool:
        if isinstance(actual, str) and isinstance(expected, str):
            return actual.strip().casefold() == expected.strip().casefold()
        return actual == expected
