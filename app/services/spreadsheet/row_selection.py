"""Exact workbook row selection, independent of AI content generation."""

import re
from copy import copy
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, field_validator

from app.services.document.exceptions import DocumentLocatorError, InvalidDocumentError


class RowSelection(BaseModel):
    model_config = ConfigDict(extra="forbid")

    ids: list[str] = Field(min_length=1, max_length=200)
    column: str | None = Field(default=None, min_length=1, max_length=128)
    sheet_name: str | None = Field(default=None, min_length=1, max_length=128)

    @field_validator("ids")
    @classmethod
    def validate_ids(cls, values: list[str]) -> list[str]:
        if any(not value.strip() or len(value) > 128 for value in values):
            raise ValueError("Provide nonempty ID values of at most 128 characters.")
        return list(dict.fromkeys(value.strip() for value in values))

    @classmethod
    def from_instruction(
        cls, instruction: str, *, workbook_context: bool = False
    ) -> "RowSelection | None":
        instruction = re.sub(r"\b(id|data|row)['’]s\b", r"\1s", instruction, flags=re.I)
        instruction = instruction.replace("\r\n", "\n")
        if re.search(
            r"\b(?:add|append|insert|update|modify|edit)\b", instruction.split("\n")[0], re.I
        ):
            return None
        if re.search(r"(?<!\w)\d+[ \t]*(?:to|through|[-–])[ \t]*\d+", instruction, re.I):
            return None
        if re.search(
            r"\b(?:pdf|docx|word|pptx|powerpoint|csv|image|photo|picture|"
            r"sort|summarize|summarise|aggregate|calculate)\b",
            instruction,
            re.I,
        ):
            # Source conversion and additional transformations need a composed
            # plan; selecting rows alone must not replace the user's final goal.
            return None
        if not (
            re.search(r"\b(?:excel|xlsx|workbook|spreadsheet)\b", instruction, re.I)
            or workbook_context
        ):
            return None
        if not (
            re.search(r"\b(?:rows?|records?|datas?|ids?|particular)\b", instruction, re.I)
            and re.search(
                r"\b(?:find|filter|select|extract|keep|return|read|get|give|show|fetch|create|generate|make|eduthu|kudu|thaa)\b",
                instruction,
                re.I,
            )
        ):
            return None
        # Read a labelled ID list, never all numbers in the prompt (which may
        # include numbered requirements, row counts, dates or column numbers).
        value = r"""(?:"[^"\n]{1,128}"|'[^'\n]{1,128}'|(?=[\w-]*\d)[A-Za-z0-9][A-Za-z0-9_-]*)"""
        separator = r"(?:[ \t]*(?:,\s*(?:(?:and|or)\s+)?|\b(?:and|or)\b\s*|\n\s*[-*•]?[ \t]*))"
        pattern = (
            r"\bIDs?\b\s*(?:(?:values?\s*)?(?::|=)|(?:exactly\s+)?match(?:es)?|are|in)?"
            rf"\s*[-*•]?[ \t]*(?P<values>{value}(?:{separator}{value})*)"
        )
        matches = list(re.finditer(pattern, instruction, re.I))
        ids = [
            token.strip("\"'") for match in matches for token in re.findall(value, match["values"])
        ]
        if not ids:
            # Also accept values before "IDs" and a separate list following
            # "particular"/"data for". Never collect arbitrary prompt numbers.
            before = re.search(
                rf"(?P<values>{value}(?:{separator}{value})*)\s+IDs?\b", instruction, re.I
            )
            labelled = re.search(
                rf"\b(?:particular|data(?:s)?\s+(?:for|of)|rows?\s+for)\s*:?\s*"
                rf"(?P<values>{value}(?:{separator}{value})*)",
                instruction,
                re.I,
            )
            match = before or labelled
            if match:
                ids = [token.strip("\"'") for token in re.findall(value, match["values"])]
        if not ids:
            return None
        # Restrict deterministic recognition to inclusion. Other transformations
        # still go through the existing planner.
        if re.search(r"\b(?:exclude|except|remove|delete)\b.{0,35}\bIDs?\b", instruction, re.I):
            return None
        sheet = re.search(r"""\b(?:sheet|worksheet)\s+["']([^"']+)["']""", instruction, re.I)
        column = re.search(r"""\bcolumn\s+["']([^"']+)["']""", instruction, re.I)
        return cls(
            ids=ids, column=column[1] if column else None, sheet_name=sheet[1] if sheet else None
        )


@dataclass(frozen=True)
class SelectedRows:
    file_data: bytes
    counts: dict[str, int]
    missing_ids: tuple[str, ...]
    formula_values_copied: bool

    def markdown(self) -> str:
        """Show actual selected cell values, without a generative summary."""

        def escape(value: object) -> str:
            text = "" if value is None else str(value)
            return (
                re.sub(r"([\\`*_{}\[\]<>()#|!])", r"\\\1", text)
                .replace("\n", " ")
                .replace("\r", " ")
            )

        lines = [self.summary]
        book = load_workbook(BytesIO(self.file_data), data_only=True)
        displayed = 0
        try:
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows)
                for index, row in enumerate(rows, 1):
                    if displayed >= 20:
                        lines.append(
                            "Only the first 20 matching records are shown here. The filtered workbook contains all matches."
                        )
                        return "\n\n".join(lines)
                    lines.append(f"### {escape(sheet.title)} — Record {index}")
                    table = ["| Column | Value |", "| --- | --- |"]
                    table.extend(
                        f"| {escape(header)} | {escape(value)} |"
                        for header, value in zip(headers, row)
                    )
                    lines.append("\n".join(table))
                    displayed += 1
            return "\n\n".join(lines)
        finally:
            book.close()

    @property
    def summary(self) -> str:
        details = ", ".join(f"{name}: {count}" for name, count in self.counts.items())
        message = f"Matched {sum(self.counts.values())} rows ({details})."
        message += (
            " IDs not found: " + ", ".join(self.missing_ids) + "."
            if self.missing_ids
            else " All requested IDs were found."
        )
        if self.formula_values_copied:
            message += " Formula cells contain their saved calculated values."
        return message


class WorkbookRowSelector:
    """Copy complete matching records; source bytes are never modified."""

    @staticmethod
    def _copy_style(source, target) -> None:
        # Style indices belong to one workbook; register the actual components
        # in the output instead of copying indices into a different style table.
        for name in ("font", "fill", "border", "alignment", "protection"):
            setattr(target, name, copy(getattr(source, name)))
        target.number_format = source.number_format

    @staticmethod
    def _key(value: object) -> str | None:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _is_id_header(value: object) -> bool:
        text = re.sub(r"([a-z])([A-Z])", r"\1 \2", str(value or ""))
        return bool(re.search(r"(?:^|[\s_*-])(?:id|ids|identifier)\s*\*?$", text, re.I))

    def _locate(self, sheet, requested: set[str], column: str | None) -> tuple[int, int] | None:
        headers = []
        for row in sheet.iter_rows(max_row=min(sheet.max_row, 50)):
            for cell in row:
                is_header = (
                    str(cell.value or "").strip().casefold() == column.strip().casefold()
                    if column
                    else self._is_id_header(cell.value)
                )
                if is_header:
                    headers.append((cell.row, cell.column))
            if headers:
                break
        if len(headers) == 1:
            return headers[0]
        if len(headers) > 1:
            matching = [
                (row, col)
                for row, col in headers
                if any(
                    self._key(sheet.cell(index, col).value) in requested
                    for index in range(row + 1, sheet.max_row + 1)
                )
            ]
            if len(matching) == 1:
                return matching[0]
            raise DocumentLocatorError(
                f"Worksheet '{sheet.title}' has multiple ID columns. Specify the column name."
            )
        if column:
            return None
        # A nonstandard ID label can be identified by exact matches in one
        # column. Ambiguous matches must not silently choose an amount column.
        candidates = {
            cell.column
            for row in sheet.iter_rows(min_row=2)
            for cell in row
            if self._key(cell.value) in requested
        }
        if len(candidates) > 1:
            raise DocumentLocatorError(
                f"Worksheet '{sheet.title}' matches IDs in multiple columns. Specify the ID column name."
            )
        return (1, next(iter(candidates))) if candidates else None

    def select(self, file_data: bytes, selection: RowSelection) -> SelectedRows:
        workbook = load_workbook(BytesIO(file_data), data_only=False)
        values = load_workbook(BytesIO(file_data), data_only=True)
        output = Workbook()
        output.remove(output.active)
        output.loaded_theme = workbook.loaded_theme
        output.epoch = workbook.epoch
        try:
            requested = set(selection.ids)
            found: set[str] = set()
            counts: dict[str, int] = {}
            formula_values = False
            worksheets = [
                sheet
                for sheet in workbook
                if selection.sheet_name is None
                or sheet.title.casefold() == selection.sheet_name.casefold()
            ]
            if not worksheets:
                raise DocumentLocatorError(f"Worksheet '{selection.sheet_name}' was not found.")
            located = []
            for sheet in worksheets:
                if sheet.max_row * sheet.max_column > 2_000_000:
                    raise InvalidDocumentError("This worksheet is too large to filter safely.")
                location = self._locate(values[sheet.title], requested, selection.column)
                if location is None:
                    continue
                header_row, column = location
                rows = []
                for row in range(header_row + 1, sheet.max_row + 1):
                    id_value = values[sheet.title].cell(row, column).value
                    if sheet.cell(row, column).data_type == "f" and id_value is None:
                        raise DocumentLocatorError(
                            f"The ID column in '{sheet.title}' contains a formula with no saved result. "
                            "Calculate and save the source workbook in Excel, then upload it again."
                        )
                    key = self._key(id_value)
                    if key in requested:
                        rows.append(row)
                        found.add(key)
                located.append((sheet, header_row, rows))
            if not located:
                raise DocumentLocatorError(
                    "I couldn't identify an ID column. Specify its exact column name."
                )
            any_matches = any(rows for _, _, rows in located)
            for sheet, header_row, rows in located:
                if any_matches and not rows:
                    continue
                target = output.create_sheet(sheet.title)
                counts[sheet.title] = len(rows)
                for name, dimension in sheet.column_dimensions.items():
                    target_dimension = target.column_dimensions[name]
                    for attribute in (
                        "width",
                        "hidden",
                        "min",
                        "max",
                        "bestFit",
                        "outlineLevel",
                        "collapsed",
                    ):
                        setattr(target_dimension, attribute, getattr(dimension, attribute))
                    self._copy_style(dimension, target_dimension)
                for target_row, source_row in enumerate([header_row, *rows], 1):
                    target.row_dimensions[target_row].height = sheet.row_dimensions[
                        source_row
                    ].height
                    for column in range(1, sheet.max_column + 1):
                        source = sheet.cell(source_row, column)
                        value = source.value
                        if source.data_type == "f":
                            value = values[sheet.title].cell(source_row, column).value
                            if value is None:
                                raise DocumentLocatorError(
                                    f"Cell {sheet.title}!{source.coordinate} has no saved formula result. "
                                    "Calculate and save the source workbook in Excel, then upload it again."
                                )
                            formula_values = True
                        cell = target.cell(target_row, column, value)
                        # Preserve literal text beginning '=' as text, not a new formula.
                        cell.data_type = values[sheet.title].cell(source_row, column).data_type
                        self._copy_style(source, cell)
                        cell.comment = copy(source.comment)
                        if source.hyperlink:
                            cell.hyperlink = source.hyperlink.target
                target.freeze_panes = "A2"
                target.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{len(rows) + 1}"
            buffer = BytesIO()
            output.save(buffer)
            return SelectedRows(
                buffer.getvalue(),
                counts,
                tuple(value for value in selection.ids if value not in found),
                formula_values,
            )
        finally:
            workbook.close()
            values.close()
            output.close()
