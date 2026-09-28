"""Render complete workbook records grouped by numeric price, without model rewriting."""

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from html import escape
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    Flowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.document.exceptions import InvalidDocumentError
from app.utils.document_output import positive_instruction, requests_pdf_output

NAVY = colors.HexColor("#1F3A5F")
GOLD = colors.HexColor("#C9A227")
GREY = colors.HexColor("#F4F6F8")


class PriceListPdfRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(default="Catering Menu", min_length=1, max_length=120)
    filename: str = Field(default="Catering_Menu_Price_Wise.pdf", max_length=180)
    sheet_name: str | None = None
    expected_count: int | None = Field(default=None, ge=1, le=100000)
    max_groups: int = Field(default=60, ge=1, le=60)

    @classmethod
    def from_instruction(cls, instruction: str) -> "PriceListPdfRequest | None":
        text = positive_instruction(instruction)
        if not requests_pdf_output(text) or not re.search(
            r"\b(?:group\w*[^.\n]{0,90}\b(?:price|rate)|(?:price|rate)[ -]wise|same\s+(?:price|rate))\b",
            text,
            re.I,
        ):
            return None
        filename = re.search(r"\b([\w-]+\.pdf)\b", text, re.I)
        title = re.search(r'\b(?:cover|title)\s*:\s*["“]([^"”\n]+)', text, re.I)
        count = re.search(r"\b(?:all|total items\s*:|has)\s*(\d+)\s*(?:rows|items)?\b", text, re.I)
        sheet = re.search(r'\b(?:from\s+)?sheet\s+["`]([^"`]+)["`]', text, re.I)
        return cls(
            title=title.group(1) if title else "Catering Menu",
            filename=filename.group(1) if filename else "Catering_Menu_Price_Wise.pdf",
            expected_count=int(count.group(1)) if count else None,
            sheet_name=sheet.group(1) if sheet else None,
        )


@dataclass(frozen=True)
class PriceListData:
    groups: dict[Decimal | None, list[tuple[str, str, str, str]]]
    count: int
    note: str

    @staticmethod
    def label(price: Decimal | None) -> str:
        return f"₹ {price:.2f}" if price is not None else "Price not provided"

    @property
    def summary(self) -> str:
        rows = ["| Price | Item Count |", "| --- | ---: |"]
        rows.extend(
            f"| {self.label(price)} | {len(items)} |" for price, items in self.groups.items()
        )
        rows.append(f"| **Total** | **{self.count}** |")
        return self.note + "\n\n" + "\n".join(rows)


def clean_price(value: object) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    cleaned = re.sub(r"₹|\bRs\.?|[,\s]", "", str(value), flags=re.I)
    try:
        number = Decimal(cleaned)
        if not number.is_finite():
            raise InvalidOperation
        return number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation as error:
        raise InvalidDocumentError(
            f"Invalid Price value: {value!s}. Supply a numeric price or leave it blank."
        ) from error


class PriceListPdf:
    def read(self, data: bytes, request: PriceListPdfRequest) -> PriceListData:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            candidates = []
            for sheet in workbook.worksheets:
                if request.sheet_name and sheet.title.casefold() != request.sheet_name.casefold():
                    continue
                for index, row in enumerate(sheet.iter_rows(max_row=25, values_only=True), 1):
                    keys = [re.sub(r"[\W_]+", "", str(value or "").casefold()) for value in row]
                    required = ("code", "name", "category", "status", "price")
                    if all(keys.count(key) == 1 for key in required):
                        candidates.append((sheet, index, [keys.index(key) for key in required]))
                        break
            if len(candidates) != 1:
                raise InvalidDocumentError(
                    "Name one worksheet with unique Code, Name, Category, Status and Price headers."
                )
            sheet, header, columns = candidates[0]
            groups: dict[Decimal | None, list[tuple[str, str, str, str]]] = defaultdict(list)
            count = 0
            raw_rows = formulas[sheet.title].iter_rows(min_row=header + 1)
            for row, raw in zip(
                sheet.iter_rows(min_row=header + 1, values_only=True), raw_rows, strict=True
            ):
                if not any(cell.value is not None for cell in raw):
                    continue
                if any(raw[col].data_type == "f" and row[col] is None for col in columns):
                    raise InvalidDocumentError(
                        "A formula has no saved value. Recalculate and save the source workbook before exporting."
                    )
                values = [row[col] for col in columns]
                if not values[0] and not values[1]:
                    raise InvalidDocumentError(
                        f"Row {raw[0].row} has data but no Code or Name. Check the item table before exporting."
                    )
                price = clean_price(values[4])
                code, name, category, status = [
                    "" if value is None else str(value) for value in values[:4]
                ]
                groups[price].append((code, name, category, status))
                count += 1
                if count > 10000:
                    raise InvalidDocumentError(
                        "This PDF export supports up to 10,000 items at a time."
                    )
            if not count:
                raise InvalidDocumentError("The worksheet contains no item rows.")
            if len(groups) > request.max_groups:
                raise InvalidDocumentError(
                    f"Found {len(groups)} price groups, above the {request.max_groups}-group limit. Check the Price column before exporting."
                )
            ordered = dict(
                sorted(groups.items(), key=lambda item: (item[0] is None, item[0] or Decimal(0)))
            )
            note = f"Grouped all {count} source items by cleaned Price only."
            if request.expected_count and count != request.expected_count:
                note += f" The workbook contains {count} items, not the requested {request.expected_count}; no missing items were invented."
            if None in groups:
                note += f" {len(groups[None])} items have blank prices and are listed separately, not as zero."
            return PriceListData(ordered, count, note)
        finally:
            workbook.close()
            formulas.close()

    def generate(
        self, data: bytes, request: PriceListPdfRequest, *, today: date | None = None
    ) -> tuple[bytes, str]:
        records = self.read(data, request)
        regular, bold = self._fonts()
        body = ParagraphStyle(
            "PriceListBody", fontName=regular, fontSize=10, leading=14, textColor=NAVY
        )
        heading = ParagraphStyle(
            "PriceListHeading", parent=body, fontName=bold, fontSize=18, leading=23, spaceAfter=16
        )
        white = ParagraphStyle("PriceListWhite", parent=body, fontName=bold, textColor=colors.white)
        buffer = BytesIO()
        # SimpleDocTemplate frames reserve six points of padding on either side.
        width = A4[0] - 36 * mm - 12
        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=23 * mm,
            bottomMargin=20 * mm,
            title=request.title,
        )
        story = [Spacer(1, 1), PageBreak(), Paragraph("Price summary", heading)]
        summary = [[Paragraph("Price", white), Paragraph("No. of Items", white)]]
        summary.extend(
            [
                [Paragraph(records.label(price), body), Paragraph(str(len(items)), body)]
                for price, items in records.groups.items()
            ]
        )
        summary.append([Paragraph("Grand total", body), Paragraph(str(records.count), body)])
        table = Table(summary, colWidths=[width * 0.7, width * 0.3], repeatRows=1)
        table.setStyle(self._table_style())
        table.setStyle(TableStyle([("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E0E8F0"))]))
        story.extend([table, Spacer(1, 14), Paragraph(escape(records.note), body), PageBreak()])
        for price, items in records.groups.items():
            # Put the section band, column headers and records in one splittable table.
            # ReportLab repeats both headers and never splits off headers without a data row.
            rows = [
                [
                    Paragraph(
                        f"{records.label(price)}  •  {len(items)} {'item' if len(items) == 1 else 'items'}",
                        white,
                    ),
                    "",
                    "",
                    "",
                ],
                [Paragraph(label, white) for label in ("Code", "Name", "Category", "Status")],
            ]
            rows.extend(
                [
                    [
                        Paragraph(escape(code), body),
                        Paragraph(escape(name), body),
                        Paragraph(escape(category), body),
                        _StatusBadge(status, regular),
                    ]
                    for code, name, category, status in items
                ]
            )
            table = Table(
                rows,
                colWidths=[width * 0.15, width * 0.43, width * 0.25, width * 0.17],
                repeatRows=2,
                hAlign="LEFT",
            )
            table.setStyle(self._table_style())
            table.setStyle(
                TableStyle(
                    [
                        ("SPAN", (0, 0), (-1, 0)),
                        ("BACKGROUND", (0, 1), (-1, 1), NAVY),
                        ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, GREY]),
                    ]
                )
            )
            story.extend([table, Spacer(1, 15)])

        def cover(canvas, doc):
            canvas.saveState()
            canvas.setFillColor(NAVY)
            canvas.rect(0, 0, *A4, stroke=0, fill=1)
            canvas.setFillColor(GOLD)
            canvas.setFont(bold, 34)
            # Wrap a caller-supplied long title without reducing body readability.
            title_style = ParagraphStyle(
                "CoverTitle", fontName=bold, fontSize=34, leading=42, textColor=GOLD, alignment=1
            )
            title = Paragraph(escape(request.title), title_style)
            _, height = title.wrap(width, 200)
            title.drawOn(canvas, 18 * mm, 520 - height)
            canvas.setFont(regular, 18)
            canvas.drawCentredString(A4[0] / 2, 400, "Price-wise Item List")
            canvas.setStrokeColor(GOLD)
            canvas.line(65 * mm, 370, A4[0] - 65 * mm, 370)
            canvas.setFont(bold, 14)
            canvas.drawCentredString(A4[0] / 2, 330, f"Total Items: {records.count}")
            canvas.setFont(regular, 11)
            canvas.drawCentredString(A4[0] / 2, 300, (today or date.today()).strftime("%d %B %Y"))
            canvas.restoreState()

        document.build(
            story,
            onFirstPage=cover,
            canvasmaker=lambda *args, **kwargs: _NumberedCanvas(
                *args, title=request.title, font=regular, **kwargs
            ),
        )
        return buffer.getvalue(), records.summary

    @staticmethod
    def _fonts() -> tuple[str, str]:
        if "PriceListSans" not in pdfmetrics.getRegisteredFontNames():
            pairs = [
                (
                    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                ),
                (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
            ]
            pair = next((pair for pair in pairs if all(path.is_file() for path in pair)), None)
            if pair is None:
                raise InvalidDocumentError(
                    "Install DejaVu Sans fonts to export Unicode prices to PDF."
                )
            for name, path in zip(("PriceListSans", "PriceListSansBold"), pair, strict=True):
                pdfmetrics.registerFont(TTFont(name, str(path)))
        return "PriceListSans", "PriceListSansBold"

    @staticmethod
    def _table_style() -> TableStyle:
        return TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GREY]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 9),
                ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E2E7EC")),
            ]
        )


class _StatusBadge(Flowable):
    def __init__(self, text: str, font: str):
        super().__init__()
        self.text, self.font = text, font
        self.width = pdfmetrics.stringWidth(text, font, 9) + 14
        self.height = 19

    def wrap(self, available_width, available_height):
        if self.width > available_width:
            raise InvalidDocumentError(
                "A Status value is too long for the PDF badge; use a short status label."
            )
        return self.width, self.height

    def draw(self):
        active = self.text.casefold() == "active"
        self.canv.setFillColor(colors.HexColor("#E4F3E8" if active else "#E8EBEF"))
        self.canv.roundRect(0, 0, self.width, self.height, 7, stroke=0, fill=1)
        self.canv.setFillColor(colors.HexColor("#22633C" if active else "#505966"))
        self.canv.setFont(self.font, 9)
        self.canv.drawString(7, 6, self.text)


class _NumberedCanvas(Canvas):
    def __init__(self, *args, title: str, font: str, **kwargs):
        super().__init__(*args, **kwargs)
        self._pages: list[dict] = []
        self._title, self._font = title, font

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        count = len(self._pages)
        for state in self._pages:
            self.__dict__.update(state)
            if self._pageNumber > 1:
                self.setFillColor(NAVY)
                self.setFont(self._font, 10)
                self.drawString(18 * mm, A4[1] - 14 * mm, self._title)
                self.setStrokeColor(colors.HexColor("#D8E0E8"))
                self.line(18 * mm, A4[1] - 17 * mm, A4[0] - 18 * mm, A4[1] - 17 * mm)
                self.setFont(self._font, 9)
                self.drawRightString(
                    A4[0] - 18 * mm, 11 * mm, f"Page {self._pageNumber} of {count}"
                )
            super().showPage()
        super().save()
