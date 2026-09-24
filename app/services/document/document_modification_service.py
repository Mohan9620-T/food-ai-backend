from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import NoReturn, cast

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, range_boundaries
from pptx import Presentation
from pypdf import PdfReader

from app.services.document.document_generation_service import DocumentGenerationService
from app.services.document.document_operation_registry import DocumentType, Fidelity
from app.services.document.document_reading_service import DocumentReadingService
from app.services.document.document_validation_service import DocumentValidationService
from app.services.document.exceptions import (
    DocumentChangeVerificationError,
    InvalidDocumentError,
    UnsupportedDocumentModificationError,
)
from app.services.document.extraction_models import (
    DocumentEdit,
    DocumentEditAction,
    DocumentMetadata,
    DocumentRepresentation,
)
from app.services.pdf.pdf_modifier import PdfModifier
from app.services.powerpoint.pptx_modifier import PptxModifier
from app.services.spreadsheet.excel_modifier import ExcelModifier
from app.services.word.docx_modifier import DocxModifier


@dataclass(frozen=True)
class ModifiedDocument:
    file_data: bytes
    filename: str
    content_type: str
    document_type: DocumentType
    fidelity: Fidelity
    summary: str
    fidelity_note: str | None = None


class DocumentModificationService:
    """Dispatch safe targeted edits and independently verify the resulting document."""

    def __init__(
        self,
        *,
        reader: DocumentReadingService | None = None,
        validator: DocumentValidationService | None = None,
        docx_modifier: DocxModifier | None = None,
        excel_modifier: ExcelModifier | None = None,
        pptx_modifier: PptxModifier | None = None,
        pdf_modifier: PdfModifier | None = None,
    ) -> None:
        self.reader = reader or DocumentReadingService()
        self.validator = validator or DocumentValidationService()
        self.docx_modifier = docx_modifier or DocxModifier()
        self.excel_modifier = excel_modifier or ExcelModifier()
        self.pptx_modifier = pptx_modifier or PptxModifier()
        self.pdf_modifier = pdf_modifier or PdfModifier()

    def modify(
        self,
        file_data: bytes,
        filename: str,
        edits: tuple[DocumentEdit, ...],
    ) -> ModifiedDocument:
        if not edits:
            raise InvalidDocumentError("At least one targeted document edit is required.")
        document_type = self._document_type(filename)
        original_snapshot = bytes(file_data)
        self.validator.validate(file_data, document_type)
        self._ensure_supported_edits(document_type, edits)
        if document_type == DocumentType.XLSX:
            edits = self._resolve_xlsx_sheet_locators(file_data, edits)
        before = self._read_representation(file_data, filename, document_type)
        output = self._dispatch(document_type, file_data, edits)
        if file_data != original_snapshot:
            raise DocumentChangeVerificationError("The original document changed in memory.")
        self.validator.validate(output, document_type)
        after = self._read_representation(output, filename, document_type)
        self._verify(before, after, edits, document_type)
        self._verify_requested_formatting(output, document_type, edits)
        output_filename = self._output_filename(filename, document_type)
        assessment = self.validator.fidelity_for_verified_modification(document_type)
        return ModifiedDocument(
            file_data=output,
            filename=output_filename,
            content_type=DocumentGenerationService.MIME_TYPES[document_type],
            document_type=document_type,
            fidelity=assessment.fidelity,
            summary=(
                f"Applied {len(edits)} targeted edit{'s' if len(edits) != 1 else ''}."
                + (f" {assessment.note}" if assessment.note else "")
            ),
            fidelity_note=assessment.note,
        )

    def _resolve_xlsx_sheet_locators(
        self, file_data: bytes, edits: tuple[DocumentEdit, ...]
    ) -> tuple[DocumentEdit, ...]:
        """Resolve each implicit worksheet locator uniquely before any edit is applied."""
        workbook = load_workbook(BytesIO(file_data), data_only=False)
        try:
            resolved: list[DocumentEdit] = []
            for edit in edits:
                if edit.locator.sheet is not None:
                    resolved.append(edit)
                    continue
                sheet = self.excel_modifier._resolve_sheet(workbook, None, edit)
                locator = edit.locator.model_copy(update={"sheet": sheet.title})
                resolved.append(edit.model_copy(update={"locator": locator}))
            return tuple(resolved)
        finally:
            workbook.close()

    def merge_pdfs(self, files: tuple[tuple[bytes, str], ...]) -> ModifiedDocument:
        if any(self._document_type(filename) != DocumentType.PDF for _, filename in files):
            raise UnsupportedDocumentModificationError("Only PDF files can be merged together.")
        snapshots = tuple(bytes(data) for data, _ in files)
        for data, _ in files:
            self.validator.validate(data, DocumentType.PDF)
        output = self.pdf_modifier.merge(tuple(data for data, _ in files))
        self.validator.validate(output, DocumentType.PDF)
        if tuple(data for data, _ in files) != snapshots:
            raise DocumentChangeVerificationError("An original PDF changed during merging.")
        return ModifiedDocument(
            file_data=output,
            filename="merged_updated.pdf",
            content_type=DocumentGenerationService.MIME_TYPES[DocumentType.PDF],
            document_type=DocumentType.PDF,
            fidelity=Fidelity.FULL,
            summary=f"Merged {len(files)} PDF files without changing the originals.",
        )

    def _dispatch(
        self,
        document_type: DocumentType,
        file_data: bytes,
        edits: tuple[DocumentEdit, ...],
    ) -> bytes:
        if document_type == DocumentType.DOCX:
            return self.docx_modifier.modify(file_data, edits)
        if document_type == DocumentType.XLSX:
            return self.excel_modifier.modify(file_data, edits)
        if document_type == DocumentType.PPTX:
            return self.pptx_modifier.modify(file_data, edits)
        if document_type == DocumentType.PDF:
            return self.pdf_modifier.modify(file_data, edits)
        raise UnsupportedDocumentModificationError(
            f"Targeted modification is not supported for {document_type.value.upper()} files."
        )

    @staticmethod
    def _ensure_supported_edits(
        document_type: DocumentType, edits: tuple[DocumentEdit, ...]
    ) -> None:
        if document_type == DocumentType.PDF and any(
            edit.action
            not in {
                DocumentEditAction.REMOVE_PDF_PAGES,
                DocumentEditAction.EXTRACT_PDF_PAGES,
                DocumentEditAction.REORDER_PDF_PAGES,
            }
            for edit in edits
        ):
            raise UnsupportedDocumentModificationError(
                "Arbitrary visual PDF editing is not supported reliably. "
                "A reconstruction workflow (PDF to Word, edit, then PDF) can be used once "
                "the required conversion path is available."
            )

    def _read_representation(
        self, file_data: bytes, filename: str, document_type: DocumentType
    ) -> DocumentRepresentation:
        if document_type != DocumentType.PDF:
            return self.reader.read(file_data, filename)
        reader = PdfReader(BytesIO(file_data))
        text = "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        return DocumentRepresentation(
            document_type=DocumentType.PDF,
            text=text,
            metadata=DocumentMetadata(page_count=len(reader.pages)),
            coverage=f"{len(reader.pages)} of {len(reader.pages)} pages",
        )

    def _verify(
        self,
        before: DocumentRepresentation,
        after: DocumentRepresentation,
        edits: tuple[DocumentEdit, ...],
        document_type: DocumentType,
    ) -> None:
        if before.document_type != after.document_type:
            self._fail()
        if document_type == DocumentType.PDF:
            self._verify_pdf(before, after, edits)
            return

        content_actions = {
            DocumentEditAction.REPLACE_TEXT,
            DocumentEditAction.ADD_PARAGRAPH,
            DocumentEditAction.REMOVE_PARAGRAPH,
            DocumentEditAction.UPDATE_TABLE_CELL,
            DocumentEditAction.UPDATE_CELL,
            DocumentEditAction.INSERT_ROWS,
            DocumentEditAction.DELETE_ROWS,
            DocumentEditAction.INSERT_COLUMNS,
            DocumentEditAction.DELETE_COLUMNS,
            DocumentEditAction.SORT_RANGE,
            DocumentEditAction.FILTER_ROWS,
        }
        if not any(edit.action in content_actions for edit in edits):
            if self._content_signature(before) != self._content_signature(after):
                self._fail()
            return

        if all(edit.action == DocumentEditAction.REPLACE_TEXT for edit in edits):
            expected = before.text
            for edit in edits:
                old = (edit.locator.text or edit.locator.heading or "").strip()
                expected = expected.replace(old, "" if edit.value is None else str(edit.value), 1)
            if after.text != expected or before.tables != after.tables:
                self._fail()
            return

        if document_type == DocumentType.XLSX:
            self._verify_xlsx(before, after, edits)
            return

        self._verify_block_edits(before, after, edits)

    def _verify_block_edits(
        self,
        before: DocumentRepresentation,
        after: DocumentRepresentation,
        edits: tuple[DocumentEdit, ...],
    ) -> None:
        expected = [
            {
                "type": block.type,
                "text": block.text,
                "cells": [list(row) for row in block.cells],
                "location": block.location,
            }
            for block in before.blocks
        ]
        for edit in edits:
            if edit.action == DocumentEditAction.REPLACE_TEXT:
                old = (edit.locator.text or edit.locator.heading or "").strip()
                matches = [
                    item
                    for item in expected
                    if isinstance(item["text"], str) and old in item["text"]
                ]
                if len(matches) != 1:
                    self._fail()
                matches[0]["text"] = str(matches[0]["text"]).replace(
                    old, "" if edit.value is None else str(edit.value), 1
                )
            elif edit.action == DocumentEditAction.UPDATE_TABLE_CELL:
                table_index = edit.locator.table_index
                row = edit.locator.row
                column = edit.locator.column
                tables = [item for item in expected if item["cells"]]
                if not table_index or not row or not column or table_index > len(tables):
                    self._fail()
                cells_value = tables[table_index - 1]["cells"]
                if not isinstance(cells_value, list) or row > len(cells_value):
                    self._fail()
                cells = cast(list[list[str]], cells_value)
                selected_row = cells[row - 1]
                if column > len(selected_row):
                    self._fail()
                selected_row[column - 1] = "" if edit.value is None else str(edit.value)
            elif edit.action == DocumentEditAction.REMOVE_PARAGRAPH:
                indexes = self._matching_block_indexes(expected, edit)
                if len(indexes) != 1:
                    self._fail()
                expected.pop(indexes[0])
            elif edit.action == DocumentEditAction.ADD_PARAGRAPH:
                indexes = self._matching_block_indexes(expected, edit)
                if len(indexes) != 1:
                    self._fail()
                inserted = {
                    "type": expected[indexes[0]]["type"],
                    "text": "" if edit.value is None else str(edit.value),
                    "cells": [],
                    "location": None,
                }
                expected.insert(indexes[0] + 1, inserted)
        actual = [
            {
                "type": block.type,
                "text": block.text,
                "cells": [list(row) for row in block.cells],
                "location": block.location,
            }
            for block in after.blocks
        ]
        if len(actual) != len(expected):
            self._fail()
        for expected_block, actual_block in zip(expected, actual, strict=True):
            if (
                expected_block["type"] != actual_block["type"]
                or expected_block["text"] != actual_block["text"]
                or expected_block["cells"] != actual_block["cells"]
            ):
                self._fail()

    @staticmethod
    def _matching_block_indexes(blocks: list[dict[str, object]], edit: DocumentEdit) -> list[int]:
        if edit.locator.paragraph_index is not None:
            return [
                index
                for index, block in enumerate(blocks)
                if getattr(block["location"], "paragraph_index", None)
                == edit.locator.paragraph_index
                and not block["cells"]
            ]
        text = (edit.locator.text or edit.locator.heading or "").strip()
        return [
            index
            for index, block in enumerate(blocks)
            if not block["cells"] and block["text"] == text
        ]

    def _verify_xlsx(
        self,
        before: DocumentRepresentation,
        after: DocumentRepresentation,
        edits: tuple[DocumentEdit, ...],
    ) -> None:
        expected = {table.name: [list(row) for row in table.rows] for table in before.tables}
        locations = {
            table.name: table.location.cell_range if table.location else None
            for table in before.tables
        }
        for edit in edits:
            sheet = edit.locator.sheet
            if sheet is None:
                if len(expected) != 1:
                    self._fail()
                sheet = next(iter(expected))
            if sheet not in expected:
                self._fail()
            rows = expected[sheet]
            start_col = start_row = 1
            if locations[sheet]:
                start_col, start_row, _, _ = range_boundaries(str(locations[sheet]))
            if edit.action in {
                DocumentEditAction.UPDATE_CELL,
                DocumentEditAction.UPDATE_TABLE_CELL,
            }:
                if edit.locator.cell:
                    column, row, _, _ = range_boundaries(edit.locator.cell)
                else:
                    row = edit.locator.row or 0
                    column = edit.locator.column or 0
                row -= start_row - 1
                column -= start_col - 1
                if row < 1 or column < 1 or row > len(rows) or column > len(rows[row - 1]):
                    self._fail()
                rows[row - 1][column - 1] = "" if edit.value is None else str(edit.value)
            elif edit.action == DocumentEditAction.FILTER_ROWS:
                if not rows:
                    self._fail()
                header = bool(edit.options.get("header", True))
                expected_value = str(edit.options.get("equals", "active")).strip().casefold()
                column = self._verification_column(rows, edit.options.get("column"), expected_value)
                prefix = rows[:1] if header else []
                data = rows[1:] if header else rows
                expected[sheet] = prefix + [
                    row for row in data if str(row[column]).strip().casefold() == expected_value
                ]
            elif edit.action == DocumentEditAction.SORT_RANGE:
                rows = expected[sheet]
                header = bool(edit.options.get("header", True))
                column = self._verification_column(rows, edit.options.get("column"), None)
                prefix = rows[:1] if header else []
                data = rows[1:] if header else rows
                data.sort(
                    key=lambda row: ExcelModifier._sort_key(row[column]),
                    reverse=bool(edit.options.get("descending", False)),
                )
                expected[sheet] = prefix + data
            elif edit.action in {
                DocumentEditAction.INSERT_ROWS,
                DocumentEditAction.DELETE_ROWS,
                DocumentEditAction.INSERT_COLUMNS,
                DocumentEditAction.DELETE_COLUMNS,
            }:
                # Structural edits are validated by the format-specific service and reopening.
                # Preservation of every other sheet is still checked below.
                continue
        actual = {table.name: [list(row) for row in table.rows] for table in after.tables}
        for sheet_name, rows in expected.items():
            structural = any(
                edit.locator.sheet == sheet_name
                and edit.action
                in {
                    DocumentEditAction.INSERT_ROWS,
                    DocumentEditAction.DELETE_ROWS,
                    DocumentEditAction.INSERT_COLUMNS,
                    DocumentEditAction.DELETE_COLUMNS,
                }
                for edit in edits
            )
            if not structural and actual.get(sheet_name) != rows:
                self._fail()
        if set(actual) != set(expected):
            self._fail()

    def _verify_requested_formatting(
        self,
        output: bytes,
        document_type: DocumentType,
        edits: tuple[DocumentEdit, ...],
    ) -> None:
        formatting = {
            DocumentEditAction.FORMAT_PARAGRAPH,
            DocumentEditAction.FORMAT_TABLE,
            DocumentEditAction.FORMAT_RANGE,
            DocumentEditAction.FREEZE_PANES,
        }
        if not any(edit.action in formatting for edit in edits):
            return
        if document_type == DocumentType.DOCX:
            document = Document(BytesIO(output))
            for edit in edits:
                if edit.action == DocumentEditAction.FORMAT_TABLE:
                    index = (edit.locator.table_index or 1) - 1
                    expected = str(edit.options.get("alignment", "center")).casefold()
                    alignments = {
                        "left": WD_TABLE_ALIGNMENT.LEFT,
                        "center": WD_TABLE_ALIGNMENT.CENTER,
                        "centre": WD_TABLE_ALIGNMENT.CENTER,
                        "right": WD_TABLE_ALIGNMENT.RIGHT,
                    }
                    if index >= len(document.tables) or document.tables[
                        index
                    ].alignment != alignments.get(expected):
                        self._fail()
            return
        if document_type == DocumentType.XLSX:
            workbook = load_workbook(BytesIO(output), data_only=False)
            try:
                for edit in edits:
                    sheet_name = edit.locator.sheet
                    if sheet_name is None:
                        if len(workbook.sheetnames) != 1:
                            self._fail()
                        sheet_name = workbook.sheetnames[0]
                    sheet = workbook[sheet_name]
                    if edit.action == DocumentEditAction.FREEZE_PANES and str(
                        sheet.freeze_panes
                    ) != str(edit.locator.cell):
                        self._fail()
                    if edit.action == DocumentEditAction.FORMAT_RANGE:
                        min_col, min_row, max_col, max_row = range_boundaries(
                            str(edit.locator.cell_range)
                        )
                        for row in sheet.iter_rows(
                            min_row=min_row,
                            max_row=max_row,
                            min_col=min_col,
                            max_col=max_col,
                        ):
                            for cell in row:
                                if (
                                    "alignment" in edit.options
                                    and cell.alignment.horizontal
                                    != str(edit.options["alignment"]).casefold()
                                ):
                                    self._fail()
                                if "bold" in edit.options and cell.font.bold != bool(
                                    edit.options["bold"]
                                ):
                                    self._fail()
                                if "wrap_text" in edit.options and cell.alignment.wrap_text != bool(
                                    edit.options["wrap_text"]
                                ):
                                    self._fail()
            finally:
                workbook.close()
            return
        if document_type == DocumentType.PPTX:
            # Reopening above already verified the presentation. Text/paragraph style edits
            # are applied only after an exactly-one locator check in PptxModifier.
            Presentation(BytesIO(output))

    @staticmethod
    def _verification_column(
        rows: list[list[str]], requested: object, expected_value: str | None
    ) -> int:
        if not rows:
            DocumentModificationService._fail()
        if isinstance(requested, int) and not isinstance(requested, bool):
            return requested - 1
        if isinstance(requested, str) and requested.strip():
            value = requested.strip()
            if value.isalpha() and len(value) <= 3:
                return column_index_from_string(value.upper()) - 1
            matches = [
                index
                for index, header in enumerate(rows[0])
                if str(header).strip().casefold() == value.casefold()
            ]
        elif expected_value is not None:
            matches = [
                index
                for index in range(len(rows[0]))
                if any(str(row[index]).strip().casefold() == expected_value for row in rows[1:])
            ]
        else:
            matches = []
        if len(matches) != 1:
            DocumentModificationService._fail()
        return matches[0]

    @staticmethod
    def _verify_pdf(
        before: DocumentRepresentation,
        after: DocumentRepresentation,
        edits: tuple[DocumentEdit, ...],
    ) -> None:
        expected_pages = before.metadata.page_count or 0
        for edit in edits:
            pages = edit.locator.pages or []
            if edit.action == DocumentEditAction.REMOVE_PDF_PAGES:
                expected_pages -= len(pages)
            elif edit.action == DocumentEditAction.EXTRACT_PDF_PAGES:
                expected_pages = len(pages)
            elif edit.action == DocumentEditAction.REORDER_PDF_PAGES:
                expected_pages = len(pages)
        if after.metadata.page_count != expected_pages:
            DocumentModificationService._fail()

    @staticmethod
    def _content_signature(document: DocumentRepresentation) -> tuple[object, ...]:
        return (
            document.text,
            tuple(table.rows for table in document.tables),
            document.metadata.page_count,
            document.metadata.sheet_names,
            document.metadata.slide_count,
        )

    @staticmethod
    def _fail() -> NoReturn:
        raise DocumentChangeVerificationError(
            "The modified document did not pass preservation checks, so no output was returned."
        )

    @staticmethod
    def _document_type(filename: str) -> DocumentType:
        extension = Path(filename).suffix.casefold().lstrip(".")
        try:
            return DocumentType(extension)
        except ValueError as error:
            raise UnsupportedDocumentModificationError(
                "Targeted modification supports DOCX, XLSX, PPTX, and PDF files."
            ) from error

    @staticmethod
    def _output_filename(filename: str, document_type: DocumentType) -> str:
        stem = Path(filename).stem or "document"
        return DocumentGenerationService.safe_filename(
            f"{stem}_updated.{document_type.value}", document_type
        )
