from io import BytesIO

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

from app.services.document.content_parser import content_lines
from app.services.document.extraction_models import (
    GeneratedTableContent,
    StructuredDocumentContent,
)


class DocxGenerator:
    """Generate a business-readable DOCX using the standard business brief preset."""

    def generate(self, content: str, *, plain_text: bool = False) -> bytes:
        document = Document()
        self._configure(document)

        for line in content_lines(content, plain_text=plain_text):
            if line.kind == "heading":
                document.add_heading(line.text, level=min(3, line.level))
            elif line.kind == "bullet":
                paragraph = document.add_paragraph(line.text, style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(6)
            elif line.kind == "number":
                paragraph = document.add_paragraph(line.text, style="List Number")
                paragraph.paragraph_format.space_after = Pt(6)
            elif line.kind == "blank":
                document.add_paragraph()
            else:
                document.add_paragraph(line.text)

        return self._save(document)

    def generate_structured(self, content: StructuredDocumentContent) -> bytes:
        document = Document()
        self._configure(document)
        document.add_paragraph(content.title, style="Title")
        self._append_blocks(document, content.paragraphs, content.bullet_lists, content.tables)
        for section in content.sections:
            document.add_heading(section.heading, level=1)
            self._append_blocks(document, section.paragraphs, section.bullet_lists, section.tables)
        return self._save(document)

    @staticmethod
    def _configure(document) -> None:
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        normal = document.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal.paragraph_format.line_spacing = 1.1
        heading_tokens = {
            1: (16, "2E74B5", 16, 8),
            2: (13, "2E74B5", 12, 6),
            3: (12, "1F4D78", 8, 4),
        }
        for level, (size, color, before, after) in heading_tokens.items():
            style = document.styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)

    @staticmethod
    def _append_blocks(
        document,
        paragraphs: list[str],
        bullet_lists: list[list[str]],
        tables: list[GeneratedTableContent],
    ) -> None:
        for text in paragraphs:
            document.add_paragraph(text)
        for items in bullet_lists:
            for item in items:
                document.add_paragraph(item, style="List Bullet")
        for content_table in tables:
            document.add_heading(content_table.title, level=2)
            width = len(content_table.headers) or len(content_table.rows[0])
            table = document.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            if content_table.headers:
                cells = table.add_row().cells
                for index, value in enumerate(content_table.headers):
                    cells[index].text = value
                    for run in cells[index].paragraphs[0].runs:
                        run.bold = True
            for values in content_table.rows:
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    cells[index].text = value

    @staticmethod
    def _save(document) -> bytes:
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()
