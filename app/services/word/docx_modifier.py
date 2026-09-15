from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.services.document.exceptions import (
    DocumentLocatorError,
    UnsupportedDocumentModificationError,
)
from app.services.document.extraction_models import (
    DocumentEdit,
    DocumentEditAction,
    DocumentLocator,
)


class DocxModifier:
    """Apply narrowly-scoped DOCX edits without asking an LLM to rewrite file bytes."""

    def modify(self, file_data: bytes, edits: tuple[DocumentEdit, ...]) -> bytes:
        document = Document(BytesIO(file_data))
        for edit in edits:
            self._apply(document, edit)
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def _apply(self, document, edit: DocumentEdit) -> None:
        action = edit.action
        if action == DocumentEditAction.REPLACE_TEXT:
            paragraph, old_text = self._resolve_text_paragraph(document, edit.locator)
            replacement = self._required_text(edit.value, "replacement text")
            self._replace_across_runs(paragraph, old_text, replacement)
        elif action == DocumentEditAction.ADD_PARAGRAPH:
            paragraph = self._resolve_paragraph(document, edit.locator)
            self._insert_paragraph_after(
                paragraph, self._required_text(edit.value, "paragraph text")
            )
        elif action == DocumentEditAction.REMOVE_PARAGRAPH:
            paragraph = self._resolve_paragraph(document, edit.locator)
            paragraph._element.getparent().remove(paragraph._element)
        elif action == DocumentEditAction.UPDATE_TABLE_CELL:
            cell = self._resolve_table_cell(document, edit.locator)
            cell.text = "" if edit.value is None else str(edit.value)
        elif action == DocumentEditAction.FORMAT_PARAGRAPH:
            paragraph = self._resolve_paragraph(document, edit.locator)
            self._format_paragraph(paragraph, edit.options)
        elif action == DocumentEditAction.FORMAT_TABLE:
            table = self._resolve_table(document, edit.locator)
            alignment = str(edit.options.get("alignment", "center")).casefold()
            values = {
                "left": WD_TABLE_ALIGNMENT.LEFT,
                "center": WD_TABLE_ALIGNMENT.CENTER,
                "centre": WD_TABLE_ALIGNMENT.CENTER,
                "right": WD_TABLE_ALIGNMENT.RIGHT,
            }
            if alignment not in values:
                raise UnsupportedDocumentModificationError(
                    "Table alignment must be left, center, or right."
                )
            table.alignment = values[alignment]
            if "vertical_alignment" in edit.options:
                vertical = str(edit.options["vertical_alignment"]).casefold()
                vertical_values = {
                    "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                    "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                    "centre": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                    "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                }
                if vertical not in vertical_values:
                    raise UnsupportedDocumentModificationError(
                        "Cell alignment must be top, center, or bottom."
                    )
                for row in table.rows:
                    for cell in row.cells:
                        cell.vertical_alignment = vertical_values[vertical]
        else:
            raise UnsupportedDocumentModificationError(
                f"The DOCX edit '{action.value}' is not supported."
            )

    @staticmethod
    def _all_paragraphs(document) -> list[Paragraph]:
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        return paragraphs

    def _resolve_text_paragraph(self, document, locator: DocumentLocator) -> tuple[Paragraph, str]:
        text = (locator.text or locator.heading or "").strip()
        if not text:
            raise DocumentLocatorError("Text replacement requires a text or heading locator.")
        paragraphs = self._all_paragraphs(document)
        matches: list[tuple[Paragraph, int]] = []
        for paragraph in paragraphs:
            count = paragraph.text.count(text)
            matches.extend((paragraph, occurrence) for occurrence in range(count))
        if locator.heading:
            matches = [
                match
                for match in matches
                if str(getattr(match[0].style, "name", "") or "").casefold().startswith("heading")
                or str(getattr(match[0].style, "name", "") or "").casefold() == "title"
            ]
        self._require_one(
            matches,
            f"I couldn't find text matching '{text}' in this document.",
            f"More than one location matches '{text}'. Please identify the paragraph or heading.",
        )
        return matches[0][0], text

    def _resolve_paragraph(self, document, locator: DocumentLocator) -> Paragraph:
        if locator.paragraph_index is not None:
            paragraphs = self._all_paragraphs(document)
            index = locator.paragraph_index - 1
            if index >= len(paragraphs):
                raise DocumentLocatorError(
                    f"Paragraph {locator.paragraph_index} does not exist in this document."
                )
            return paragraphs[index]
        paragraph, _ = self._resolve_text_paragraph(document, locator)
        return paragraph

    @staticmethod
    def _resolve_table(document, locator: DocumentLocator):
        if locator.table_index is None:
            if not document.tables:
                raise DocumentLocatorError("I couldn't find a table in this document.")
            if len(document.tables) > 1:
                raise DocumentLocatorError(
                    "This document has more than one table. Please identify the table number."
                )
            return document.tables[0]
        index = locator.table_index - 1
        if index >= len(document.tables):
            raise DocumentLocatorError(
                f"Table {locator.table_index} does not exist in this document."
            )
        return document.tables[index]

    def _resolve_table_cell(self, document, locator: DocumentLocator):
        table = self._resolve_table(document, locator)
        if locator.row is None or locator.column is None:
            raise DocumentLocatorError("A table-cell edit requires row and column numbers.")
        row_index = locator.row - 1
        column_index = locator.column - 1
        if row_index >= len(table.rows) or column_index >= len(table.columns):
            raise DocumentLocatorError(
                f"Table {locator.table_index} does not contain cell "
                f"({locator.row}, {locator.column})."
            )
        return table.cell(row_index, column_index)

    @staticmethod
    def _replace_across_runs(paragraph: Paragraph, old: str, new: str) -> None:
        runs: list[Run] = list(paragraph.runs)
        combined = "".join(run.text for run in runs)
        start = combined.find(old)
        if start < 0:
            raise DocumentLocatorError(f"I couldn't find text matching '{old}'.")
        end = start + len(old)
        offset = 0
        for run in runs:
            run_start = offset
            run_end = offset + len(run.text)
            offset = run_end
            if run_end <= start or run_start >= end:
                continue
            local_start = max(start - run_start, 0)
            local_end = min(end - run_start, len(run.text))
            prefix = run.text[:local_start]
            suffix = run.text[local_end:]
            if run_start <= start < run_end:
                run.text = prefix + new + suffix
            else:
                run.text = suffix

    @staticmethod
    def _insert_paragraph_after(paragraph: Paragraph, text: str) -> None:
        element = OxmlElement("w:p")
        paragraph._p.addnext(element)
        inserted = Paragraph(element, paragraph._parent)
        inserted.style = paragraph.style
        inserted.add_run(text)

    @staticmethod
    def _format_paragraph(paragraph: Paragraph, options: dict[str, object]) -> None:
        if "alignment" in options:
            alignment = str(options["alignment"]).casefold()
            values = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "centre": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            if alignment not in values:
                raise UnsupportedDocumentModificationError(
                    "Paragraph alignment must be left, center, right, or justify."
                )
            paragraph.alignment = values[alignment]
        if "space_before_pt" in options:
            paragraph.paragraph_format.space_before = Pt(float(str(options["space_before_pt"])))
        if "space_after_pt" in options:
            paragraph.paragraph_format.space_after = Pt(float(str(options["space_after_pt"])))
        if "line_spacing" in options:
            paragraph.paragraph_format.line_spacing = float(str(options["line_spacing"]))
        for run in paragraph.runs:
            if "bold" in options:
                run.bold = bool(options["bold"])
            if "italic" in options:
                run.italic = bool(options["italic"])
            if "font_size_pt" in options:
                run.font.size = Pt(float(str(options["font_size_pt"])))

    @staticmethod
    def _required_text(value: object | None, label: str) -> str:
        if not isinstance(value, str):
            raise UnsupportedDocumentModificationError(f"The {label} must be text.")
        return value

    @staticmethod
    def _require_one(matches, missing: str, ambiguous: str) -> None:
        if not matches:
            raise DocumentLocatorError(missing)
        if len(matches) > 1:
            raise DocumentLocatorError(ambiguous)
