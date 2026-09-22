import asyncio
import json
from contextlib import closing
from io import BytesIO

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

from app.api import chat_documents
from app.main import app
from app.services.chat_document_service import ChatDocumentService
from app.services.document.document_automation_service import (
    AvailableDocument,
    DocumentAutomationService,
)
from app.services.document.document_intent_service import InvalidDocumentParametersError
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.document_pipeline_service import (
    DocumentPipelineService,
    StructuredPipelineStep,
)
from app.services.document.extraction_models import (
    GeneratedTableContent,
    StructuredDocumentContent,
)


def _login(client, email: str) -> dict[str, str]:
    client.post(
        "/users/",
        json={"fullname": "Phase Five", "email": email, "password": "secret123"},
    )
    token = client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


def _pdf_bytes(label: str) -> bytes:
    output = BytesIO()
    canvas = Canvas(output)
    canvas.drawString(72, 760, label)
    canvas.save()
    return output.getvalue()


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("Operations Report", level=1)
    document.add_paragraph("Customer: Atlas Foods")
    document.save(output)
    return output.getvalue()


def _upload(client, headers, filename: str, data: bytes, session_id: int | None = None):
    form = {"message": f"Upload {filename}", "analyze": "false"}
    if session_id is not None:
        form["session_id"] = str(session_id)
    response = client.post(
        "/chat/documents",
        headers=headers,
        data=form,
        files={"file": (filename, data, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_create_plan_normalizes_target_type_out_of_source_type_field():
    service = DocumentPipelineService()
    steps = service.plan_structured(
        (
            StructuredPipelineStep(
                document_type=DocumentType.DOCX,
                operation=DocumentOperation.CREATE_DOCUMENT,
                input_file="sample-1.pdf",
                output_type=DocumentType.DOCX,
                parameters={},
            ),
        ),
        input_filename="sample-1.pdf",
        request_instruction="can you create the word document on this content",
    )

    assert len(steps) == 1
    assert steps[0].intent.document_type == DocumentType.PDF
    assert steps[0].intent.output_type == DocumentType.DOCX


def test_simple_word_creation_uses_deterministic_validated_plan(monkeypatch):
    service = DocumentAutomationService()

    def unexpected_provider(*args, **kwargs):
        raise AssertionError("An unambiguous creation request must not need the AI planner")

    monkeypatch.setattr(service.chat_service, "chat", unexpected_provider)
    plan = service.plan(
        "can you create the word document on this content",
        (
            AvailableDocument(
                document_id=7,
                filename="sample-1.pdf",
                document_type=DocumentType.PDF,
                is_latest=True,
            ),
        ),
    )

    assert plan.status == "ready"
    assert len(plan.steps) == 1
    assert plan.steps[0].intent.operation == DocumentOperation.CONVERT_DOCUMENT
    assert plan.steps[0].intent.document_type == DocumentType.PDF
    assert plan.steps[0].intent.output_type == DocumentType.DOCX
    assert plan.steps[0].source_filenames == ("sample-1.pdf",)


def test_simple_document_question_uses_deterministic_validated_plan(monkeypatch):
    service = DocumentAutomationService()

    def unexpected_provider(*args, **kwargs):
        raise AssertionError("An unambiguous document question must not need the AI planner")

    monkeypatch.setattr(service.chat_service, "chat", unexpected_provider)
    plan = service.plan(
        "What date appears in this document?",
        (
            AvailableDocument(
                document_id=7,
                filename="sample-1.pdf",
                document_type=DocumentType.PDF,
                is_latest=True,
            ),
        ),
    )

    assert plan.status == "ready"
    assert len(plan.steps) == 1
    assert plan.steps[0].intent.operation == DocumentOperation.ANSWER_DOCUMENT_QUESTION
    assert plan.steps[0].intent.document_type == DocumentType.PDF
    assert plan.steps[0].intent.output_type is None
    assert plan.steps[0].intent.parameters == {"question": "What date appears in this document?"}


@pytest.mark.parametrize(
    "instruction",
    [
        "create a Word document from this",
        "make this into a Word document",
        "create a docx from this content",
        "itha Word document ah create panni kudu",
        "indha content ah Word file ah maathi kudu",
        "itha docx ah create pannu",
        "indha document ah Word file ah ready pannu",
    ],
)
def test_word_creation_variants_use_the_same_validated_operation(instruction, monkeypatch):
    service = DocumentAutomationService()
    monkeypatch.setattr(
        service.chat_service,
        "chat",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("The deterministic intent engine should resolve this request")
        ),
    )
    plan = service.plan(
        instruction,
        (
            AvailableDocument(
                document_id=7,
                filename="sample-1.pdf",
                document_type=DocumentType.PDF,
                is_latest=True,
            ),
        ),
    )

    assert plan.steps[0].intent.operation == DocumentOperation.CONVERT_DOCUMENT
    assert plan.steps[0].intent.output_type == DocumentType.DOCX
    assert plan.steps[0].source_filenames == ("sample-1.pdf",)


def test_explicit_conversion_variant_uses_conversion_pipeline_without_ai(monkeypatch):
    service = DocumentAutomationService()
    monkeypatch.setattr(
        service.chat_service,
        "chat",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("An unambiguous conversion must not need the AI planner")
        ),
    )
    plan = service.plan(
        "convert this to Word",
        (
            AvailableDocument(
                document_id=7,
                filename="sample-1.pdf",
                document_type=DocumentType.PDF,
                is_latest=True,
            ),
        ),
    )

    assert plan.steps[0].intent.operation == DocumentOperation.CONVERT_DOCUMENT
    assert plan.steps[0].intent.document_type == DocumentType.PDF
    assert plan.steps[0].intent.output_type == DocumentType.DOCX


def _generated_content() -> StructuredDocumentContent:
    return StructuredDocumentContent(
        title="Customer Report",
        paragraphs=["Prepared from the uploaded source."],
        tables=[
            GeneratedTableContent(
                title="Customers",
                headers=["Customer", "Status"],
                rows=[["Atlas Foods", "Active"]],
            )
        ],
    )


def test_read_extract_create_excel_and_word_report_end_to_end(client, monkeypatch):
    headers = _login(client, "phase5-create@example.com")
    uploaded = _upload(client, headers, "customers.pdf", _pdf_bytes("Atlas Foods Active"))

    async def generate_content(self, *args, **kwargs):
        return _generated_content()

    monkeypatch.setattr(ChatDocumentService, "generate_content", generate_content)
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "pdf",
                        "operation": "read_document",
                        "input_file": "customers.pdf",
                        "output_type": "text_response",
                        "parameters": {},
                    },
                    {
                        "document_type": "pdf",
                        "operation": "extract_document",
                        "input_file": "customers.pdf",
                        "output_type": "text_response",
                        "parameters": {},
                    },
                    {
                        "document_type": None,
                        "operation": "create_document",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {
                            "filename": "extracted-data.xlsx",
                            "title": "Extracted Data",
                        },
                    },
                    {
                        "document_type": None,
                        "operation": "create_document",
                        "input_file": "customers.pdf",
                        "output_type": "docx",
                        "parameters": {
                            "filename": "document-report.docx",
                            "title": "Document Report",
                        },
                    },
                ],
            }
        ),
    )
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "confirm": True,
            "session_id": uploaded["session_id"],
            "instruction": (
                "Read this PDF, extract the customer data, create an Excel and prepare a Word "
                "report."
            ),
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "done"
    assert [step["status"] for step in body["steps"]] == ["completed"] * 4
    assert [item["filename"] for item in body["attachments"]] == [
        "extracted-data.xlsx",
        "document-report.docx",
    ]
    assert "All files are attached" in body["response"]

    excel = client.get(f"/chat/documents/{body['attachments'][0]['id']}/download", headers=headers)
    with closing(load_workbook(BytesIO(excel.content), read_only=True)) as workbook:
        assert workbook.sheetnames
        assert "Atlas Foods" in {
            str(cell.value)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
    word = client.get(f"/chat/documents/{body['attachments'][1]['id']}/download", headers=headers)
    assert "Customer Report" in "\n".join(
        paragraph.text for paragraph in Document(BytesIO(word.content)).paragraphs
    )


def test_update_word_and_return_final_pdf_end_to_end(client, monkeypatch):
    headers = _login(client, "phase5-word@example.com")
    uploaded = _upload(client, headers, "operations.docx", _docx_bytes())

    async def generate_content(self, *args, **kwargs):
        return _generated_content()

    monkeypatch.setattr(ChatDocumentService, "generate_content", generate_content)
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "docx",
                        "operation": "read_document",
                        "input_file": "operations.docx",
                        "output_type": "text_response",
                        "parameters": {},
                    },
                    {
                        "document_type": None,
                        "operation": "create_document",
                        "input_file": "operations.docx",
                        "output_type": "docx",
                        "parameters": {
                            "filename": "updated-document.docx",
                            "title": "Updated Document",
                        },
                    },
                    {
                        "document_type": "docx",
                        "operation": "convert_document",
                        "input_file": None,
                        "output_type": "pdf",
                        "parameters": {},
                    },
                ],
            }
        ),
    )
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "confirm": True,
            "session_id": uploaded["session_id"],
            "instruction": (
                "Update this Word document, improve the formatting and give me the final PDF."
            ),
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "done"
    assert [item["filename"] for item in body["attachments"]] == [
        "updated-document.docx",
        "updated-document-converted.pdf",
    ]
    downloaded = client.get(
        f"/chat/documents/{body['latest_document_id']}/download", headers=headers
    )
    assert len(PdfReader(BytesIO(downloaded.content)).pages) >= 1


def test_merge_uses_multiple_named_inputs_and_preserves_both_sources(client, monkeypatch):
    headers = _login(client, "phase5-merge@example.com")
    first = _upload(client, headers, "north.pdf", _pdf_bytes("North"))
    second = _upload(
        client,
        headers,
        "south.pdf",
        _pdf_bytes("South"),
        session_id=first["session_id"],
    )
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "pdf",
                        "operation": "merge_pdf",
                        "input_file": ["north.pdf", "south.pdf"],
                        "output_type": "pdf",
                        "parameters": {},
                    }
                ],
            }
        ),
    )

    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "confirm": True,
            "session_id": first["session_id"],
            "instruction": "Combine both PDF files.",
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["steps"][0]["source_document_ids"] == [
        first["attachment"]["id"],
        second["attachment"]["id"],
    ]
    merged = client.get(f"/chat/documents/{body['latest_document_id']}/download", headers=headers)
    assert len(PdfReader(BytesIO(merged.content)).pages) == 2


def test_clarification_reply_reuses_persisted_conversation(client, monkeypatch):
    headers = _login(client, "phase5-clarification@example.com")
    first = _upload(client, headers, "first.pdf", _pdf_bytes("First customer"))
    _upload(
        client,
        headers,
        "second.pdf",
        _pdf_bytes("Second customer"),
        session_id=first["session_id"],
    )
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "clarification_required",
                "clarifying_question": "Which uploaded PDF should I use for this request?",
                "steps": [],
            }
        ),
    )
    ambiguous = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "confirm": True,
            "session_id": first["session_id"],
            "instruction": "Extract the customer data and create Excel and Word reports.",
        },
    )
    assert ambiguous.json()["status"] == "clarification_required"

    prompts = []

    def plan(message, history, reference_history, **kwargs):
        prompts.append(message)
        return json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "pdf",
                        "operation": "read_document",
                        "input_file": "first.pdf",
                        "output_type": "text_response",
                        "parameters": {},
                    }
                ],
            }
        )

    monkeypatch.setattr(chat_documents.automation_service.chat_service, "chat", plan)
    resolved = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "confirm": True,
            "session_id": first["session_id"],
            "instruction": "Use the first PDF.",
        },
    )

    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["status"] == "done"
    assert "First customer" in resolved.json()["response"]
    assert prompts and "Which uploaded PDF should I use" in prompts[0]


def test_concurrent_same_session_pipelines_keep_a_coherent_latest_pointer(client, db_session):
    headers = _login(client, "phase5-concurrency@example.com")
    uploaded = _upload(client, headers, "items.csv", b"Item,Count\nRice,2\n")
    payload = {
        "session_id": uploaded["session_id"],
        "source_document_id": uploaded["attachment"]["id"],
        "instruction": "Convert items.csv to Excel",
    }

    async def run_both():
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as async_client:
            return await asyncio.gather(
                async_client.post("/chat/documents/pipeline", headers=headers, json=payload),
                async_client.post("/chat/documents/pipeline", headers=headers, json=payload),
            )

    responses = asyncio.run(run_both())
    assert [response.status_code for response in responses] == [200, 200]
    ids = [response.json()["latest_document_id"] for response in responses]
    db_session.expire_all()
    session = chat_documents.repository.get_session(db_session, uploaded["session_id"], 1)
    assert session is not None
    assert int(session.latest_document_id) in ids
    assert len(set(ids)) == 2
    for document_id in ids:
        download = client.get(f"/chat/documents/{document_id}/download", headers=headers)
        with closing(load_workbook(BytesIO(download.content), read_only=True)) as workbook:
            assert workbook.sheetnames


def test_pipeline_rejects_an_overlong_plan_before_execution():
    service = DocumentPipelineService(max_steps=2)
    steps = tuple(
        StructuredPipelineStep(
            document_type=DocumentType.PDF,
            operation=DocumentOperation.READ_DOCUMENT,
            input_file="source.pdf",
            output_type=None,
            parameters={},
        )
        for _ in range(3)
    )

    with pytest.raises(InvalidDocumentParametersError, match="at most 2 steps"):
        service.plan_structured(steps, input_filename="source.pdf")


def test_pipeline_rejects_multiple_inputs_for_a_single_input_operation():
    service = DocumentPipelineService()
    step = StructuredPipelineStep(
        document_type=DocumentType.PDF,
        operation=DocumentOperation.READ_DOCUMENT,
        input_file=["first.pdf", "second.pdf"],
        output_type=None,
        parameters={},
    )

    with pytest.raises(InvalidDocumentParametersError, match="exactly one source"):
        service.plan_structured((step,), input_filename="first.pdf")
