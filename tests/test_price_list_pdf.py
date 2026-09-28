from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook
from pypdf import PdfReader

from app.api import chat_documents
from app.services.chat_document_service import ChatDocumentService
from app.services.document.document_automation_service import (
    AvailableDocument,
    DocumentAutomationService,
)
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.exceptions import InvalidDocumentError
from app.services.pdf.price_list_pdf import PriceListPdf, PriceListPdfRequest, clean_price
from app.utils.document_output import requests_pdf_output
from tests.test_document_automation import _login, _upload

REQUEST = """Covert to pdf
TASK: Generate a PDF file. The output MUST be a .pdf file.
Do NOT create or modify any Excel file. Do NOT split the workbook into sheets.
Read all 127 rows. Group by price ONLY. Ignore Category completely for grouping.
Show a check table: Price | Item Count | Total.
Page 1, cover: "Catering Menu", "Price-wise Item List".
Repeat the header row and format the columns.
Give me ONE downloadable PDF file named Catering_Menu_Price_Wise.pdf."""


def workbook_bytes(rows, extra_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "menus"
    sheet.append(["#", "Code", "Name", "Category", "Status", "Price"])
    for index, row in enumerate(rows, 1):
        sheet.append([index, *row])
    if extra_sheet:
        book.copy_worksheet(sheet)
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


@pytest.mark.parametrize(
    "value,expected",
    [
        (" ₹ 1,234.565 ", Decimal("1234.57")),
        ("Rs. 3.0", Decimal("3.00")),
        ("Rs 3", Decimal("3.00")),
        (3, Decimal("3.00")),
        (0, Decimal("0.00")),
        (None, None),
        (" ", None),
    ],
)
def test_clean_numeric_prices(value, expected):
    assert clean_price(value) == expected


@pytest.mark.parametrize("value", ["ask us", "NaN", "Infinity", True])
def test_invalid_price_is_not_silently_dropped_or_zero(value):
    with pytest.raises(InvalidDocumentError, match="Invalid Price"):
        clean_price(value)


def test_pdf_grouping_uses_actual_rows_and_reports_count_mismatch():
    rows = [
        ("A1", "Tomato Shorba", "Soup", "Active", "Rs 3.00"),
        ("A2", "Asst. Naan", "Bread", "Inactive", 3),
        ("A3", "Sauce", "Sides", "Active", None),
        ("A4", "Water", "Drink", "Active", 0),
    ]
    data = workbook_bytes(rows)
    request = PriceListPdfRequest.from_instruction(REQUEST)
    assert request and request.expected_count == 127
    result = PriceListPdf().read(data, request)
    assert result.count == 4
    assert list(result.groups) == [Decimal("0.00"), Decimal("3.00"), None]
    assert len(result.groups[Decimal("3.00")]) == 2
    pdf, summary = PriceListPdf().generate(data, request, today=date(2026, 9, 28))
    reader = PdfReader(BytesIO(pdf))
    assert len(reader.pages) == 3
    assert "28 September 2026" in reader.pages[0].extract_text()
    assert "Total Items: 4" in reader.pages[0].extract_text()
    assert "Grand total" in reader.pages[1].extract_text()
    details = reader.pages[2].extract_text()
    for row in rows:
        assert details.count(row[0]) == 1
        assert row[1] in details
    assert "₹ 3.00" in details
    assert "4 items, not the requested 127" in summary
    assert "| **Total** | **4** |" in summary
    assert "Price not provided" in summary
    for index, page in enumerate(reader.pages[1:], 2):
        assert f"Page {index} of 3" in page.extract_text()


def test_table_continuations_repeat_price_band_and_column_headers():
    rows = [
        (f"ITEM-{index:04}", "Long item name with details " * 2, "Soup", "Active", 3)
        for index in range(95)
    ]
    pdf, _ = PriceListPdf().generate(workbook_bytes(rows), PriceListPdfRequest())
    reader = PdfReader(BytesIO(pdf))
    text = "\n".join(page.extract_text() for page in reader.pages[2:])
    for row in rows:
        assert text.count(row[0]) == 1
    for page in reader.pages[2:]:
        content = page.extract_text()
        assert "₹ 3.00" in content and "Code" in content and "Status" in content
        assert "ITEM-" in content  # no orphan section/header-only pages


def test_empty_ambiguous_formula_and_excessive_groups_fail_clearly():
    request = PriceListPdfRequest()
    cases = [
        (workbook_bytes([]), "no item rows"),
        (workbook_bytes([("A", "Soup", "Food", "Active", "=1+1")]), "formula has no saved value"),
        (workbook_bytes([("", "", "Food", "Active", 2)]), "no Code or Name"),
        (
            workbook_bytes([("A", "Soup", "Food", "Active", 2)], extra_sheet=True),
            "Name one worksheet",
        ),
        (
            workbook_bytes([(str(i), "Soup", "Food", "Active", i) for i in range(61)]),
            "61 price groups",
        ),
    ]
    for data, message in cases:
        with pytest.raises(InvalidDocumentError, match=message):
            PriceListPdf().read(data, request)
    assert PriceListPdf().read(cases[3][0], PriceListPdfRequest(sheet_name="menus")).count == 1


@pytest.mark.parametrize(
    "instruction",
    [REQUEST, "Covert to pdf", "Convert my Excel to PDF", "same price items PDF ah kudu"],
)
def test_pdf_output_bypasses_legacy_excel_edits(instruction):
    assert requests_pdf_output(instruction)
    assert ChatDocumentService.spreadsheet_operation(instruction) is None


def test_pdf_as_input_does_not_override_excel_output():
    assert not requests_pdf_output("Convert this PDF to Excel")
    assert (
        PriceListPdfRequest.from_instruction("Split the workbook by Price into separate sheets")
        is None
    )
    assert (
        ChatDocumentService.spreadsheet_operation(
            "Do not split by Category. Add a column filter to Category"
        )
        is not None
    )


def test_short_conversion_and_missing_source_never_generate_a_guide(monkeypatch):
    service = DocumentAutomationService()
    monkeypatch.setattr(
        service.chat_service, "chat", lambda *a, **k: pytest.fail("No model needed")
    )
    source = AvailableDocument(1, "menu.xlsx", DocumentType.XLSX, is_latest=True)
    for instruction in ("Covert to pdf", "Convert to PDF"):
        plan = service.plan(instruction, (source,), explicit_source=True)
        assert plan.steps[0].intent.operation == DocumentOperation.CONVERT_DOCUMENT
        assert plan.steps[0].source_filenames == ("menu.xlsx",)
    assert service.plan(REQUEST, ()).status == "clarification_required"
    assert service.plan("Convert to PDF", ()).status == "clarification_required"


def test_symbol_number_header_does_not_match_arbitrary_by_instruction():
    book = Workbook()
    book.active.append(["#", "Code", "Name", "Category", "Price"])
    selected = ChatDocumentService._find_grouping_header(tuple(book.active[1]), "split by category")
    assert selected.column == 4
    book.close()


def test_pdf_automation_returns_only_downloadable_pdf_and_counts(client, monkeypatch):
    headers = _login(client, "price-pdf@example.com")
    data = workbook_bytes(
        [
            ("A", "Tomato Shorba", "Soup", "Active", 3),
            ("B", "Asst. Naan", "Bread", "Active", "₹ 3.00"),
        ]
    )
    uploaded = _upload(
        client,
        headers,
        "menu.xlsx",
        data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *a, **k: pytest.fail("No model needed for source-based PDF"),
    )
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={"session_id": uploaded["session_id"], "instruction": REQUEST},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "done", result
    assert "| **Total** | **2** |" in result["response"]
    assert "not the requested 127" in result["response"]
    assert len(result["steps"]) == 1
    assert result["steps"][0]["output_type"] == "pdf"
    assert result["steps"][0]["filename"] == "Catering_Menu_Price_Wise.pdf"
    download = client.get(
        f"/chat/documents/{result['latest_document_id']}/download", headers=headers
    )
    assert download.content.startswith(b"%PDF-")
    details = PdfReader(BytesIO(download.content)).pages[-1].extract_text()
    assert "Tomato Shorba" in details and "Asst. Naan" in details
    original = client.get(
        f"/chat/documents/{uploaded['attachment']['id']}/download", headers=headers
    )
    assert original.content == data
