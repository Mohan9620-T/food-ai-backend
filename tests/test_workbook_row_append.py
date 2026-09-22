import json
from copy import copy
from io import BytesIO
from unittest.mock import MagicMock
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from app.api import chat_documents
from app.services.document.document_automation_service import DocumentAutomationService
from app.services.document.exceptions import DocumentLocatorError
from app.services.spreadsheet.row_append import AppendRowsRequest, WorkbookRowAppender

LEAD = "can you add the new data's on the sheet and update and give me the updated sheet"
ROWS = [
    [
        "DRK0836",
        "MAGNOLIA FRESH UHT MILK BANANA_200ML × 24 EL",
        "Consumable",
        "Beverages",
        "CTN",
        "12 PACK",
        "2",
    ],
    [
        "DRY3886",
        "JINTAN MANIS (FENNEL SEED) EL",
        "Consumable",
        "Dry Goods / Grocery",
        "1 KG",
        "1 KG ( 12 PACK / CTN )",
        "2",
    ],
]
DATA = "\n".join("\t".join(row) for row in ROWS)


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "Items"
    sheet.append(
        [
            "ItemCode*",
            "ItemName*",
            "ItemType*",
            "Category*",
            "Uom*",
            "BaseUom",
            "UomFactor",
            "Formula",
        ]
    )
    sheet.append(["001", "First", "Consumable", "Beverages", "CTN", "CTN", 1, '=IF(A2<>"",A2,"")'])
    sheet.append(
        ["002", "Second", "Consumable", "Dry Goods / Grocery", "KG", "KG", 1, '=IF(A3<>"",A3,H2)']
    )
    for row in range(4, 8):
        sheet.cell(row, 8, f'=IF(A{row}<>"",A{row},H{row - 1})')
    sheet["A1"].font = Font(bold=True)
    sheet["B3"].font = Font(name="Arial", color="115599")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H3"
    validation = DataValidation(type="list", formula1='"CTN,KG"')
    validation.add("E2:E3")
    sheet.add_data_validation(validation)
    book.create_sheet("Reference").append(["Instructions", "Do not modify", "=1+2"])
    audit = book.create_sheet("Upload_Audit")
    audit.append(["Status meaning:", "OK or FAIL"])
    audit.append([])
    audit.append(
        ["ItemCode", "ItemName", "Category", "Uom", "BaseUom", "UomFactor", "Status", "Notes"]
    )
    audit.append(["001", "First", "Beverages", "CTN", "CTN", 1, "OK", None])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.mark.parametrize("data", [DATA, "\n".join(" ".join(row) for row in ROWS)])
def test_append_catalog_rows_preserves_existing_data_formulas_and_other_sheets(data):
    raw = workbook_bytes()
    result, resolved = WorkbookRowAppender().append(raw, AppendRowsRequest(data=data))
    assert resolved.start_row == 4 and resolved.sheet_name == "Items"
    before, after = load_workbook(BytesIO(raw)), load_workbook(BytesIO(result))
    assert before.sheetnames == after.sheetnames
    for sheet in before:
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    actual = after[sheet.title][cell.coordinate]
                    assert (actual.value, actual.data_type) == (cell.value, cell.data_type)
    for index, row in enumerate(ROWS, 4):
        assert [after["Items"].cell(index, col).value for col in range(1, 8)] == [*row[:6], 2]
    assert copy(after["Items"]["B4"].font) == copy(after["Items"]["B3"].font)
    assert after["Items"]["A1"].font.bold
    assert after["Items"].freeze_panes == "A2"
    assert after["Items"].auto_filter.ref == "A1:H5"
    assert "E4" in after["Items"].data_validations.dataValidation[0]
    before.close()
    after.close()


@pytest.mark.parametrize(
    "data",
    [
        'ItemName*,ItemCode*,UomFactor\n"New, item",0009,0',
        "| ItemName* | ItemCode* | UomFactor |\n| --- | --- | --- |\n| New, item | 0009 | 0 |",
        json.dumps([{"ItemName*": "New, item", "ItemCode*": "0009", "UomFactor": 0}]),
    ],
)
def test_named_columns_and_literal_text_are_preserved(data):
    output, resolved = WorkbookRowAppender().append(
        workbook_bytes(), AppendRowsRequest(data=data, sheet_name="Items")
    )
    book = load_workbook(BytesIO(output))
    assert book["Items"].cell(resolved.start_row, 1).value == "0009"
    assert book["Items"].cell(resolved.start_row, 2).value == "New, item"
    assert book["Items"].cell(resolved.start_row, 7).value == 0
    book.close()


def test_ambiguous_unseparated_values_request_columns_instead_of_guessing():
    with pytest.raises(DocumentLocatorError, match="column headers"):
        WorkbookRowAppender().resolve(
            workbook_bytes(), AppendRowsRequest(data="A new item unknown type 20")
        )
    assert AppendRowsRequest.from_instruction(LEAD) is None
    assert AppendRowsRequest.is_request("Append these rows using the column headers")
    assert not AppendRowsRequest.is_request("Add a new column with these data values")
    with pytest.raises(DocumentLocatorError, match="Items, Upload_Audit"):
        WorkbookRowAppender().resolve(
            workbook_bytes(), AppendRowsRequest(data='[{"ItemCode": "009", "ItemName": "New"}]')
        )


def test_plain_text_formulas_are_not_executed_and_table_range_extends():
    book = Workbook()
    book.active.append(["Code", "Description"])
    book.active.append(["001", "Before"])
    book.active.add_table(Table(displayName="Items", ref="A1:B2"))
    raw = BytesIO()
    book.save(raw)
    output, _ = WorkbookRowAppender().append(raw.getvalue(), AppendRowsRequest(data="002\t=1+2"))
    updated = load_workbook(BytesIO(output))
    assert updated.active["B3"].value == "=1+2" and updated.active["B3"].data_type == "s"
    assert updated.active.tables["Items"].ref == "A1:B3"
    updated.close()
    book.close()


def test_preserves_exact_existing_decimal_values_and_unrelated_package_parts():
    book = Workbook()
    book.active.title = "Items"
    book.active.append(["Code", "Cost"])
    book.active.append(["001", 11.5])
    book.create_sheet("Reference").append(["Leave this sheet intact"])
    stream = BytesIO()
    book.save(stream)
    original = BytesIO()
    with ZipFile(stream) as source, ZipFile(original, "w") as archive:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v>11.5</v>", b"<v>11.666666666666666</v>")
                # Some Excel writers store blank template cells as an empty value node.
                data = data.replace(
                    b"</sheetData>", b'<row r="3"><c r="A3"><v/></c></row></sheetData>'
                )
            archive.writestr(info, data)
    updated, _ = WorkbookRowAppender().append(
        original.getvalue(), AppendRowsRequest(data="002\t15", sheet_name="Items")
    )
    reopened = load_workbook(BytesIO(updated))
    assert reopened["Items"]["B2"].value == 11.666666666666666
    assert reopened["Items"]["B3"].value == 15
    with ZipFile(original) as source, ZipFile(BytesIO(updated)) as after:
        assert source.namelist() == after.namelist()
        for name in source.namelist():
            if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                assert source.read(name) == after.read(name)
    reopened.close()
    book.close()


def test_trailing_blank_cells_and_invalid_cell_values():
    request = AppendRowsRequest.from_instruction('Add rows to sheet "Items"\nCode\tCost\n002\t')
    assert request.data.endswith("\t")
    book = Workbook()
    book.active.title = "Items"
    book.active.append(["Code", "Cost"])
    book.active.append(["001", 11])
    raw = BytesIO()
    book.save(raw)
    result, _ = WorkbookRowAppender().append(raw.getvalue(), request)
    after = load_workbook(BytesIO(result))
    assert after.active["A3"].value == "002" and after.active["B3"].value is None
    after.close()
    for value in [float("nan"), float("inf"), "a" * 32768]:
        with pytest.raises(DocumentLocatorError, match="finite number"):
            WorkbookRowAppender().resolve(
                raw.getvalue(), AppendRowsRequest(data=json.dumps([{"Code": value}]))
            )
    book.close()


def test_upload_append_missing_data_clarification_direct_creation_and_download(client, monkeypatch):
    def no_ai(*args, **kwargs):
        raise AssertionError("Appending supplied rows must not call AI")

    monkeypatch.setattr(chat_documents.automation_service, "chat_service", MagicMock(chat=no_ai))
    monkeypatch.setattr(type(chat_documents.service), "generate_content", no_ai)
    monkeypatch.setattr(type(chat_documents.service), "summarize", no_ai)
    credentials = {"email": "append@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Append test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    raw = workbook_bytes()
    # Legacy clients no longer receive a 422 for missing rows; the original is saved.
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"message": LEAD},
        files={"file": ("items.xlsx", raw)},
    )
    assert upload.status_code == 200, upload.text
    body = upload.json()
    assert body["response"] == DocumentAutomationService.APPEND_NEEDS_DATA
    session, original = body["session_id"], body["attachment"]["id"]
    base = {"session_id": session, "source_document_id": original}
    question = client.post(
        "/chat/documents/automate", headers=headers, json={**base, "instruction": LEAD}
    ).json()
    assert question["status"] == "clarification_required"
    result = client.post(
        "/chat/documents/automate", headers=headers, json={**base, "instruction": DATA}
    ).json()
    assert result["status"] == "done", result
    assert result["steps"][0]["operation"] == "append_workbook_rows"
    assert "Added 2 new rows" in result["response"]
    generated = result["attachments"][0]
    assert generated["source_document_ids"] == [original]
    file = client.get(f"/chat/documents/{generated['id']}/download", headers=headers)
    book = load_workbook(BytesIO(file.content))
    assert book["Items"]["A4"].value == "DRK0836" and book["Items"]["A5"].value == "DRY3886"
    book.close()
    assert client.get(f"/chat/documents/{original}/download", headers=headers).content == raw
    assert client.get(f"/chat/documents/{generated['id']}/data", headers=headers).json()[
        "extracted_data"
    ]["tables"]
    # Direct upload with supplied data also remains supported.
    direct = client.post(
        "/chat/documents",
        headers=headers,
        data={"message": LEAD + "\n" + DATA},
        files={"file": ("items.xlsx", raw)},
    )
    assert direct.status_code == 200 and "Added 2 new rows" in direct.json()["response"]
