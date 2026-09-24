from collections import Counter
from contextlib import closing
from copy import copy
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.api import chat_documents
from app.services.document.exceptions import DocumentLocatorError, InvalidDocumentError
from app.services.spreadsheet.column_grouping import ColumnGroupingRequest, WorkbookColumnGrouper
from app.services.web_lookup import automatic_web_question

PROMPTS = [
    "Put items with the same price into separate sheets",
    "same rate items thani thani sheet la list out panni kudu",
    "Split my uploaded Excel into one worksheet for each price",
    "List same price items in different tabs",
    "ஒரே விலையில் உள்ள பொருட்களை தனித்தனி தாள்களில் பிரித்து கொடு",
    "Use only the Catering Menu worksheet. Split rows based on Price into separate worksheets. Do not split by Category. Keep the original worksheet.",
]


def workbook_bytes(prices=(2, "1.50", 1.5, 0, None, 10)):
    book = Workbook()
    sheet = book.active
    sheet.title = "menus"
    sheet.append(["#", "Code", "Name", "Category", "Status", "Price", None])
    for index, price in enumerate(prices):
        sheet.append(
            [
                index + 1,
                f"FM{index:03}",
                f"Item {index}",
                ["Rice", "Bread"][index % 2],
                "Active",
                price,
                None if index % 2 == 0 else "Keep notes",
            ]
        )
    sheet["C2"].font = Font(bold=True, color="FF0000")
    sheet["F2"].number_format = "0.00"
    sheet.column_dimensions["C"].width = 32
    overview = book.create_sheet("Overview")
    overview.append(["Existing note", "Do not change"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.mark.parametrize("prompt", PROMPTS)
def test_plain_grouping_requests_are_not_web_queries(prompt):
    request = ColumnGroupingRequest.from_instruction(prompt)
    assert request is not None and request.column == "Price"
    assert not automatic_web_question(prompt)


@pytest.mark.parametrize(
    "prompt",
    [
        "What is the current price of rice?",
        "Sort the items price-wise in the same sheet",
        "Do not split by Price into separate sheets. Just describe the file.",
    ],
)
def test_does_not_group_unrelated_requests(prompt):
    assert ColumnGroupingRequest.from_instruction(prompt) is None


def test_all_rows_and_columns_preserved_with_numeric_blank_and_style_groups():
    source = workbook_bytes()
    result = WorkbookColumnGrouper().split(
        source, ColumnGroupingRequest(column="Price"), "Catering Menu.xlsx"
    )
    with (
        closing(load_workbook(BytesIO(source))) as before,
        closing(load_workbook(BytesIO(result.file_data))) as after,
    ):
        assert after.sheetnames == ["menus", "Overview", "0.00", "1.50", "2.00", "10.00", "Blank"]
        for name in before.sheetnames:
            assert list(before[name].values) == list(after[name].values)
        original = list(before["menus"].values)
        collected = []
        for name in after.sheetnames[2:]:
            rows = list(after[name].values)
            assert rows[0] == original[0]
            collected.extend(rows[1:])
        assert Counter(collected) == Counter(original[1:])
        assert len(list(after["1.50"].values)) == 3
        assert copy(after["2.00"]["C2"].font) == copy(before["menus"]["C2"].font)
        assert after["2.00"]["F2"].number_format == "0.00"
        assert after["menus"].column_dimensions["C"].width == 32
    assert "6 item rows" in result.summary and "5 sheets" in result.summary


def test_recovers_filename_used_as_sheet_name_only_when_source_is_unique():
    result = WorkbookColumnGrouper().split(
        workbook_bytes(),
        ColumnGroupingRequest(column="Rate", sheet_name="Catering Menu"),
        "Catering Menu.xlsx",
    )
    assert "from 'menus'" in result.summary
    with pytest.raises(DocumentLocatorError):
        WorkbookColumnGrouper().split(
            workbook_bytes(),
            ColumnGroupingRequest(column="Price", sheet_name="Nonexistent"),
            "Catering Menu.xlsx",
        )


def test_ambiguous_sources_require_a_sheet_and_existing_group_tabs_are_not_overwritten():
    with closing(load_workbook(BytesIO(workbook_bytes()))) as book:
        other = book.copy_worksheet(book["menus"])
        other.title = "2.00"
        stream = BytesIO()
        book.save(stream)
    with pytest.raises(DocumentLocatorError, match="Which sheet"):
        WorkbookColumnGrouper().split(
            stream.getvalue(), ColumnGroupingRequest(column="Price"), "source.xlsx"
        )
    result = WorkbookColumnGrouper().split(
        stream.getvalue(), ColumnGroupingRequest(column="Price", sheet_name="menus"), "source.xlsx"
    )
    with closing(load_workbook(BytesIO(result.file_data))) as book:
        assert "2.00 (2)" in book.sheetnames
        assert book["2.00"].max_row == 7


@pytest.mark.parametrize("prices,match", [([], "item rows"), (["=1+1"], "Recalculate")])
def test_incomplete_sources_are_not_returned_as_valid_groups(prices, match):
    with pytest.raises(InvalidDocumentError, match=match):
        WorkbookColumnGrouper().split(
            workbook_bytes(prices), ColumnGroupingRequest(column="Price"), "source.xlsx"
        )


def test_formula_records_reference_the_retained_original_sheet():
    with closing(load_workbook(BytesIO(workbook_bytes()))) as book:
        book["menus"]["G2"] = "=F2*2"
        stream = BytesIO()
        book.save(stream)
    result = WorkbookColumnGrouper().split(
        stream.getvalue(), ColumnGroupingRequest(column="Price"), "source.xlsx"
    )
    with closing(load_workbook(BytesIO(result.file_data))) as book:
        assert book["menus"]["G2"].value == "=F2*2"
        assert book["2.00"]["G2"].value == "='menus'!G2"


@pytest.mark.parametrize(
    "prompt",
    [
        *PROMPTS[:4],
        "The source worksheet is named:\nCatering Menu\n"
        + "Put items with the same price into separate sheets. "
        + "Keep all source rows and columns unchanged. " * 150,
    ],
)
def test_automation_endpoint_copies_rows_instead_of_calling_the_model(client, monkeypatch, prompt):
    def forbidden(*args, **kwargs):
        pytest.fail("Source rows must not be regenerated by AI")

    monkeypatch.setattr(chat_documents.automation_service.chat_service, "chat", forbidden)
    client.post(
        "/users/",
        json={
            "fullname": "Grouping Test",
            "email": "grouping@example.com",
            "password": "test-password",
        },
    )
    token = client.post(
        "/users/login", json={"email": "grouping@example.com", "password": "test-password"}
    ).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    original = workbook_bytes()
    upload = client.post(
        "/chat/documents",
        headers=headers,
        files={
            "file": (
                "Catering Menu.xlsx",
                original,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"message": "Save workbook", "analyze": "false"},
    )
    assert upload.status_code == 200, upload.text
    for _ in range(2):
        response = client.post(
            "/chat/documents/automate",
            headers=headers,
            json={"session_id": upload.json()["session_id"], "instruction": prompt},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == "done", body
        assert "6 item rows" in body["response"]
        download = client.get(
            f"/chat/documents/{body['latest_document_id']}/download", headers=headers
        )
        with closing(load_workbook(BytesIO(download.content))) as book:
            assert book.sheetnames == [
                "menus",
                "Overview",
                "0.00",
                "1.50",
                "2.00",
                "10.00",
                "Blank",
            ]
        assert body["attachments"][0]["source_document_ids"] == [upload.json()["attachment"]["id"]]
    assert (
        client.get(
            f"/chat/documents/{upload.json()['attachment']['id']}/download", headers=headers
        ).content
        == original
    )
