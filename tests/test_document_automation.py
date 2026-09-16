import json
from contextlib import closing
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.api import chat_documents
from app.services.document.document_pipeline_service import DocumentPipelineResult
from app.services.document.exceptions import InvalidDocumentError

CSV_DATA = b"Item,Category,Price\nApple,Fruit,1.25\nCarrot,Vegetable,0.80\n"


def _login(client, email: str) -> dict[str, str]:
    client.post(
        "/users/",
        json={"fullname": "Automation User", "email": email, "password": "secret123"},
    )
    token = client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(["Item", "Category", "Price"])
    sheet.append(["Apple", "Fruit", 1.25])
    sheet.append(["Carrot", "Vegetable", 0.80])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _upload(
    client,
    headers,
    filename: str,
    data: bytes,
    content_type: str,
    *,
    session_id: int | None = None,
) -> dict:
    form = {"message": "Import this source file", "analyze": "false"}
    if session_id is not None:
        form["session_id"] = str(session_id)
    response = client.post(
        "/chat/documents",
        files={"file": (filename, data, content_type)},
        data=form,
        headers=headers,
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_natural_language_csv_automation_runs_three_real_steps(client, monkeypatch):
    headers = _login(client, "automation-csv@example.com")
    uploaded = _upload(client, headers, "inventory.csv", CSV_DATA, "text/csv")
    provider_calls = []

    def plan_with_provider(message, history, reference_history):
        provider_calls.append(message)
        return json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "csv",
                        "operation": "convert_document",
                        "input_file": "inventory.csv",
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                    {
                        "document_type": "xlsx",
                        "operation": "format_workbook",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                    {
                        "document_type": "xlsx",
                        "operation": "filter_column",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {"column": "Category"},
                    },
                ],
            }
        )

    monkeypatch.setattr(chat_documents.automation_service.chat_service, "chat", plan_with_provider)
    response = client.post(
        "/chat/documents/automate",
        json={
            "confirm": True,
            "session_id": uploaded["session_id"],
            "instruction": (
                "Turn inventory.csv into a polished workbook and make Category easy to search."
            ),
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "done"
    assert body["response"].startswith("Done")
    assert len(body["attachments"]) == 3
    assert body["attachments"][-1]["filename"] == "inventory-converted_updated_filter_updated.xlsx"
    assert [step["status"] for step in body["steps"]] == [
        "completed",
        "completed",
        "completed",
    ]
    assert provider_calls and "Available files" in provider_calls[0]
    download = client.get(f"/chat/documents/{body['latest_document_id']}/download", headers=headers)
    with closing(load_workbook(BytesIO(download.content))) as workbook:
        assert workbook.active.auto_filter.ref == "B1:B3"

    history = client.get(f"/chat/sessions/{uploaded['session_id']}", headers=headers).json()[
        "messages"
    ]
    visible_automation_messages = [item["content"] for item in history[-2:]]
    assert visible_automation_messages[0] == (
        "Turn inventory.csv into a polished workbook and make Category easy to search."
    )
    assert visible_automation_messages[1].startswith("Done")
    assert history[-1]["attachments"] == body["attachments"]
    assert history[-1]["automation"]["response"] == body


def test_natural_language_workbook_automation_splits_and_formats(client, monkeypatch):
    headers = _login(client, "automation-workbook@example.com")
    uploaded = _upload(
        client,
        headers,
        "products.xlsx",
        _workbook_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "xlsx",
                        "operation": "split_by_category",
                        "input_file": "products.xlsx",
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                    {
                        "document_type": "xlsx",
                        "operation": "format_workbook",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                ],
            }
        ),
    )
    response = client.post(
        "/chat/documents/automate",
        json={
            "confirm": True,
            "session_id": uploaded["session_id"],
            "instruction": (
                "Organize products.xlsx into a tab for each category and make it presentation-ready."
            ),
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "done"
    assert len(body["attachments"]) == 2
    assert body["attachments"][-1]["filename"] == "products_category_wise_updated.xlsx"
    download = client.get(f"/chat/documents/{body['latest_document_id']}/download", headers=headers)
    with closing(load_workbook(BytesIO(download.content))) as workbook:
        assert set(workbook.sheetnames) == {"Fruit", "Vegetable"}
        assert all(sheet.freeze_panes == "A2" for sheet in workbook.worksheets)


def test_ambiguous_automation_returns_and_persists_one_clarifying_question(client, monkeypatch):
    headers = _login(client, "automation-clarify@example.com")
    first = _upload(client, headers, "first.csv", CSV_DATA, "text/csv")
    _upload(
        client,
        headers,
        "second.csv",
        CSV_DATA,
        "text/csv",
        session_id=first["session_id"],
    )
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "csv",
                        "operation": "convert_document",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {},
                    }
                ],
            }
        ),
    )

    response = client.post(
        "/chat/documents/automate",
        json={
            "confirm": True,
            "session_id": first["session_id"],
            "instruction": "Make it a workbook.",
        },
        headers=headers,
    )

    assert response.status_code == 200
    assert response.json()["status"] == "clarification_required"
    assert response.json()["attachments"] == []
    assert response.json()["response"] == "Which uploaded document should I use for this request?"


def test_partial_failure_never_claims_done(client, monkeypatch):
    headers = _login(client, "automation-partial@example.com")
    uploaded = _upload(client, headers, "inventory.csv", CSV_DATA, "text/csv")
    monkeypatch.setattr(
        chat_documents.automation_service.chat_service,
        "chat",
        lambda *args, **kwargs: json.dumps(
            {
                "status": "ready",
                "clarifying_question": None,
                "steps": [
                    {
                        "document_type": "csv",
                        "operation": "convert_document",
                        "input_file": "inventory.csv",
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                    {
                        "document_type": "xlsx",
                        "operation": "format_workbook",
                        "input_file": None,
                        "output_type": "xlsx",
                        "parameters": {},
                    },
                ],
            }
        ),
    )
    original_execute = chat_documents.pipeline_service.execute_step
    calls = 0

    def fail_second_step(step, source) -> DocumentPipelineResult:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise InvalidDocumentError("The generated workbook could not be validated.")
        return original_execute(step, source)

    monkeypatch.setattr(chat_documents.pipeline_service, "execute_step", fail_second_step)
    response = client.post(
        "/chat/documents/automate",
        json={
            "confirm": True,
            "session_id": uploaded["session_id"],
            "instruction": "Create a polished workbook from inventory.csv.",
        },
        headers=headers,
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "partial"
    assert body["attachments"][0]["filename"].endswith(".xlsx")
    assert "couldn't finish" in body["response"]
    assert "Done" not in body["response"]
