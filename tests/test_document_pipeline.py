from contextlib import closing
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.conversion.document_conversion_service import (
    DocumentConversionService,
    LibreOfficeConverter,
    UnsupportedDocumentConversionError,
)
from app.services.document.document_generation_service import GeneratedDocument
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.document_pipeline_service import (
    DocumentPipelineService,
    PipelineDocument,
)
from app.services.document.document_validation_service import GeneratedDocumentValidationError
from app.services.document.exceptions import DocumentProcessingUnavailableError
from app.services.pdf.pdf_generator import PdfGenerator

CSV_DATA = b"Item,Category,Price\nApple,Fruit,1.25\nCarrot,Vegetable,0.80\n"


def _login(client, email: str) -> str:
    client.post(
        "/users/", json={"fullname": "Pipeline User", "email": email, "password": "secret123"}
    )
    return client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]


def _run(service: DocumentPipelineService, instruction: str) -> list:
    steps = service.plan(instruction, input_filename="items.csv")
    current = PipelineDocument(CSV_DATA, "items.csv")
    results = []
    for step in steps:
        result = service.execute_step(step, current)
        results.append(result)
        current = PipelineDocument(result.document.file_data, result.document.filename)
    return results


def test_real_csv_to_xlsx_conversion_reopens_with_preserved_rows():
    generated = DocumentConversionService().convert(CSV_DATA, "items.csv", DocumentType.XLSX)

    assert generated.filename == "items-converted.xlsx"
    with closing(
        load_workbook(BytesIO(generated.file_data), read_only=True, data_only=True)
    ) as workbook:
        assert workbook.active.title == "CSV data"
        assert list(workbook.active.values) == [
            ("Item", "Category", "Price"),
            ("Apple", "Fruit", "1.25"),
            ("Carrot", "Vegetable", "0.80"),
        ]


def test_xlsx_to_csv_and_text_markdown_conversion_pairs_are_real_and_reopenable():
    workbook = Workbook()
    workbook.active.append(["Item", "Count"])
    workbook.active.append(["Apple", 2])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    csv_document = DocumentConversionService().convert(
        source.getvalue(), "items.xlsx", DocumentType.CSV
    )
    assert csv_document.file_data.decode("utf-8-sig").splitlines() == [
        "Item,Count",
        "Apple,2",
    ]
    markdown = DocumentConversionService().convert(
        b"# Inventory\n\nApples", "notes.txt", DocumentType.MARKDOWN
    )
    assert markdown.filename == "notes-converted.md"
    plain_text = DocumentConversionService().convert(
        markdown.file_data, markdown.filename, DocumentType.TXT
    )
    assert plain_text.file_data == b"# Inventory\n\nApples"


def test_multisheet_xlsx_to_csv_is_rejected_instead_of_dropping_sheets():
    workbook = Workbook()
    workbook.active.append(["First"])
    workbook.create_sheet("Second").append(["Second"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    with pytest.raises(UnsupportedDocumentConversionError, match="exactly one"):
        DocumentConversionService().convert(source.getvalue(), "multiple.xlsx", DocumentType.CSV)


def test_libreoffice_conversion_uses_generated_pdf_and_validates_it(monkeypatch):
    pdf_bytes = PdfGenerator().generate("Converted office document")

    def run(command, **kwargs):
        source = Path(command[-1])
        assert source.read_bytes() == b"office bytes"
        assert command[2].startswith("-env:UserInstallation=file://")
        source.with_suffix(".pdf").write_bytes(pdf_bytes)
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    monkeypatch.setattr("app.services.conversion.document_conversion_service.subprocess.run", run)
    converted = LibreOfficeConverter(binary="soffice").convert_to_pdf(
        b"office bytes", "report.docx"
    )

    assert converted == pdf_bytes


def test_two_step_pipeline_runs_in_declared_order():
    results = _run(
        DocumentPipelineService(),
        "Convert items.csv to XLSX, then format the workbook professionally",
    )

    assert [result.step.intent.operation for result in results] == [
        DocumentOperation.CONVERT_DOCUMENT,
        DocumentOperation.FORMAT_WORKBOOK,
    ]
    assert results[-1].document.document_type == DocumentType.XLSX


def test_three_step_pipeline_passes_each_output_to_the_next_step():
    results = _run(
        DocumentPipelineService(),
        "Convert items.csv to XLSX, then format the workbook professionally, "
        "then add a filter to the Category column",
    )

    assert [result.step.intent.operation for result in results] == [
        DocumentOperation.CONVERT_DOCUMENT,
        DocumentOperation.FORMAT_WORKBOOK,
        DocumentOperation.FILTER_COLUMN,
    ]
    with closing(load_workbook(BytesIO(results[-1].document.file_data))) as workbook:
        assert workbook.active.auto_filter.ref == "B1:B3"


def test_invalid_intermediate_output_is_rejected_before_it_can_continue():
    class InvalidConversionService(DocumentConversionService):
        def convert(self, file_data, filename, output_type):
            return GeneratedDocument(
                file_data=b"not an xlsx",
                filename="broken.xlsx",
                content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                document_type=DocumentType.XLSX,
            )

    service = DocumentPipelineService(conversion_service=InvalidConversionService())
    step = service.plan("Convert items.csv to XLSX", input_filename="items.csv")[0]

    with pytest.raises(GeneratedDocumentValidationError, match="could not be reopened"):
        service.execute_step(step, PipelineDocument(CSV_DATA, "items.csv"))


def test_unsupported_conversion_is_detected_while_planning():
    with pytest.raises(UnsupportedDocumentConversionError, match="CSV to PDF"):
        DocumentPipelineService().plan("Convert items.csv to PDF", input_filename="items.csv")


def test_office_to_pdf_reports_libreoffice_requirement_when_unavailable(monkeypatch):
    monkeypatch.setattr(LibreOfficeConverter, "_discover_binary", classmethod(lambda cls: None))
    monkeypatch.setattr("app.config.settings.LIBREOFFICE_BINARY", "")
    converter = LibreOfficeConverter()

    with pytest.raises(DocumentProcessingUnavailableError, match="requires headless LibreOffice"):
        converter.convert_to_pdf(b"office bytes", "report.docx")


def test_pipeline_endpoint_persists_latest_document_between_requests(client):
    token = _login(client, "pipeline-latest@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("items.csv", CSV_DATA, "text/csv")},
        data={"message": "Import inventory", "analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200
    session_id = uploaded.json()["session_id"]

    pipeline = client.post(
        "/chat/documents/pipeline",
        json={
            "session_id": session_id,
            "instruction": (
                "Convert items.csv to XLSX, then format the workbook professionally, "
                "then add a filter to the Category column"
            ),
        },
        headers=headers,
    )
    assert pipeline.status_code == 200, pipeline.text
    body = pipeline.json()
    assert [step["operation"] for step in body["steps"]] == [
        "convert_document",
        "format_workbook",
        "filter_column",
    ]
    assert body["steps"][1]["source_document_id"] == body["steps"][0]["output_document_id"]
    assert body["steps"][2]["source_document_id"] == body["steps"][1]["output_document_id"]
    assert body["latest_document_id"] == body["attachment"]["id"]

    follow_up = client.post(
        "/chat/documents/pipeline",
        json={"session_id": session_id, "instruction": "Convert the workbook to CSV"},
        headers=headers,
    )
    assert follow_up.status_code == 200, follow_up.text
    assert follow_up.json()["steps"][0]["source_document_id"] == body["latest_document_id"]
    assert follow_up.json()["attachment"]["filename"].endswith(".csv")


def test_pipeline_endpoint_reports_missing_input_and_unsupported_pair(client):
    token = _login(client, "pipeline-errors@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    session_id = client.post(
        "/chat/sessions", json={"title": "Pipeline errors"}, headers=headers
    ).json()["id"]

    missing = client.post(
        "/chat/documents/pipeline",
        json={"session_id": session_id, "instruction": "Convert to XLSX"},
        headers=headers,
    )
    assert missing.status_code == 422
    assert "Upload or select" in missing.json()["detail"]

    uploaded = client.post(
        "/chat/documents",
        files={"file": ("items.csv", CSV_DATA, "text/csv")},
        data={"session_id": str(session_id), "analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200
    unsupported = client.post(
        "/chat/documents/pipeline",
        json={"session_id": session_id, "instruction": "Convert items.csv to PDF"},
        headers=headers,
    )
    assert unsupported.status_code == 422
    assert "CSV to PDF" in unsupported.json()["detail"]
