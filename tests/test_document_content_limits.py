import asyncio
import json
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from app.models.chat import ChatDocumentAttachment, ChatMessageRecord, ChatSession
from app.services.chat_document_service import ChatDocumentService, InvalidDocumentError
from app.services.chat_service import ChatService


def _headers(client):
    email = "document-content-limits@example.com"
    registered = client.post(
        "/users/", json={"fullname": "Document User", "email": email, "password": "secret123"}
    )
    assert registered.status_code in (200, 201)
    login = client.post("/users/login", json={"email": email, "password": "secret123"})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def test_generation_source_is_bounded_and_discloses_omitted_content(monkeypatch):
    prompts = []

    async def completed(self, message, history, references, **kwargs):
        prompts.append(message)
        return json.dumps(
            {
                "title": "Source report",
                "paragraphs": ["Summary based only on the supplied excerpt."],
                "bullet_lists": [],
                "tables": [],
                "sections": [],
            }
        )

    monkeypatch.setattr(ChatService, "complete_chat", completed)
    source = "VISIBLE SOURCE\n" + "A" * 200_000 + "OMITTED SOURCE TAIL"
    sources = [source, "ANOTHER OMITTED DOCUMENT"]
    original_sources = sources.copy()
    service = ChatDocumentService()
    request = service.generate_content("Make a report", sources, [])
    result = asyncio.run(request)
    assert result.paragraphs == ["Summary based only on the supplied excerpt."]
    prompt = prompts[0]
    source_block = prompt.split("--- UNTRUSTED DOCUMENT DATA ---\n", 1)[1].split(
        "\n--- END UNTRUSTED DOCUMENT DATA ---", 1
    )[0]
    _, excerpt = source_block.split("\n\n", 1)
    assert excerpt == source[: service.MAX_DOCUMENT_CONTEXT_CHARS]
    assert "Source excerpt truncated: first 30,000" in prompt
    assert "The remaining source content was not provided" in prompt
    assert "do not claim to have read the full document" in prompt
    assert "VISIBLE SOURCE" in prompt
    assert "OMITTED SOURCE TAIL" not in prompt
    assert "ANOTHER OMITTED DOCUMENT" not in prompt
    assert sources == original_sources


def test_large_summary_processes_every_chunk_and_reports_full_coverage(monkeypatch):
    prompts = []

    async def completed(self, message, history, references, **kwargs):
        del self, history, references
        prompts.append(message)
        return "chunk result"

    monkeypatch.setattr(ChatService, "complete_chat", completed)
    source = "VISIBLE SOURCE\n" + "A" * 65_000 + "FINAL SOURCE TAIL"
    result = asyncio.run(ChatDocumentService().summarize(source, "Summarize this file"))

    assert result == "chunk result"
    assert len(prompts) == 4  # three map calls and one reduce call
    supplied_chunks = "\n".join(prompts[:-1])
    assert "VISIBLE SOURCE" in supplied_chunks
    assert "FINAL SOURCE TAIL" in supplied_chunks
    assert "processed 3 of 3 chunks" in prompts[-1]
    assert "--- UNTRUSTED DOCUMENT DATA ---" in supplied_chunks


def test_source_budget_is_shared_across_documents_and_marks_partial_later_document():
    sources = ["A" * 20_000, "B" * 20_000, "NEVER INCLUDED"]
    excerpt = ChatDocumentService._source_excerpt(sources)
    assert excerpt.count("A") == 20_000
    assert excerpt.count("B") == 9_998
    assert "NEVER INCLUDED" not in excerpt
    assert "40,018 characters supplied" in excerpt


def test_short_and_exact_limit_sources_are_unchanged():
    assert ChatDocumentService._source_excerpt([]) == ""
    assert ChatDocumentService._source_excerpt(["First", "Second"]) == "First\n\nSecond"
    source = "A" * ChatDocumentService.MAX_DOCUMENT_CONTEXT_CHARS
    assert ChatDocumentService._source_excerpt([source]) == source


@pytest.mark.parametrize("text", ["தமிழ் உணவு", "漢字", "العربية", "Ready 😊"])
@pytest.mark.parametrize("plain_text", [False, True])
def test_pdf_rejects_unrepresentable_text_instead_of_silent_black_squares(text, plain_text):
    with pytest.raises(InvalidDocumentError, match=r"choose Word \(\.docx\)"):
        ChatDocumentService().render(text, "pdf", plain_text=plain_text)


def test_pdf_keeps_supported_winansi_characters():
    text = "Café – €12.50; voilà!"
    content, _, _ = ChatDocumentService().render(text, "pdf", plain_text=True)
    extracted = "\n".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    assert text in extracted


def test_docx_preserves_full_unicode_source_without_ai_context_limit():
    text = "தமிழ் 漢字 العربية 😊\n" + "Source line\n" * 3000 + "FINAL SOURCE LINE"
    assert len(text) > ChatDocumentService.MAX_DOCUMENT_CONTEXT_CHARS
    content, _, _ = ChatDocumentService().render(text, "docx", plain_text=True)
    extracted = "\n".join(paragraph.text for paragraph in Document(BytesIO(content)).paragraphs)
    assert extracted == text


def test_unrepresentable_pdf_text_export_rejects_before_saving_artifacts(client, db_session):
    response = client.post(
        "/chat/documents/generate",
        json={"mode": "export", "instruction": "தமிழ் 漢字", "output_format": "pdf"},
        headers=_headers(client),
    )
    assert response.status_code == 422, response.text
    assert "choose Word (.docx)" in response.json()["detail"]
    assert db_session.query(ChatSession).count() == 0
    assert db_session.query(ChatMessageRecord).count() == 0
    assert db_session.query(ChatDocumentAttachment).count() == 0


def test_unrepresentable_pdf_source_export_preserves_original_then_allows_docx(client, db_session):
    headers = _headers(client)
    text = "தமிழ் 漢字 العربية 😊\n" + "Source line\n" * 3000 + "FINAL SOURCE LINE"
    original_bytes = text.encode("utf-8")
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("source.txt", original_bytes, "text/plain")},
        data={"analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text
    upload = uploaded.json()
    source_id = upload["attachment"]["id"]
    payload = {"mode": "export", "source_document_id": source_id, "output_format": "pdf"}
    rejected = client.post("/chat/documents/generate", json=payload, headers=headers)
    assert rejected.status_code == 422, rejected.text
    assert "choose Word (.docx)" in rejected.json()["detail"]
    assert db_session.query(ChatSession).count() == 1
    assert db_session.query(ChatMessageRecord).count() == 2
    assert db_session.query(ChatDocumentAttachment).count() == 1
    original = db_session.query(ChatDocumentAttachment).one()
    assert original.raw_text == text
    assert original.file_data == original_bytes
    assert (
        client.get(f"/chat/documents/{source_id}/download", headers=headers).content
        == original_bytes
    )

    exported = client.post(
        "/chat/documents/generate", json={**payload, "output_format": "docx"}, headers=headers
    )
    assert exported.status_code == 200, exported.text
    exported_id = exported.json()["attachment"]["id"]
    download = client.get(f"/chat/documents/{exported_id}/download", headers=headers)
    extracted = "\n".join(
        paragraph.text for paragraph in Document(BytesIO(download.content)).paragraphs
    )
    assert extracted == text
