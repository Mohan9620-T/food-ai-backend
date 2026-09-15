import asyncio
import csv
import json
from io import BytesIO, StringIO

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pydantic import ValidationError
from pypdf import PdfReader

from app.models.chat import ChatDocumentAttachment
from app.services.chat_document_service import ChatDocumentService
from app.services.chat_service import ChatModelUnavailableError, ChatService
from app.services.document.document_generation_service import DocumentGenerationService
from app.services.document.extraction_models import (
    GeneratedSectionContent,
    GeneratedTableContent,
    StructuredDocumentContent,
)


def _headers(client, email: str) -> dict[str, str]:
    response = client.post(
        "/users/",
        json={"fullname": "Phase Three", "email": email, "password": "secret123"},
    )
    assert response.status_code in (200, 201)
    token = client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


def _structured_content() -> StructuredDocumentContent:
    return StructuredDocumentContent(
        title="Professional Project Report",
        paragraphs=["This report summarizes the approved project facts."],
        bullet_lists=[["Protect the original source", "Validate every generated file"]],
        tables=[
            GeneratedTableContent(
                title="Milestones",
                headers=["Code", "Status"],
                rows=[["M1", "Complete"], ["M2", "Planned"]],
            )
        ],
        sections=[
            GeneratedSectionContent(
                heading="Next Steps",
                paragraphs=["Continue only after review."],
                bullet_lists=[["Review", "Approve"]],
                tables=[],
            )
        ],
    )


@pytest.mark.parametrize(
    ("output_format", "content_type"),
    [
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("pdf", "application/pdf"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("csv", "text/csv"),
        ("pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        ("txt", "text/plain"),
        ("markdown", "text/markdown"),
    ],
)
def test_structured_content_renders_and_reopens_in_every_creation_format(
    output_format, content_type
):
    generated = DocumentGenerationService().generate(_structured_content(), output_format)

    assert generated.content_type == content_type
    assert generated.filename.startswith("Professional-Project-Report")
    if output_format == "docx":
        document = Document(BytesIO(generated.file_data))
        assert document.paragraphs[0].text == "Professional Project Report"
        assert document.tables[0].cell(1, 0).text == "M1"
    elif output_format == "pdf":
        text = "\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(generated.file_data)).pages
        )
        assert "Professional Project Report" in text
        assert "M1" in text
    elif output_format == "xlsx":
        workbook = load_workbook(BytesIO(generated.file_data), data_only=False)
        try:
            assert workbook.sheetnames == ["Overview", "Milestones"]
            assert workbook["Milestones"]["A2"].value == "M1"
        finally:
            workbook.close()
    elif output_format == "csv":
        rows = list(csv.reader(StringIO(generated.file_data.decode("utf-8-sig"))))
        assert ["Code", "Status"] in rows
        assert ["M1", "Complete"] in rows
    elif output_format == "pptx":
        presentation = Presentation(BytesIO(generated.file_data))
        assert presentation.slides[0].shapes.title.text == "Professional Project Report"
        assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)
    elif output_format == "txt":
        text = generated.file_data.decode("utf-8")
        assert "Professional Project Report" in text
        assert "M1\tComplete" in text
    else:
        text = generated.file_data.decode("utf-8")
        assert text.startswith("# Professional Project Report")
        assert "| M1 | Complete |" in text


def test_ai_generation_requires_and_parses_the_structured_json_contract(monkeypatch):
    captured: list[str] = []
    payload = _structured_content().model_dump(mode="json")

    async def completed(self, message, history, reference_history, **kwargs):
        captured.append(message)
        return f"```json\n{json.dumps(payload)}\n```"

    monkeypatch.setattr(ChatService, "complete_chat", completed)
    result = asyncio.run(
        ChatDocumentService().generate_content(
            "Create a professional project report in Word.", [], []
        )
    )

    assert result == _structured_content()
    assert '"sections"' in captured[0]
    assert "Do not return Markdown fences" in captured[0]
    assert "file bytes" in captured[0]


def test_invalid_ai_document_structure_is_rejected_without_rendering(monkeypatch):
    async def invalid(self, message, history, reference_history, **kwargs):
        return '{"title":"Incomplete","paragraphs":[]}'

    monkeypatch.setattr(ChatService, "complete_chat", invalid)
    with pytest.raises(ChatModelUnavailableError, match="invalid document structure"):
        asyncio.run(ChatDocumentService().generate_content("Create a report", [], []))


def test_structured_contract_rejects_mismatched_table_rows_and_empty_content():
    with pytest.raises(ValidationError, match="same width"):
        GeneratedTableContent(title="Broken", headers=["One", "Two"], rows=[["only-one"]])
    with pytest.raises(ValidationError, match="at least one content block"):
        StructuredDocumentContent(title="Empty")


def test_professional_word_request_returns_stored_attached_downloadable_document(
    client, db_session, monkeypatch
):
    headers = _headers(client, "phase3-word@example.com")

    async def generated(self, instruction, summaries, history, profile=None):
        assert instruction == "Create a professional project report in Word."
        return _structured_content()

    monkeypatch.setattr(ChatDocumentService, "generate_content", generated)
    response = client.post(
        "/chat/documents/generate",
        json={
            "instruction": "Create a professional project report in Word.",
            "output_format": "docx",
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["attachment"]["filename"] == "Professional-Project-Report.docx"
    assert body["attachment"]["kind"] == "generated"
    assert body["response"].startswith("# Professional Project Report")
    attachment = db_session.get(ChatDocumentAttachment, body["attachment"]["id"])
    assert attachment is not None
    assert attachment.raw_text == body["response"]
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    document = Document(BytesIO(download.content))
    assert document.tables[0].cell(2, 1).text == "Planned"


def test_selected_source_document_is_extracted_for_ai_creation_and_not_overwritten(
    client, db_session, monkeypatch
):
    headers = _headers(client, "phase3-source@example.com")
    original = b"Customer,Total\nAcme,120\nBeta,90\n"
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("customers.csv", original, "text/csv")},
        data={"analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text
    source_id = uploaded.json()["attachment"]["id"]
    captured = []

    async def generated(self, instruction, summaries, history, profile=None, documents=None):
        captured.extend(documents or [])
        return _structured_content()

    monkeypatch.setattr(ChatDocumentService, "generate_content", generated)
    response = client.post(
        "/chat/documents/generate",
        json={
            "session_id": uploaded.json()["session_id"],
            "source_document_id": source_id,
            "instruction": "Create an Excel project report from these customer records.",
            "output_format": "xlsx",
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    assert captured[0].tables[0].rows[1] == ("Acme", "120")
    assert response.json()["attachment"]["filename"] == "customers_generated.xlsx"
    source = db_session.get(ChatDocumentAttachment, source_id)
    assert source is not None
    assert source.kind == "uploaded"
    assert source.file_data == original
    assert response.json()["attachment"]["id"] != source_id
