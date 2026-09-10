from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

from app.services.chat_document_service import (
    ChatDocumentService,
    DocumentProcessingUnavailableError,
)
from app.services.chat_service import ChatModelUnavailableError

DOCUMENT_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DISH_EXPANSION_MESSAGE = (
    "Create separate rows based on Regular, Easy to Chew, Soft & Bite, Minced & Moist and Pureed."
)


def _login(client, email: str) -> str:
    client.post("/users/", json={"fullname": "Doc User", "email": email, "password": "secret123"})
    return client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]


def _dish_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dish Master"
    sheet.append(
        [
            "Id",
            "Code",
            "Label",
            "Regular",
            "Easy to chew",
            "Soft and bite-sized",
            "Minced and Moist",
            "Pureed",
            "Price",
        ]
    )
    sheet.append([1, "DISH-A", "Dish A", "x", None, " X ", None, None, 12.5])
    sheet.append([2, "DISH-B", "Dish B", None, None, None, None, None, 8.0])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_extracts_txt_docx_csv_and_xlsx():
    service = ChatDocumentService()
    assert service.extract(b"Calories: 1800", "notes.txt") == "Calories: 1800"
    assert "protein" in service.extract(b"name,value\nprotein,90", "data.csv")
    document = Document()
    document.add_paragraph("Allergy: peanuts")
    buffer = BytesIO()
    document.save(buffer)
    assert "Allergy: peanuts" in service.extract(buffer.getvalue(), "report.docx")


def test_upload_document_persists_attachment(client, monkeypatch):
    token = _login(client, "document-upload@example.com")
    monkeypatch.setattr(
        ChatDocumentService, "summarize", AsyncMock(return_value="Structured nutrition summary")
    )
    response = client.post(
        "/chat/documents",
        files={"file": ("notes.txt", b"Calories: 1800", "text/plain")},
        data={"message": "Read my notes"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["attachment"]["filename"] == "notes.txt"
    assert body["attachment"]["kind"] == "uploaded"
    history = client.get(
        f"/chat/sessions/{body['session_id']}", headers={"Authorization": f"Bearer {token}"}
    ).json()
    assert history["messages"][0]["document_attachment"]["filename"] == "notes.txt"


def test_document_upload_rejects_invalid_and_empty_files(client):
    token = _login(client, "document-validation@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    invalid = client.post(
        "/chat/documents",
        files={"file": ("bad.exe", b"x", "application/octet-stream")},
        headers=headers,
    )
    empty = client.post(
        "/chat/documents", files={"file": ("empty.txt", b"", "text/plain")}, headers=headers
    )
    assert invalid.status_code == 415
    assert empty.status_code == 422


def test_generate_and_download_enforces_ownership(client, monkeypatch):
    first_token = _login(client, "document-owner@example.com")
    first_headers = {"Authorization": f"Bearer {first_token}"}
    session_id = client.post(
        "/chat/sessions", json={"title": "Documents"}, headers=first_headers
    ).json()["id"]
    monkeypatch.setattr(
        ChatDocumentService,
        "generate_content",
        AsyncMock(return_value="# Doctor Summary\nBalanced nutrition plan"),
    )
    generated = client.post(
        "/chat/documents/generate",
        json={"session_id": session_id, "instruction": "Doctor summary", "output_format": "pdf"},
        headers=first_headers,
    )
    assert generated.status_code == 200
    document_id = generated.json()["attachment"]["id"]
    download = client.get(f"/chat/documents/{document_id}/download", headers=first_headers)
    assert download.status_code == 200
    assert download.headers["content-type"] == "application/pdf"
    second_token = _login(client, "document-stranger@example.com")
    denied = client.get(
        f"/chat/documents/{document_id}/download",
        headers={"Authorization": f"Bearer {second_token}"},
    )
    assert denied.status_code == 404


def _document_bytes(extension):
    content = "Product details: oats 100"
    buffer = BytesIO()
    if extension == "pdf":
        canvas = Canvas(buffer)
        canvas.drawString(72, 750, content)
        canvas.save()
    elif extension == "docx":
        document = Document()
        document.add_heading("Product details", level=1)
        document.add_paragraph(content)
        document.save(buffer)
    elif extension == "xlsx":
        workbook = Workbook()
        workbook.active.append([content])
        workbook.save(buffer)
        workbook.close()
    else:
        return content.encode("utf-8")
    return buffer.getvalue()


@pytest.mark.parametrize("extension", ["pdf", "docx", "txt", "csv", "xlsx"])
def test_all_supported_uploads_can_be_reloaded_and_downloaded(client, monkeypatch, extension):
    token = _login(client, f"upload-{extension}@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    captured = []

    async def summarize(self, raw, note):
        captured.append((raw, note))
        return "# Product details\nOats: 100"

    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)
    file_data = _document_bytes(extension)
    result = client.post(
        "/chat/documents",
        files={"file": (f"REPORT.{extension.upper()}", file_data, "application/octet-stream")},
        data={"message": "Show the headings and details"},
        headers=headers,
    )
    assert result.status_code == 200, result.text
    body = result.json()
    assert "oats 100" in captured[0][0]
    assert captured[0][1] == "Show the headings and details"
    reloaded = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    assert len(reloaded["messages"]) == 2
    assert reloaded["messages"][0]["document_attachment"]["id"] == body["attachment"]["id"]
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.content == file_data


@pytest.mark.parametrize("mime", ["application/vnd.ms-excel", "text/plain", "text/csv"])
def test_csv_accepts_windows_browser_mime_types(client, monkeypatch, mime):
    token = _login(client, "csv-mime@example.com")
    monkeypatch.setattr(ChatDocumentService, "summarize", AsyncMock(return_value="Oats 100"))
    response = client.post(
        "/chat/documents",
        files={"file": ("data.csv", b"name,value\noats,100", mime)},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200


def test_xlsx_formatting_instruction_returns_downloadable_updated_workbook_without_ai(
    client, monkeypatch
):
    token = _login(client, "spreadsheet-format@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    first = workbook.active
    first.title = "Sales"
    first.append(["Item", "Amount"])
    first.append(["Oats", 100])
    first.append(["Rice", 250])
    first["B4"] = "=SUM(B2:B3)"
    second = workbook.create_sheet("Notes")
    second.append(["Status", "Comment"])
    second.append(["Ready", "This is a long note that should be wrapped"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    monkeypatch.setattr(
        ChatDocumentService,
        "summarize",
        AsyncMock(side_effect=AssertionError("Spreadsheet formatting must not call AI")),
    )
    response = client.post(
        "/chat/documents",
        files={
            "file": (
                "sales.xlsx",
                source.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "message": "Center align all cells, style the headers, auto fit columns, and wrap text"
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "skipped"
    assert body["attachment"]["kind"] == "generated"
    assert body["attachment"]["filename"] == "sales_updated.xlsx"
    assert "original upload is unchanged" in body["response"]
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert revised.sheetnames == ["Sales", "Notes"]
        assert revised["Sales"]["A2"].value == "Oats"
        assert revised["Sales"]["B3"].value == 250
        assert revised["Sales"]["B4"].value == "=SUM(B2:B3)"
        assert revised["Sales"]["A2"].alignment.horizontal == "center"
        assert revised["Notes"]["B2"].alignment.wrap_text is True
        assert revised["Sales"]["A1"].font.bold is True
        assert revised["Sales"]["A1"].fill.fgColor.rgb == "001F4E78"
        assert revised["Sales"].column_dimensions["A"].width >= 10
    finally:
        revised.close()

    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    assert history["messages"][0]["document_attachment"]["kind"] == "uploaded"
    assert history["messages"][1]["document_attachment"]["kind"] == "generated"


def test_xlsx_category_instruction_creates_one_sheet_per_category(client, monkeypatch):
    token = _login(client, "spreadsheet-categories@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    source_sheet = workbook.active
    source_sheet.title = "Items"
    source_sheet.append(["Item", "Category", "Price"])
    source_sheet.append(["Apple", "Fruit", 20])
    source_sheet.append(["Carrot", "Vegetable", 30])
    source_sheet.append(["Banana", "Fruit", 40])
    source_sheet.append(["Mystery item", None, 50])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    monkeypatch.setattr(
        ChatDocumentService,
        "summarize",
        AsyncMock(side_effect=AssertionError("Category splitting must not call AI")),
    )
    response = client.post(
        "/chat/documents",
        files={
            "file": (
                "inventory.xlsx",
                source.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={
            "message": (
                "I have many categories. Split the item list category-wise and create a "
                "separate sheet for each category."
            )
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["attachment"]["filename"] == "inventory_category_wise.xlsx"
    assert "3 category worksheets" in body["response"]
    assert "separate sheet for each category" in body["response"]
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert revised.sheetnames == ["Fruit", "Vegetable", "Uncategorized"]
        assert list(revised["Fruit"].values) == [
            ("Item", "Category", "Price"),
            ("Apple", "Fruit", 20),
            ("Banana", "Fruit", 40),
        ]
        assert list(revised["Vegetable"].values) == [
            ("Item", "Category", "Price"),
            ("Carrot", "Vegetable", 30),
        ]
        assert revised["Uncategorized"]["A2"].value == "Mystery item"
        assert revised["Fruit"].freeze_panes == "A2"
        assert revised["Fruit"]["A1"].font.bold is True
    finally:
        revised.close()


def test_xlsx_filter_upload_returns_downloadable_workbook_without_ai(client, monkeypatch):
    token = _login(client, "spreadsheet-filter@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    sheet.append(["ItemCode", "ItemName", "Category*", "Uom"])
    sheet.append(["BNS001", "Coffee", "Beverages", "CTN"])
    sheet.append(["BNS002", "Tea", "Beverages", "PKT"])
    sheet.append(["CHC001", "Chicken", "Meat", "KG"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    original_bytes = source.getvalue()
    summarize = AsyncMock(side_effect=AssertionError("Filter operations must not call AI"))
    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)

    response = client.post(
        "/chat/documents",
        files={"file": ("items.xlsx", original_bytes, DOCUMENT_XLSX_MIME)},
        data={
            "message": (
                "The Excel sheet contains different types of Category values. "
                "Please add a filter only to the Category* column."
            )
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "skipped"
    assert body["attachment"]["kind"] == "generated"
    assert body["attachment"]["filename"] == "items_filter_updated.xlsx"
    assert "requested column filter" in body["response"]
    assert summarize.await_count == 0
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert revised["Items"].auto_filter.ref == "C1:C4"
        assert list(revised["Items"].values) == [
            ("ItemCode", "ItemName", "Category*", "Uom"),
            ("BNS001", "Coffee", "Beverages", "CTN"),
            ("BNS002", "Tea", "Beverages", "PKT"),
            ("CHC001", "Chicken", "Meat", "KG"),
        ]
    finally:
        revised.close()
    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    original_id = history["messages"][0]["document_attachment"]["id"]
    original = client.get(f"/chat/documents/{original_id}/download", headers=headers)
    assert original.content == original_bytes


def test_dish_category_expansion_upload_is_stored_and_downloadable_without_ai(client, monkeypatch):
    from app.api import chat_documents

    token = _login(client, "dish-category-expansion@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    original_bytes = _dish_workbook_bytes()
    summarize = AsyncMock(
        side_effect=AssertionError("Dish category expansion must not call document AI")
    )
    provider_call = MagicMock(
        side_effect=AssertionError("Dish category expansion must not call a provider")
    )
    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)
    monkeypatch.setattr(chat_documents.service.chat_service, "stream_chat", provider_call)

    response = client.post(
        "/chat/documents",
        files={"file": ("dish-master.xlsx", original_bytes, DOCUMENT_XLSX_MIME)},
        data={"message": DISH_EXPANSION_MESSAGE},
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "skipped"
    assert body["attachment"]["kind"] == "generated"
    assert body["attachment"]["filename"] == "dish-master_category_expanded.xlsx"
    assert "processed 2 original dish rows" in body["response"]
    assert "generated 3 category rows" in body["response"]
    assert summarize.await_count == 0
    provider_call.assert_not_called()

    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(DOCUMENT_XLSX_MIME)
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        rows = list(revised["Dish Master"].iter_rows(min_row=2, values_only=True))
        assert len(rows) == 3
        assert [row[1] for row in rows] == ["DISH-A-7R", "DISH-A-6SB", "DISH-B"]
        assert [row[3:8] for row in rows] == [
            ("x", None, None, None, None),
            (None, None, "x", None, None),
            (None, None, None, None, None),
        ]
    finally:
        revised.close()

    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    original_id = history["messages"][0]["document_attachment"]["id"]
    original = client.get(f"/chat/documents/{original_id}/download", headers=headers)
    assert original.status_code == 200
    assert original.content == original_bytes


def test_dish_category_expansion_missing_column_returns_clean_error_without_ai(client, monkeypatch):
    token = _login(client, "dish-category-expansion-missing@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    workbook.active.append(
        ["Code", "Label", "Regular", "Easy to chew", "Soft and bite-sized", "Pureed"]
    )
    workbook.active.append(["DISH-A", "Dish A", "x", None, None, None])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    summarize = AsyncMock(
        side_effect=AssertionError("Dish category expansion must not call document AI")
    )
    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)

    response = client.post(
        "/chat/documents",
        files={"file": ("dish-master.xlsx", source.getvalue(), DOCUMENT_XLSX_MIME)},
        data={"message": DISH_EXPANSION_MESSAGE},
        headers=headers,
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": (
            "I couldn't find all five required dish category columns in the uploaded Excel "
            "file: Regular, Easy to chew, Soft and bite-sized, Minced and Moist, and Pureed."
        )
    }
    assert summarize.await_count == 0


def test_csv_filter_upload_returns_downloadable_xlsx_without_ai(client, monkeypatch):
    token = _login(client, "spreadsheet-filter-csv@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    source = b"ItemCode,ItemName,Category*,Uom\r\nBNS001,Coffee,Beverages,CTN\r\n"
    summarize = AsyncMock(side_effect=AssertionError("Filter operations must not call AI"))
    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)

    response = client.post(
        "/chat/documents",
        files={"file": ("items.csv", source, "text/csv")},
        data={"message": "Please add a filter only to the Category* column."},
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "skipped"
    assert body["attachment"]["filename"] == "items_filter_updated.xlsx"
    assert summarize.await_count == 0
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert revised.active.auto_filter.ref == "C1:C2"
        assert list(revised.active.values)[1] == ("BNS001", "Coffee", "Beverages", "CTN")
    finally:
        revised.close()


def test_filter_follow_up_targets_previous_generated_category_workbook(client, monkeypatch):
    token = _login(client, "spreadsheet-filter-chain@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Category*", "Price"])
    sheet.append(["Apple", "Fruit", 20])
    sheet.append(["Carrot", "Vegetable", 30])
    sheet.append(["Banana", "Fruit", 40])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(
        ChatDocumentService,
        "summarize",
        AsyncMock(side_effect=AssertionError("Spreadsheet operations must not call AI")),
    )
    split = client.post(
        "/chat/documents",
        files={"file": ("inventory.xlsx", source.getvalue(), DOCUMENT_XLSX_MIME)},
        data={"message": "Split the item list category-wise into one sheet per category."},
        headers=headers,
    )
    assert split.status_code == 200, split.text

    filtered = client.post(
        "/chat/documents/spreadsheet",
        json={
            "session_id": split.json()["session_id"],
            "instruction": "Please add a filter only to the Category* column.",
        },
        headers=headers,
    )

    assert filtered.status_code == 200, filtered.text
    body = filtered.json()
    assert body["attachment"]["filename"] == ("inventory_category_wise_filter_updated.xlsx")
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    revised = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert revised.sheetnames == ["Fruit", "Vegetable"]
        assert revised["Fruit"].auto_filter.ref == "B1:B3"
        assert revised["Vegetable"].auto_filter.ref == "B1:B2"
    finally:
        revised.close()


def test_filter_follow_up_targets_original_upload_when_no_generated_xlsx(client, monkeypatch):
    token = _login(client, "spreadsheet-filter-original-source@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Category*", "Price"])
    sheet.append(["Apple", "Fruit", 20])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    monkeypatch.setattr(
        ChatDocumentService,
        "summarize",
        AsyncMock(side_effect=AssertionError("Filter operations must not call AI")),
    )
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("original.xlsx", source.getvalue(), DOCUMENT_XLSX_MIME)},
        data={"analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text

    filtered = client.post(
        "/chat/documents/spreadsheet",
        json={
            "session_id": uploaded.json()["session_id"],
            "instruction": "Add an Excel filter to Category*.",
        },
        headers=headers,
    )

    assert filtered.status_code == 200, filtered.text
    body = filtered.json()
    assert body["attachment"]["filename"] == "original_filter_updated.xlsx"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    revised = load_workbook(BytesIO(download.content))
    try:
        assert revised.active.auto_filter.ref == "B1:B2"
    finally:
        revised.close()


def test_filter_follow_up_without_spreadsheet_returns_clean_error(client):
    token = _login(client, "spreadsheet-filter-missing-source@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    session_id = client.post(
        "/chat/sessions", json={"title": "No spreadsheet"}, headers=headers
    ).json()["id"]

    response = client.post(
        "/chat/documents/spreadsheet",
        json={
            "session_id": session_id,
            "instruction": "Add a filter to the Category* column.",
        },
        headers=headers,
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "Upload an Excel spreadsheet first, then ask me to add the filter."
    }


def test_filter_upload_missing_category_returns_clean_error_without_ai(client, monkeypatch):
    token = _login(client, "spreadsheet-filter-missing-column@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    summarize = AsyncMock(side_effect=AssertionError("Filter operations must not call AI"))
    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)

    response = client.post(
        "/chat/documents",
        files={"file": ("items.xlsx", _document_bytes("xlsx"), DOCUMENT_XLSX_MIME)},
        data={"message": "Please add a filter only to the Category* column."},
        headers=headers,
    )

    assert response.status_code == 422
    assert response.json() == {
        "detail": "I couldn't find a Category* column in the uploaded Excel file."
    }
    assert summarize.await_count == 0


def test_xlsx_category_split_keeps_original_attachment_bytes(client, monkeypatch):
    token = _login(client, "spreadsheet-original@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    source = _document_bytes("xlsx")
    workbook = load_workbook(BytesIO(source))
    sheet = workbook.active
    sheet.delete_rows(1, sheet.max_row)
    sheet.append(["Item", "Category"])
    sheet.append(["Apple", "Fruit"])
    revised_source = BytesIO()
    workbook.save(revised_source)
    workbook.close()
    original_bytes = revised_source.getvalue()

    monkeypatch.setattr(
        ChatDocumentService,
        "summarize",
        AsyncMock(side_effect=AssertionError("Category splitting must not call AI")),
    )
    response = client.post(
        "/chat/documents",
        files={"file": ("items.xlsx", original_bytes, DOCUMENT_XLSX_MIME)},
        data={
            "message": "Split the item list category-wise into a separate sheet for each category"
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    history = client.get(f"/chat/sessions/{response.json()['session_id']}", headers=headers).json()
    original_id = history["messages"][0]["document_attachment"]["id"]
    original = client.get(f"/chat/documents/{original_id}/download", headers=headers)
    assert original.content == original_bytes
    assert response.json()["attachment"]["id"] != original_id


def test_xlsx_category_split_returns_plain_missing_column_error(client):
    token = _login(client, "spreadsheet-missing-category@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    source = _document_bytes("xlsx")
    response = client.post(
        "/chat/documents",
        files={"file": ("items.xlsx", source, DOCUMENT_XLSX_MIME)},
        data={
            "message": "Split the item list category-wise into a separate sheet for each category"
        },
        headers=headers,
    )

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "Category column" in detail
    assert "{" not in detail
    assert '"error"' not in detail


def test_xlsx_category_split_returns_clean_error_when_generated_file_cannot_be_saved(
    client, monkeypatch
):
    from app.api import chat_documents

    token = _login(client, "spreadsheet-storage@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    workbook = Workbook()
    workbook.active.append(["Item", "Category"])
    workbook.active.append(["Apple", "Fruit"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    def fail_storage(*args, **kwargs):
        raise RuntimeError("database detail that must not reach the browser")

    monkeypatch.setattr(chat_documents.repository, "add_document_attachment", fail_storage)
    response = client.post(
        "/chat/documents",
        files={"file": ("items.xlsx", source.getvalue(), DOCUMENT_XLSX_MIME)},
        data={
            "message": "Split the item list category-wise into a separate sheet for each category"
        },
        headers=headers,
    )

    assert response.status_code == 500
    assert response.json() == {
        "detail": "The generated Excel workbook could not be saved. Please try again."
    }


def test_missing_ocr_is_not_reported_as_corrupt_document(client, monkeypatch):
    token = _login(client, "ocr-unavailable@example.com")

    def unavailable(*args):
        raise DocumentProcessingUnavailableError(
            "Scanned PDFs require Tesseract OCR on the backend."
        )

    monkeypatch.setattr(ChatDocumentService, "extract", unavailable)
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post(
        "/chat/documents",
        files={"file": ("scan.pdf", b"scan", "application/pdf")},
        headers=headers,
    )
    assert response.status_code == 503
    assert "Tesseract" in response.json()["detail"]
    assert client.get("/chat/sessions", headers=headers).json() == []


@pytest.mark.parametrize("output_format", ["pdf", "docx"])
def test_generates_document_from_instruction_in_new_chat(client, monkeypatch, output_format):
    token = _login(client, f"new-{output_format}@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    captured = []

    async def generate(self, instruction, summaries, history, profile=None):
        captured.append((instruction, summaries, history))
        return "# Project Summary\nFirst task: import a file.\nUse <label> & <value> safely."

    monkeypatch.setattr(ChatDocumentService, "generate_content", generate)
    response = client.post(
        "/chat/documents/generate",
        json={
            "session_id": None,
            "instruction": "  Create a project summary  ",
            "output_format": output_format,
        },
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert captured[0] == ("Create a project summary", [], [])
    body = response.json()
    downloaded = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    if output_format == "pdf":
        text = "".join(page.extract_text() for page in PdfReader(BytesIO(downloaded.content)).pages)
    else:
        text = "\n".join(p.text for p in Document(BytesIO(downloaded.content)).paragraphs)
    assert "Project Summary" in text
    assert "<label> & <value>" in text
    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    assert history["messages"][1]["document_attachment"]["kind"] == "generated"


def test_generate_rejects_empty_instruction_and_does_not_create_failed_sessions(
    client, monkeypatch
):
    token = _login(client, "failed-generation@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    invalid = client.post("/chat/documents/generate", json={"instruction": "   "}, headers=headers)
    assert invalid.status_code == 422

    def unavailable(*args):
        raise ChatModelUnavailableError("Model unavailable")

    monkeypatch.setattr(ChatDocumentService, "generate_content", AsyncMock(side_effect=unavailable))
    failed = client.post(
        "/chat/documents/generate", json={"instruction": "Create a summary"}, headers=headers
    )
    assert failed.status_code == 503
    assert client.get("/chat/sessions", headers=headers).json() == []


def test_document_routes_check_session_ownership_before_processing(client, monkeypatch):
    owner = _login(client, "session-owner@example.com")
    owner_headers = {"Authorization": f"Bearer {owner}"}
    session_id = client.post(
        "/chat/sessions", json={"title": "Private"}, headers=owner_headers
    ).json()["id"]
    other = _login(client, "session-other@example.com")
    headers = {"Authorization": f"Bearer {other}"}

    def unexpected(*args):
        raise AssertionError("Do not process another user's session")

    monkeypatch.setattr(ChatDocumentService, "extract", unexpected)
    monkeypatch.setattr(ChatDocumentService, "generate_content", unexpected)
    upload = client.post(
        "/chat/documents",
        files={"file": ("note.txt", b"Note", "text/plain")},
        data={"session_id": str(session_id)},
        headers=headers,
    )
    generate = client.post(
        "/chat/documents/generate",
        json={"session_id": session_id, "instruction": "Summarize"},
        headers=headers,
    )
    assert upload.status_code == generate.status_code == 404
