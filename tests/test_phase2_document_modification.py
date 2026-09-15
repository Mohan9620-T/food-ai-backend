from io import BytesIO

import pytest
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pypdf import PdfReader, PdfWriter

from app.services.document.document_intent_service import DocumentIntentService
from app.services.document.document_modification_service import DocumentModificationService
from app.services.document.document_operation_registry import (
    DocumentOperation,
    DocumentOperationRegistry,
    InputArity,
)
from app.services.document.document_pipeline_service import (
    DocumentPipelineService,
    PipelineDocument,
    StructuredPipelineStep,
)
from app.services.document.exceptions import (
    DocumentLocatorError,
    UnsupportedDocumentModificationError,
)
from app.services.document.extraction_models import (
    DocumentEdit,
    DocumentEditAction,
    DocumentLocator,
)
from app.services.powerpoint.pptx_modifier import PptxModifier
from app.services.spreadsheet.excel_modifier import ExcelModifier
from app.services.word.docx_modifier import DocxModifier

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _login(client, email: str) -> str:
    client.post(
        "/users/",
        json={"fullname": "Phase Two", "email": email, "password": "secret123"},
    )
    return client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]


def _docx_bytes() -> bytes:
    document = Document()
    heading = document.add_heading(level=1)
    first = heading.add_run("Old ")
    first.bold = True
    heading.add_run("Title")
    document.add_paragraph("Keep this paragraph exactly.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "10"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Code", "Status", "Name"])
    sheet.append(["C3", "Active", "Third"])
    sheet.append(["C1", "Inactive", "First"])
    sheet.append(["C2", "Active", "Second"])
    other = workbook.create_sheet("Untouched")
    other.append(["Keep", "Formula"])
    other.append([42, "=A2*2"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 2_000_000, 800_000)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "Quarterly "
    paragraph.add_run().text = "Draft"
    table_shape = slide.shapes.add_table(2, 2, 0, 900_000, 2_000_000, 800_000)
    table_shape.table.cell(0, 0).text = "Name"
    table_shape.table.cell(1, 0).text = "Old"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _pdf_bytes(page_count: int = 3) -> bytes:
    writer = PdfWriter()
    for number in range(1, page_count + 1):
        writer.add_blank_page(width=300 + number, height=400)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_docx_split_run_title_and_table_alignment_are_targeted_and_original_is_immutable():
    source = _docx_bytes()
    original = bytes(source)
    edits = (
        DocumentEdit(
            action=DocumentEditAction.REPLACE_TEXT,
            locator=DocumentLocator(heading="Old Title"),
            value="Final Title",
        ),
        DocumentEdit(
            action=DocumentEditAction.FORMAT_TABLE,
            locator=DocumentLocator(table_index=1),
            options={"alignment": "center"},
        ),
    )

    result = DocumentModificationService().modify(source, "report.docx", edits)

    assert source == original
    assert result.filename == "report_updated.docx"
    revised = Document(BytesIO(result.file_data))
    assert revised.paragraphs[0].text == "Final Title"
    assert revised.paragraphs[0].runs[0].bold is True
    assert revised.paragraphs[1].text == "Keep this paragraph exactly."
    assert revised.tables[0].alignment == WD_TABLE_ALIGNMENT.CENTER
    assert revised.tables[0].cell(1, 1).text == "10"


@pytest.mark.parametrize(
    ("paragraphs", "message"),
    [
        (("One",), "couldn't find"),
        (("Repeat", "Repeat"), "More than one"),
    ],
)
def test_docx_locator_requires_exactly_one_match(paragraphs, message):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    requested = "Missing" if paragraphs == ("One",) else "Repeat"

    with pytest.raises(DocumentLocatorError, match=message):
        DocumentModificationService().modify(
            output.getvalue(),
            "ambiguous.docx",
            (
                DocumentEdit(
                    action=DocumentEditAction.REPLACE_TEXT,
                    locator=DocumentLocator(text=requested),
                    value="New",
                ),
            ),
        )


def test_excel_filters_active_records_sorts_by_code_and_preserves_other_sheet():
    source = _xlsx_bytes()
    original = bytes(source)
    edits = (
        DocumentEdit(
            action=DocumentEditAction.FILTER_ROWS,
            locator=DocumentLocator(sheet="Records"),
            options={"equals": "active", "header": True},
        ),
        DocumentEdit(
            action=DocumentEditAction.SORT_RANGE,
            locator=DocumentLocator(sheet="Records"),
            options={"column": "Code", "header": True},
        ),
    )

    result = DocumentModificationService().modify(source, "records.xlsx", edits)

    assert source == original
    revised = load_workbook(BytesIO(result.file_data), data_only=False)
    try:
        assert list(revised["Records"].values) == [
            ("Code", "Status", "Name"),
            ("C2", "Active", "Second"),
            ("C3", "Active", "Third"),
        ]
        assert revised["Records"].auto_filter.ref == "A1:C3"
        assert revised["Untouched"]["A2"].value == 42
        assert revised["Untouched"]["B2"].value == "=A2*2"
    finally:
        revised.close()


def test_excel_cell_and_format_edits_reopen_with_requested_style():
    source = _xlsx_bytes()
    result = DocumentModificationService().modify(
        source,
        "records.xlsx",
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_CELL,
                locator=DocumentLocator(sheet="Records", cell="C2"),
                value="Updated",
            ),
            DocumentEdit(
                action=DocumentEditAction.FORMAT_RANGE,
                locator=DocumentLocator(sheet="Records", cell_range="A1:C1"),
                options={"bold": True, "alignment": "center", "border": True},
            ),
            DocumentEdit(
                action=DocumentEditAction.FREEZE_PANES,
                locator=DocumentLocator(sheet="Records", cell="A2"),
            ),
        ),
    )
    revised = load_workbook(BytesIO(result.file_data))
    try:
        assert revised["Records"]["C2"].value == "Updated"
        assert revised["Records"]["A1"].font.bold is True
        assert revised["Records"]["A1"].alignment.horizontal == "center"
        assert revised["Records"].freeze_panes == "A2"
    finally:
        revised.close()


def test_pptx_split_run_replacement_and_table_update_preserve_slide():
    source = _pptx_bytes()
    result = DocumentModificationService().modify(
        source,
        "deck.pptx",
        (
            DocumentEdit(
                action=DocumentEditAction.REPLACE_TEXT,
                locator=DocumentLocator(slide=1, text="Quarterly Draft"),
                value="Quarterly Final",
            ),
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(slide=1, table_index=1, row=2, column=1),
                value="New",
            ),
        ),
    )
    revised = Presentation(BytesIO(result.file_data))
    assert len(revised.slides) == 1
    assert any(
        getattr(shape, "has_text_frame", False) and "Quarterly Final" in shape.text
        for shape in revised.slides[0].shapes
    )
    table = next(shape.table for shape in revised.slides[0].shapes if shape.has_table)
    assert table.cell(1, 0).text == "New"


def test_pdf_page_operations_reopen_and_original_is_immutable():
    source = _pdf_bytes()
    original = bytes(source)
    result = DocumentModificationService().modify(
        source,
        "pages.pdf",
        (
            DocumentEdit(
                action=DocumentEditAction.REMOVE_PDF_PAGES,
                locator=DocumentLocator(pages=[2]),
            ),
        ),
    )
    assert source == original
    reader = PdfReader(BytesIO(result.file_data))
    assert len(reader.pages) == 2
    assert [float(page.mediabox.width) for page in reader.pages] == [301, 303]


@pytest.mark.parametrize(
    ("action", "pages", "expected_widths"),
    [
        (DocumentEditAction.EXTRACT_PDF_PAGES, [3, 1], [303, 301]),
        (DocumentEditAction.REORDER_PDF_PAGES, [2, 3, 1], [302, 303, 301]),
    ],
)
def test_pdf_extract_and_reorder_are_deterministic(action, pages, expected_widths):
    result = DocumentModificationService().modify(
        _pdf_bytes(),
        "pages.pdf",
        (DocumentEdit(action=action, locator=DocumentLocator(pages=pages)),),
    )
    reader = PdfReader(BytesIO(result.file_data))
    assert [float(page.mediabox.width) for page in reader.pages] == expected_widths


def test_pdf_merge_keeps_both_sources_immutable_and_reopens():
    first = _pdf_bytes(1)
    second = _pdf_bytes(2)
    snapshots = (bytes(first), bytes(second))

    result = DocumentModificationService().merge_pdfs(
        ((first, "first.pdf"), (second, "second.pdf"))
    )

    assert (first, second) == snapshots
    assert len(PdfReader(BytesIO(result.file_data)).pages) == 3


def test_arbitrary_pdf_visual_edit_returns_honest_limitation():
    with pytest.raises(UnsupportedDocumentModificationError, match="visual PDF editing"):
        DocumentModificationService().modify(
            _pdf_bytes(),
            "pages.pdf",
            (
                DocumentEdit(
                    action=DocumentEditAction.REPLACE_TEXT,
                    locator=DocumentLocator(text="Old"),
                    value="New",
                ),
            ),
        )


def test_phase2_operations_are_registered_with_correct_arity():
    registry = DocumentOperationRegistry()
    assert registry.get(DocumentOperation.MODIFY_DOCUMENT).input_arity == InputArity.SINGLE
    assert registry.get(DocumentOperation.MODIFY_PDF_PAGES).input_arity == InputArity.SINGLE
    assert registry.get(DocumentOperation.MERGE_PDF).input_arity == InputArity.MULTI


def test_clear_phase2_requests_produce_validated_targeted_plans_without_ai():
    service = DocumentIntentService()
    docx_plan = service.resolve_plan(
        'Change the title from "Old Title" to "Final Title" and center the table',
        input_file="report.docx",
    )
    excel_plan = service.resolve_plan(
        "Filter active records and sort them by Code", input_file="records.xlsx"
    )

    assert docx_plan.steps[0].operation == DocumentOperation.MODIFY_DOCUMENT
    assert len(docx_plan.steps[0].parameters["edits"]) == 2
    assert excel_plan.steps[0].operation == DocumentOperation.FILTER_AND_SORT_WORKBOOK
    assert len(excel_plan.steps[0].parameters["edits"]) == 2


def test_docx_upload_modification_is_stored_attached_and_downloadable(client):
    token = _login(client, "phase2-docx@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    source = _docx_bytes()

    response = client.post(
        "/chat/documents",
        files={"file": ("report.docx", source, DOCX_MIME)},
        data={
            "message": ('Change the title from "Old Title" to "Final Title" and center the table')
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "complete"
    assert body["attachment"]["kind"] == "generated"
    assert body["attachment"]["filename"] == "report_updated.docx"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    revised = Document(BytesIO(download.content))
    assert revised.paragraphs[0].text == "Final Title"
    assert revised.tables[0].alignment == WD_TABLE_ALIGNMENT.CENTER

    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    original_id = history["messages"][0]["document_attachment"]["id"]
    original = client.get(f"/chat/documents/{original_id}/download", headers=headers)
    assert original.content == source


def test_pipeline_executes_a_registered_targeted_modification():
    edit = DocumentEdit(
        action=DocumentEditAction.REPLACE_TEXT,
        locator=DocumentLocator(heading="Old Title"),
        value="Pipeline Title",
    )
    pipeline = DocumentPipelineService()
    steps = pipeline.plan_structured(
        (
            StructuredPipelineStep(
                document_type=None,
                operation=DocumentOperation.MODIFY_DOCUMENT,
                input_file="report.docx",
                output_type=None,
                parameters={"edits": [edit.model_dump(mode="json")]},
            ),
        ),
        input_filename="report.docx",
    )

    result = pipeline.execute_step(
        steps[0], PipelineDocument(file_data=_docx_bytes(), filename="report.docx")
    )

    revised = Document(BytesIO(result.document.file_data))
    assert revised.paragraphs[0].text == "Pipeline Title"


def test_docx_modifier_supports_paragraph_table_and_formatting_edits():
    source = _docx_bytes()
    result = DocxModifier().modify(
        source,
        (
            DocumentEdit(
                action=DocumentEditAction.ADD_PARAGRAPH,
                locator=DocumentLocator(paragraph_index=2),
                value="Inserted paragraph",
            ),
            DocumentEdit(
                action=DocumentEditAction.REMOVE_PARAGRAPH,
                locator=DocumentLocator(text="Inserted paragraph"),
            ),
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(table_index=1, row=2, column=2),
                value="25",
            ),
            DocumentEdit(
                action=DocumentEditAction.FORMAT_PARAGRAPH,
                locator=DocumentLocator(paragraph_index=2),
                options={
                    "alignment": "justify",
                    "space_before_pt": 4,
                    "space_after_pt": 6,
                    "line_spacing": 1.5,
                    "bold": True,
                    "italic": True,
                    "font_size_pt": 14,
                },
            ),
            DocumentEdit(
                action=DocumentEditAction.FORMAT_TABLE,
                locator=DocumentLocator(table_index=1),
                options={"alignment": "right", "vertical_alignment": "bottom"},
            ),
        ),
    )

    revised = Document(BytesIO(result))
    paragraph = revised.paragraphs[1]
    assert [item.text for item in revised.paragraphs] == [
        "Old Title",
        "Keep this paragraph exactly.",
    ]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[0].italic is True
    assert paragraph.runs[0].font.size.pt == 14
    assert revised.tables[0].cell(1, 1).text == "25"
    assert revised.tables[0].alignment == WD_TABLE_ALIGNMENT.RIGHT
    assert revised.tables[0].cell(0, 0).vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


@pytest.mark.parametrize(
    "edit, message",
    [
        (
            DocumentEdit(
                action=DocumentEditAction.REMOVE_PARAGRAPH,
                locator=DocumentLocator(paragraph_index=99),
            ),
            "does not exist",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(table_index=1, row=9, column=9),
                value="New",
            ),
            "does not contain cell",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.FORMAT_PARAGRAPH,
                locator=DocumentLocator(paragraph_index=1),
                options={"alignment": "diagonal"},
            ),
            "Paragraph alignment",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.FORMAT_TABLE,
                locator=DocumentLocator(table_index=1),
                options={"alignment": "diagonal"},
            ),
            "Table alignment",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.FORMAT_TABLE,
                locator=DocumentLocator(table_index=1),
                options={"vertical_alignment": "middle-ish"},
            ),
            "Cell alignment",
        ),
    ],
)
def test_docx_modifier_rejects_unsafe_or_invalid_locators(edit, message):
    with pytest.raises((DocumentLocatorError, UnsupportedDocumentModificationError), match=message):
        DocxModifier().modify(_docx_bytes(), (edit,))


def test_excel_modifier_supports_structural_and_extended_formatting_edits():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Code", "Status", "Name"])
    sheet.append(["C2", "Active", "Second"])
    sheet.append(["C1", "Active", "First"])
    sheet["A2"].fill = PatternFill("solid", fgColor="00FF00")
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    result = ExcelModifier().modify(
        source.getvalue(),
        (
            DocumentEdit(
                action=DocumentEditAction.SORT_RANGE,
                locator=DocumentLocator(sheet="data"),
                options={"column": "A", "header": True, "descending": True},
            ),
            DocumentEdit(
                action=DocumentEditAction.INSERT_ROWS,
                locator=DocumentLocator(sheet="Data", row=2),
                options={"amount": 1},
            ),
            DocumentEdit(
                action=DocumentEditAction.DELETE_ROWS,
                locator=DocumentLocator(sheet="Data", row=2),
            ),
            DocumentEdit(
                action=DocumentEditAction.INSERT_COLUMNS,
                locator=DocumentLocator(sheet="Data", column=2),
            ),
            DocumentEdit(
                action=DocumentEditAction.DELETE_COLUMNS,
                locator=DocumentLocator(sheet="Data", column=2),
            ),
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(sheet="Data", table_index=1, row=2, column=3),
                value="Changed",
            ),
            DocumentEdit(
                action=DocumentEditAction.FORMAT_RANGE,
                locator=DocumentLocator(sheet="Data", cell_range="A1:C1"),
                options={
                    "vertical_alignment": "center",
                    "wrap_text": True,
                    "fill": "#112233",
                    "column_width": 22,
                },
            ),
        ),
    )

    revised = load_workbook(BytesIO(result), data_only=False)
    try:
        assert list(revised["Data"].values) == [
            ("Code", "Status", "Name"),
            ("C2", "Active", "Changed"),
            ("C1", "Active", "First"),
        ]
        assert revised["Data"]["A1"].alignment.vertical == "center"
        assert revised["Data"]["A1"].alignment.wrap_text is True
        assert revised["Data"]["A1"].fill.fgColor.rgb == "00112233"
        assert revised["Data"].column_dimensions["A"].width == 22
    finally:
        revised.close()


@pytest.mark.parametrize(
    "edit, message",
    [
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_CELL,
                locator=DocumentLocator(sheet="Missing", cell="A1"),
                value="x",
            ),
            "worksheet named",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_CELL,
                locator=DocumentLocator(sheet="Records", cell="A1:B2"),
                value="x",
            ),
            "exactly one cell",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.INSERT_ROWS,
                locator=DocumentLocator(sheet="Records", row=2),
                options={"amount": 0},
            ),
            "positive integer",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.DELETE_ROWS,
                locator=DocumentLocator(sheet="Records", row=4),
                options={"amount": 2},
            ),
            "do not exist",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.DELETE_COLUMNS,
                locator=DocumentLocator(sheet="Records", column=3),
                options={"amount": 2},
            ),
            "do not exist",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.SORT_RANGE,
                locator=DocumentLocator(sheet="Records", cell_range="A1:Z9"),
                options={"column": "Code"},
            ),
            "extends beyond",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.FILTER_ROWS,
                locator=DocumentLocator(sheet="Records"),
                options={"equals": "unknown"},
            ),
            "couldn't find a column",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.SORT_RANGE,
                locator=DocumentLocator(sheet="Records"),
                options={"column": "Missing"},
            ),
            "column named",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.SORT_RANGE,
                locator=DocumentLocator(sheet="Records", cell_range="A1:B4"),
                options={"column": "C"},
            ),
            "outside the selected range",
        ),
    ],
)
def test_excel_modifier_rejects_invalid_or_ambiguous_targets(edit, message):
    with pytest.raises((DocumentLocatorError, UnsupportedDocumentModificationError), match=message):
        ExcelModifier().modify(_xlsx_bytes(), (edit,))


def test_excel_modifier_requires_sheet_name_and_rejects_merged_sort_ranges():
    with pytest.raises(DocumentLocatorError, match="multiple worksheets"):
        ExcelModifier().modify(
            _xlsx_bytes(),
            (
                DocumentEdit(
                    action=DocumentEditAction.UPDATE_CELL,
                    locator=DocumentLocator(cell="A1"),
                    value="x",
                ),
            ),
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Code", "Name"])
    sheet.append(["C1", "First"])
    sheet.merge_cells("A1:B1")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(UnsupportedDocumentModificationError, match="merged cells"):
        ExcelModifier().modify(
            output.getvalue(),
            (
                DocumentEdit(
                    action=DocumentEditAction.SORT_RANGE,
                    locator=DocumentLocator(cell_range="A1:B2"),
                    options={"column": "A"},
                ),
            ),
        )


def test_pptx_modifier_applies_basic_formatting_and_exact_shape_locator():
    source = _pptx_bytes()
    presentation = Presentation(BytesIO(source))
    shape_name = next(
        shape.name
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame and "Quarterly" in shape.text
    )
    result = PptxModifier().modify(
        source,
        (
            DocumentEdit(
                action=DocumentEditAction.FORMAT_PARAGRAPH,
                locator=DocumentLocator(slide=1, shape_name=shape_name, text="Quarterly Draft"),
                options={
                    "alignment": "center",
                    "bold": True,
                    "italic": True,
                    "font_size_pt": 18,
                },
            ),
        ),
    )
    revised = Presentation(BytesIO(result))
    paragraph = next(
        shape.text_frame.paragraphs[0]
        for shape in revised.slides[0].shapes
        if shape.name == shape_name
    )
    assert paragraph.alignment == PP_ALIGN.CENTER
    assert all(run.font.bold is True for run in paragraph.runs)
    assert all(run.font.italic is True for run in paragraph.runs)
    assert all(run.font.size.pt == 18 for run in paragraph.runs)


@pytest.mark.parametrize(
    "edit, message",
    [
        (
            DocumentEdit(
                action=DocumentEditAction.REPLACE_TEXT,
                locator=DocumentLocator(slide=2, text="Quarterly Draft"),
                value="New",
            ),
            "slide 2",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(slide=9, table_index=1, row=1, column=1),
                value="New",
            ),
            "Slide 9",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(slide=1, table_index=9, row=1, column=1),
                value="New",
            ),
            "Table 9",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.UPDATE_TABLE_CELL,
                locator=DocumentLocator(slide=1, table_index=1, row=9, column=1),
                value="New",
            ),
            "cell does not exist",
        ),
        (
            DocumentEdit(
                action=DocumentEditAction.FORMAT_PARAGRAPH,
                locator=DocumentLocator(slide=1, text="Quarterly Draft"),
                options={"alignment": "diagonal"},
            ),
            "Paragraph alignment",
        ),
    ],
)
def test_pptx_modifier_rejects_unsafe_or_invalid_locators(edit, message):
    with pytest.raises((DocumentLocatorError, UnsupportedDocumentModificationError), match=message):
        PptxModifier().modify(_pptx_bytes(), (edit,))
