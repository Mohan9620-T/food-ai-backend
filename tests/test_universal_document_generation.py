import csv
from io import BytesIO, StringIO

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.models.chat import ChatDocumentAttachment
from app.services.chat_document_service import ChatDocumentService
from app.services.document.document_generation_service import DocumentGenerationService
from app.services.document.document_operation_registry import DocumentType

CONTENT = """# Inventory Report

Item,Category,Quantity
Apple,Fruit,2
Carrot,Vegetable,3
"""

EXPECTED_TYPES = {
    "pdf": (".pdf", "application/pdf"),
    "docx": (
        ".docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    "xlsx": (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv": (".csv", "text/csv"),
    "pptx": (
        ".pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ),
    "txt": (".txt", "text/plain"),
    "markdown": (".md", "text/markdown"),
}


def _headers(client, email: str) -> dict[str, str]:
    registered = client.post(
        "/users/",
        json={"fullname": "Universal Document User", "email": email, "password": "secret123"},
    )
    assert registered.status_code in (200, 201)
    login = client.post("/users/login", json={"email": email, "password": "secret123"})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _reopen(file_data: bytes, output_format: str) -> None:
    if output_format == "pdf":
        assert len(PdfReader(BytesIO(file_data)).pages) >= 1
    elif output_format == "docx":
        assert Document(BytesIO(file_data)).paragraphs
    elif output_format == "xlsx":
        workbook = load_workbook(BytesIO(file_data), read_only=True)
        try:
            assert workbook.sheetnames == ["Document"]
        finally:
            workbook.close()
    elif output_format == "csv":
        assert list(csv.reader(StringIO(file_data.decode("utf-8-sig"))))
    elif output_format == "pptx":
        assert len(Presentation(BytesIO(file_data)).slides) >= 1
    else:
        assert file_data.decode("utf-8").strip()


@pytest.mark.parametrize("output_format", EXPECTED_TYPES)
def test_each_document_type_generates_with_valid_metadata_and_reopens(output_format):
    generated = DocumentGenerationService().generate(CONTENT, output_format)
    extension, content_type = EXPECTED_TYPES[output_format]

    assert generated.filename.endswith(extension)
    assert generated.content_type == content_type
    assert generated.document_type == DocumentType(output_format)
    _reopen(generated.file_data, output_format)


@pytest.mark.parametrize(
    ("requested_filename", "output_format", "expected"),
    [
        ("../../Quarterly report?.exe", "xlsx", "Quarterly-report.xlsx"),
        (r"..\..\CON.pdf", "pdf", "generated-con.pdf"),
        ("🔥🔥.txt", "txt", "generated-document.txt"),
        ("notes.markdown", "markdown", "notes.md"),
    ],
)
def test_generated_filename_is_safe_and_uses_requested_output_extension(
    requested_filename, output_format, expected
):
    generated = DocumentGenerationService().generate(
        CONTENT, output_format, requested_filename=requested_filename
    )

    assert generated.filename == expected
    assert "/" not in generated.filename
    assert "\\" not in generated.filename


@pytest.mark.parametrize("output_format", EXPECTED_TYPES)
def test_each_type_is_stored_attached_and_downloadable(client, db_session, output_format):
    headers = _headers(client, f"universal-{output_format}@example.com")
    extension, content_type = EXPECTED_TYPES[output_format]
    response = client.post(
        "/chat/documents/generate",
        json={
            "instruction": CONTENT,
            "output_format": output_format,
            "mode": "export",
            "filename": "../../Inventory Report.unsafe",
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["attachment"]["kind"] == "generated"
    assert body["attachment"]["filename"] == f"Inventory-Report{extension}"
    assert body["attachment"]["content_type"] == content_type
    attachment = db_session.get(ChatDocumentAttachment, body["attachment"]["id"])
    assert attachment is not None
    assert attachment.file_data
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    assert download.headers["content-type"].split(";", 1)[0] == content_type
    assert download.content == attachment.file_data
    _reopen(download.content, output_format)


def test_generation_creates_a_new_attachment_without_overwriting_original_upload(
    client, db_session
):
    headers = _headers(client, "original-preserved@example.com")
    original = b"Item,Category\nApple,Fruit\n"
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("inventory.csv", original, "text/csv")},
        data={"analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text
    source_id = uploaded.json()["attachment"]["id"]
    generated = client.post(
        "/chat/documents/generate",
        json={
            "session_id": uploaded.json()["session_id"],
            "source_document_id": source_id,
            "mode": "export",
            "output_format": "xlsx",
        },
        headers=headers,
    )

    assert generated.status_code == 200, generated.text
    assert generated.json()["attachment"]["id"] != source_id
    source = db_session.get(ChatDocumentAttachment, source_id)
    assert source is not None
    assert source.kind == "uploaded"
    assert source.filename == "inventory.csv"
    assert source.file_data == original


def test_existing_dish_category_excel_generation_regression_is_unchanged():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Id",
            "Code",
            "Label",
            "Regular",
            "Easy to chew",
            "Soft and bite-sized",
            "Minced and Moist",
            "Pureed",
            "Price",
        ]
    )
    sheet.append([1, "DISH-A", "Dish A", "x", None, "X", None, None, 12.5])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    original = source.getvalue()

    generated, filename, _ = ChatDocumentService().format_spreadsheet(
        original,
        "dish-master.xlsx",
        "Create one row for every category marked x.",
    )

    assert original == source.getvalue()
    assert filename == "dish-master_category_expanded.xlsx"
    revised = load_workbook(BytesIO(generated), read_only=True)
    try:
        rows = list(revised.active.iter_rows(min_row=2, values_only=True))
        assert [row[1] for row in rows] == ["DISH-A-7R", "DISH-A-6SB"]
        assert [row[3:8].index("x") for row in rows] == [0, 2]
    finally:
        revised.close()
