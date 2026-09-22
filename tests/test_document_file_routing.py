import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from openpyxl import load_workbook
from reportlab.pdfgen.canvas import Canvas

from app.api import chat as chat_api
from app.api import chat_documents
from app.config import settings
from app.models.chat import ChatDocumentAttachment, ChatSession
from app.schemas.chat import ChatHistoryMessage
from app.services.document.document_automation_service import (
    AvailableDocument,
    DocumentAutomationService,
)
from app.services.document.document_generation_service import DocumentGenerationService
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.extraction_models import StructuredDocumentContent


def test_dotted_title_keeps_full_document_filename():
    assert (
        DocumentGenerationService.safe_filename(
            "U.S. Presidents from 1900 to 2026", DocumentType.XLSX
        )
        == "U-S-Presidents-from-1900-to-2026.xlsx"
    )
    assert DocumentGenerationService.safe_filename("report.docx", DocumentType.PDF) == "report.pdf"


def test_legacy_chat_consolidation_preserves_uploaded_bytes(client, headers, db_session):
    first = client.post("/chat/sessions", headers=headers, json={"title": "First chat"}).json()
    second = client.post("/chat/sessions", headers=headers, json={"title": "Second chat"}).json()
    source = b"Project Cedar launches in October."
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"session_id": second["id"], "analyze": "false"},
        files={"file": ("notes.txt", source, "text/plain")},
    ).json()
    document_id = upload["attachment"]["id"]
    result = client.post("/chat/sessions/consolidate", headers=headers)
    assert result.status_code == 200
    document = db_session.get(ChatDocumentAttachment, document_id)
    assert document.session_id == first["id"]
    download = client.get(f"/chat/documents/{document_id}/download", headers=headers)
    assert download.status_code == 200 and download.content == source
    assert db_session.get(ChatSession, first["id"]).latest_document_id == document_id


def test_read_only_document_answer_does_not_require_a_build_click(client, headers, monkeypatch):
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("notes.txt", b"Project Cedar launches in October.", "text/plain")},
    ).json()
    analyze = AsyncMock(return_value="**Project Cedar** launches in October.")
    monkeypatch.setattr(type(chat_documents.pipeline_service.document_service), "analyze", analyze)
    result = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={"session_id": upload["session_id"], "instruction": "What is this document about?"},
    ).json()
    assert result["status"] == "done", result
    assert "Project Cedar" in result["response"]
    assert result["attachments"] == []
    analyze.assert_awaited_once()


def test_planning_deadline_returns_error_without_a_file(client, headers, monkeypatch):
    import time

    def slow_plan(*args, **kwargs):
        time.sleep(0.03)
        raise AssertionError("The deadline should expire first")

    sid = client.post("/chat/sessions", headers=headers, json={}).json()["id"]
    monkeypatch.setattr(chat_documents.automation_service, "plan", slow_plan)
    monkeypatch.setattr(settings, "DOCUMENT_AI_TIMEOUT_SECONDS", 0.005)
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={"session_id": sid, "instruction": "Create an Excel file", "confirm": True},
    )
    assert response.status_code == 503
    assert "planner took too long" in response.json()["detail"]


def test_source_free_clarification_keeps_source_free_plan():
    model = MagicMock()
    model.chat.return_value = json.dumps(
        {
            "status": "ready",
            "steps": [{"operation": "create_document", "output_type": "xlsx", "parameters": {}}],
        }
    )
    plan = DocumentAutomationService(chat_service=model).plan(
        "All planets",
        (AvailableDocument(1, "menu.xlsx", DocumentType.XLSX, is_latest=True),),
        conversation_history=[
            ChatHistoryMessage(role="user", content="Create an Excel file about planets"),
            ChatHistoryMessage(role="assistant", content="Which planets should I include?"),
        ],
    )
    assert plan.steps[0].source_filenames == ()
    assert "Available files: []" in model.chat.call_args.args[0]


@pytest.mark.parametrize("verb", ["create", "crate", "creat", "generate"])
def test_independent_excel_request_does_not_reuse_old_upload(client, headers, monkeypatch, verb):
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("old-menu.txt", b"Dish: Idli. Price: 30 INR.", "text/plain")},
    ).json()
    planner = MagicMock(
        return_value=json.dumps(
            {
                "status": "ready",
                "steps": [
                    {"operation": "create_document", "output_type": "xlsx", "parameters": {}}
                ],
            }
        )
    )
    monkeypatch.setattr(chat_documents.automation_service, "chat_service", MagicMock(chat=planner))
    generate = AsyncMock(
        return_value=StructuredDocumentContent.model_validate(
            {"title": "Planets", "tables": [{"headers": ["Planet"], "rows": [["Earth"]]}]}
        )
    )
    monkeypatch.setattr(
        type(chat_documents.pipeline_service.document_service), "generate_content", generate
    )
    payload = {
        "session_id": upload["session_id"],
        "instruction": f"List the planets and {verb} the excel file",
    }
    result = client.post("/chat/documents/automate", headers=headers, json=payload).json()
    assert result["status"] == "done", result
    assert result["steps"][0]["source_document_ids"] == []
    attachment = result["attachments"][0]
    assert attachment["provenance"] == "general_knowledge"
    assert attachment["source_document_ids"] == []
    assert not generate.call_args.kwargs["documents"]
    assert "Available files: []" in planner.call_args.args[0]
    download = client.get(f"/chat/documents/{attachment['id']}/download", headers=headers)
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content))
    assert any("Earth" in row for sheet in workbook for row in sheet.values)
    workbook.close()


@pytest.mark.parametrize("path", ["/chat/", "/chat/stream"])
def test_general_question_after_upload_does_not_receive_document_context(
    client, headers, monkeypatch, path
):
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("menu.txt", b"Dish: Idli.", "text/plain")},
    ).json()
    calls = []

    def answer(message, history, reference_history):
        calls.append(reference_history)
        return "**Yes.** George Washington was the first U.S. president."

    async def stream(message, history, reference_history):
        yield answer(message, history, reference_history)

    monkeypatch.setattr(type(chat_api.service), "chat", lambda self, *args, **kwargs: answer(*args))
    monkeypatch.setattr(
        type(chat_api.service), "stream_chat", lambda self, *args, **kwargs: stream(*args)
    )
    response = client.post(
        path,
        headers=headers,
        params={"session_id": upload["session_id"]},
        json={"message": "you know us presidents"},
    )
    assert response.status_code == 200
    assert "George Washington" in response.text
    assert calls == [[]]


def test_document_reference_still_supplies_uploaded_content(client, headers, monkeypatch):
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("menu.txt", b"Dish: Idli.", "text/plain")},
    ).json()
    answer = MagicMock(return_value="The menu contains Idli.")
    monkeypatch.setattr(type(chat_api.service), "chat", lambda self, *args, **kwargs: answer(*args))
    response = client.post(
        "/chat/",
        headers=headers,
        params={"session_id": upload["session_id"]},
        json={"message": "What is in this document?"},
    )
    assert response.status_code == 200
    assert "Dish: Idli" in answer.call_args.args[2][0].content


def pdf_bytes():
    output = BytesIO()
    canvas = Canvas(output)
    for text in ("First page: Project Cedar", "Last page: Delivery in October"):
        canvas.drawString(72, 720, text)
        canvas.showPage()
    canvas.save()
    return output.getvalue()


@pytest.fixture
def headers(client):
    credentials = {"email": "file-routing@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "File routing"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.parametrize(
    "instruction",
    [
        'can you add this content "My design focuses on text" and update the excel document',
        "Create a Word document from this PDF",
        "Create separate rows based on Regular, Easy to Chew, Soft & Bite, Minced & Moist and Pureed",
    ],
)
def test_save_only_pdf_does_not_attempt_spreadsheet_edits(
    client, headers, monkeypatch, db_session, instruction
):
    summarize = AsyncMock(side_effect=AssertionError("Save-only must not call AI"))
    modify = MagicMock(side_effect=AssertionError("Save-only must not modify a file"))
    monkeypatch.setattr(type(chat_documents.service), "summarize", summarize)
    monkeypatch.setattr(type(chat_documents.intent_service), "resolve_plan", modify)
    source = pdf_bytes()
    response = client.post(
        "/chat/documents",
        headers=headers,
        files={"file": ("source.pdf", source, "application/pdf")},
        data={"message": instruction, "analyze": "false"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["analysis_status"] == "skipped"
    attachment = db_session.query(ChatDocumentAttachment).one()
    assert attachment.kind == "uploaded" and attachment.file_data == source
    assert "First page: Project Cedar" in attachment.raw_text
    assert "Last page: Delivery in October" in attachment.raw_text
    summarize.assert_not_called()
    modify.assert_not_called()


@pytest.mark.parametrize(
    "request_text,uses_ai",
    [
        ("can you create the word document on this content", False),
        ("Create a Word summary of this PDF", True),
    ],
)
def test_uploaded_pdf_word_direct_creation_and_download(
    client, headers, monkeypatch, db_session, request_text, uses_ai
):
    source = pdf_bytes()
    upload = client.post(
        "/chat/documents",
        headers=headers,
        files={"file": ("source.pdf", source, "application/pdf")},
        data={"message": "Create the Word document on this content", "analyze": "false"},
    )
    assert upload.status_code == 200
    data = upload.json()
    content = StructuredDocumentContent(
        title="Project Cedar",
        paragraphs=["First page: Project Cedar", "Last page: Delivery in October"],
    )
    generate = (
        AsyncMock(return_value=content)
        if uses_ai
        else AsyncMock(side_effect=AssertionError("A Word copy must work without AI"))
    )
    monkeypatch.setattr(
        type(chat_documents.pipeline_service.document_service), "generate_content", generate
    )
    instruction = {
        "session_id": data["session_id"],
        "source_document_id": data["attachment"]["id"],
        "instruction": request_text,
    }
    built = client.post("/chat/documents/automate", headers=headers, json=instruction)
    assert built.status_code == 200, built.text
    assert built.json()["status"] == "done", built.text
    assert db_session.query(ChatDocumentAttachment).filter_by(kind="generated").count() == 1
    attachment = built.json()["attachments"][0]
    assert attachment["filename"].endswith(".docx")
    assert attachment["source_document_ids"] == [data["attachment"]["id"]]
    if uses_ai:
        supplied_sources = generate.call_args.kwargs["documents"]
        assert "First page: Project Cedar" in supplied_sources[0].text
        assert "Last page: Delivery in October" in supplied_sources[0].text
    else:
        generate.assert_not_awaited()
    download = client.get(f"/chat/documents/{attachment['id']}/download", headers=headers)
    assert download.status_code == 200
    document = Document(BytesIO(download.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "First page: Project Cedar" in text and "Last page: Delivery in October" in text


def test_pdf_add_content_with_excel_target_creates_a_new_file():
    model = MagicMock()
    planner = DocumentAutomationService(chat_service=model)
    plan = planner.plan(
        'can you add this content "My design focuses on providing text-based responses" '
        "and update the excel document",
        (AvailableDocument(1, "sample-1.pdf", DocumentType.PDF, is_latest=True),),
        explicit_source=True,
    )
    assert plan.status == "ready"
    assert plan.steps[0].intent.operation == DocumentOperation.CREATE_DOCUMENT
    assert plan.steps[0].intent.output_type == DocumentType.XLSX
    assert plan.steps[0].source_filenames == ("sample-1.pdf",)
    model.chat.assert_not_called()


def test_question_wording_does_not_turn_an_edit_into_an_answer():
    model = MagicMock()
    model.chat.return_value = (
        '{"status":"clarification_required",'
        '"clarifying_question":"Where should the new paragraph go?","steps":[]}'
    )
    plan = DocumentAutomationService(chat_service=model).plan(
        'Can you add "Project Cedar" to this document?',
        (AvailableDocument(1, "source.docx", DocumentType.DOCX, is_latest=True),),
    )
    assert plan.status == "clarification_required"
    assert plan.steps == ()
    model.chat.assert_called_once()
