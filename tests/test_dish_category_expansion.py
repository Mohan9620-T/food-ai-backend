from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Sequence

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table

from app.services.chat_document_service import (
    ChatDocumentService,
    InvalidDocumentError,
    SpreadsheetOperation,
)

EXPANSION_PROMPT = (
    "Create separate rows based on Regular, Easy to Chew, Soft & Bite, Minced & Moist and Pureed."
)
CATEGORY_HEADERS = [
    "Regular",
    "Easy to chew",
    "Soft and bite-sized",
    "Minced and Moist",
    "Pureed",
]


def _workbook_bytes(
    rows: Sequence[Sequence[object]],
    *,
    headers: Sequence[str] | None = None,
    sheet_name: str = "Sheet1",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(headers or ["Id", "Code", "Label", *CATEGORY_HEADERS, "Price", "Date"]))
    for row in rows:
        sheet.append(list(row))
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    if sheet.max_row >= 2:
        sheet.cell(row=2, column=sheet.max_column).number_format = "yyyy-mm-dd"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _base_row(
    categories: Sequence[object],
    *,
    code: str = "CSPL-THK-1-Staple-DN-CH-FMTM",
    label: str = "Fried Mee Tai Mak",
) -> list[object]:
    return [101, code, label, *categories, 12.5, date(2026, 9, 10)]


@pytest.mark.parametrize(
    "instruction",
    [
        "Transform the Excel and generate a new Excel file",
        EXPANSION_PROMPT,
        "Expand each dish row based on the five texture categories",
        "Create one row for every category marked x",
        "Split each dish into rows based on the five category columns",
    ],
)
def test_dish_category_expansion_intent_is_detected(instruction):
    assert ChatDocumentService.spreadsheet_operation(instruction) == (
        SpreadsheetOperation.EXPAND_DISH_BY_DIETARY_CATEGORY
    )


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        ("Regular", "regular"),
        ("REGULAR", "regular"),
        (" regular ", "regular"),
        ("Easy   to CHEW", "easy_to_chew"),
        ("Soft & Bite", "soft_and_bite_sized"),
        ("Soft and bite-sized", "soft_and_bite_sized"),
        ("Minced & Moist", "minced_and_moist"),
        ("Minced and Moist", "minced_and_moist"),
        ("PUREED", "pureed"),
    ],
)
def test_dish_category_header_matching_is_case_whitespace_and_punctuation_tolerant(
    header, expected
):
    assert ChatDocumentService._dish_category_header_key(header) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("x", True),
        ("X", True),
        ("  x  ", True),
        ("\tX\n", True),
        ("yes", False),
        ("true", False),
        ("1", False),
        (1, False),
        ("Y", False),
        ("", False),
        (None, False),
    ],
)
def test_only_trimmed_x_activates_dish_category(value, expected):
    assert ChatDocumentService._is_active_dish_category(value) is expected


def test_two_active_categories_generate_two_isolated_rows_and_preserve_other_fields():
    source = _workbook_bytes([_base_row(["x", "", " X ", None, "yes"])])
    original_snapshot = bytes(source)

    generated, filename, actions = ChatDocumentService().format_spreadsheet(
        source, "dish-master.xlsx", EXPANSION_PROMPT
    )

    assert source == original_snapshot
    assert filename == "dish-master_category_expanded.xlsx"
    assert "processed 1 original dish rows" in " ".join(actions)
    assert "generated 2 category rows" in " ".join(actions)
    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2
        assert [row[3:8] for row in rows] == [
            ("x", None, None, None, None),
            (None, None, "x", None, None),
        ]
        assert rows[0][1:3] == (
            "CSPL-THK-1-Staple-DN-CH-FMTM-7R",
            "Fried Mee Tai Mak (7R - Regular)",
        )
        assert rows[1][1:3] == (
            "CSPL-THK-1-Staple-DN-CH-FMTM-6SB",
            "Fried Mee Tai Mak (6SB - Chopped)",
        )
        assert rows[0][0] == rows[1][0] == 101
        assert rows[0][8:] == rows[1][8:] == (12.5, datetime(2026, 9, 10))
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("categories", "expected_rows"),
    [
        (["x", "x", "x", "x", "x"], 5),
        ([None, None, "x", None, None], 1),
    ],
)
def test_active_category_count_controls_generated_row_count(categories, expected_rows):
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes([_base_row(categories)]), "dish.xlsx", EXPANSION_PROMPT
    )

    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        assert len(rows) == expected_rows
        assert all(sum(value == "x" for value in row[3:8]) == 1 for row in rows)
    finally:
        workbook.close()


def test_zero_active_category_row_remains_completely_unchanged():
    source_row = _base_row(["yes", "true", "1", "Y", None])
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes([source_row]), "dish.xlsx", EXPANSION_PROMPT
    )

    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert list(workbook.active.values)[1] == tuple(source_row[:-1]) + (datetime(2026, 9, 10),)
    finally:
        workbook.close()


def test_all_code_and_label_suffixes_are_exact_and_in_fixed_category_order():
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes([_base_row(["x"] * 5)]), "dish.xlsx", EXPANSION_PROMPT
    )

    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        assert [row[1] for row in rows] == [
            "CSPL-THK-1-Staple-DN-CH-FMTM-7R",
            "CSPL-THK-1-Staple-DN-CH-FMTM-7EC",
            "CSPL-THK-1-Staple-DN-CH-FMTM-6SB",
            "CSPL-THK-1-Staple-DN-CH-FMTM-5MM",
            "CSPL-THK-1-Staple-DN-CH-FMTM-4Pu",
        ]
        assert [row[2] for row in rows] == [
            "Fried Mee Tai Mak (7R - Regular)",
            "Fried Mee Tai Mak (7EC - Easy)",
            "Fried Mee Tai Mak (6SB - Chopped)",
            "Fried Mee Tai Mak (5MM - Minced)",
            "Fried Mee Tai Mak (4Pu - Pureed)",
        ]
        assert not any(" - 7" in str(row[1]) for row in rows)
        assert not any(" - Regular - 7R" in str(row[2]) for row in rows)
        assert [row[3:8].index("x") for row in rows] == [0, 1, 2, 3, 4]
    finally:
        workbook.close()


def test_existing_suffix_and_old_sample_format_are_replaced_not_duplicated():
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes(
            [
                _base_row(
                    [None, "x", None, None, None],
                    code="CSPL-THK-1-Staple-DN-CH-FMTM - 7R",
                    label="Fried Mee Tai Mak - Regular - 7R",
                )
            ]
        ),
        "sample1.xlsx",
        EXPANSION_PROMPT,
    )

    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        row = list(workbook.active.values)[1]
        assert row[1] == "CSPL-THK-1-Staple-DN-CH-FMTM-7EC"
        assert row[2] == "Fried Mee Tai Mak (7EC - Easy)"
        assert "-7R-7EC" not in row[1]
    finally:
        workbook.close()


def test_non_category_values_styles_formulas_order_and_unrelated_sheet_are_preserved():
    headers = ["Id", "Code", "Label", *CATEGORY_HEADERS, "Formula", "Notes"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dish Master"
    sheet.append(headers)
    sheet.append([1, "DISH-A", "Dish A", "x", "x", None, None, None, "=A2*10", "keep"])
    sheet.append([2, "DISH-B", "Dish B", None, None, None, "x", None, "=A3*10", None])
    sheet["J2"].font = Font(italic=True, color="FF0000")
    sheet.add_table(Table(displayName="DishMasterTable", ref="A1:J3"))
    other = workbook.create_sheet("Lookup")
    other.append(["Key", "Value"])
    other.append(["A", 99])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    generated, _, _ = ChatDocumentService().format_spreadsheet(
        buffer.getvalue(),
        "dishes.xlsx",
        f"{EXPANSION_PROMPT} Use worksheet Dish Master.",
    )

    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised.sheetnames == ["Dish Master", "Lookup"]
        assert list(revised["Lookup"].values) == [("Key", "Value"), ("A", 99)]
        rows = list(revised["Dish Master"].iter_rows(min_row=2, values_only=True))
        assert [row[0] for row in rows] == [1, 1, 2]
        assert [row[8] for row in rows] == ["=A2*10", "=A2*10", "=A3*10"]
        assert [row[9] for row in rows] == ["keep", "keep", None]
        assert revised["Dish Master"]["J3"].font.italic is True
        assert revised["Dish Master"]["J3"].font.color.rgb == "00FF0000"
        assert revised["Dish Master"].tables["DishMasterTable"].ref == "A1:J4"
        assert list(revised["Dish Master"].values)[0] == tuple(headers)
    finally:
        revised.close()


def test_multiple_qualifying_sheets_require_explicit_sheet_name():
    workbook = Workbook()
    for index, title in enumerate(("First", "Second")):
        sheet = workbook.active if index == 0 else workbook.create_sheet(title)
        sheet.title = title
        sheet.append(["Code", "Label", *CATEGORY_HEADERS])
        sheet.append([f"D-{index}", f"Dish {index}", "x", None, None, None, None])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(InvalidDocumentError, match="More than one worksheet"):
        ChatDocumentService().format_spreadsheet(buffer.getvalue(), "multi.xlsx", EXPANSION_PROMPT)

    generated, _, _ = ChatDocumentService().format_spreadsheet(
        buffer.getvalue(), "multi.xlsx", f"{EXPANSION_PROMPT} Use worksheet Second."
    )
    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised["First"]["C2"].value == "x"
        assert revised["Second"]["A2"].value == "D-1-7R"
    finally:
        revised.close()


def test_header_can_follow_a_nonempty_title_row():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Dish master export"])
    sheet.append(["Code", "Label", *CATEGORY_HEADERS])
    sheet.append(["DISH-A", "Dish A", "x", None, None, None, None])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    generated, _, _ = ChatDocumentService().format_spreadsheet(
        source.getvalue(), "dish.xlsx", EXPANSION_PROMPT
    )

    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised.active["A1"].value == "Dish master export"
        assert revised.active["A3"].value == "DISH-A-7R"
    finally:
        revised.close()


def test_missing_category_column_and_no_rows_return_clean_errors():
    with pytest.raises(InvalidDocumentError, match="all five required dish category columns"):
        ChatDocumentService().format_spreadsheet(
            _workbook_bytes([_base_row(["x"] * 5)], headers=["Id", "Code", "Label"]),
            "missing.xlsx",
            EXPANSION_PROMPT,
        )
    with pytest.raises(InvalidDocumentError, match="no dish rows"):
        ChatDocumentService().format_spreadsheet(
            _workbook_bytes([]), "empty.xlsx", EXPANSION_PROMPT
        )


def _large_fixture_bytes() -> bytes:
    extra_headers = [f"Field {number}" for number in range(1, 71)]
    headers = ["Code", "Label", *extra_headers[:48], *CATEGORY_HEADERS, *extra_headers[48:]]
    assert len(headers) == 77
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    distributions = [(3, 211), (1, 637), (2, 38), (4, 35), (5, 330)]
    row_number = 0
    for active_count, count in distributions:
        for _ in range(count):
            row_number += 1
            code = "CSPL-THK-1-Staple-DN-CH-FMTM" if row_number == 1 else f"DISH-{row_number:04d}"
            label = "Fried Mee Tai Mak" if row_number == 1 else f"Dish {row_number}"
            row: list[object] = [
                code,
                label,
                *[f"value-{row_number}-{column}" for column in range(48)],
            ]
            row.extend(["x" if category < active_count else None for category in range(5)])
            row.extend([row_number + column for column in range(22)])
            assert len(row) == 77
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_full_scale_1251_row_77_column_fixture_generates_3136_rows():
    source = _large_fixture_bytes()
    original_snapshot = bytes(source)

    generated, filename, actions = ChatDocumentService().format_spreadsheet(
        source, "UAT-EMOS-Dish_Template_10092026.xlsx", EXPANSION_PROMPT
    )

    assert source == original_snapshot
    assert filename == "UAT-EMOS-Dish_Template_10092026_category_expanded.xlsx"
    assert "processed 1,251 original dish rows" in " ".join(actions)
    assert "generated 3,136 category rows" in " ".join(actions)
    workbook = load_workbook(BytesIO(generated), read_only=True, data_only=False)
    try:
        sheet = workbook["Sheet1"]
        assert sheet.max_row == 3137
        assert sheet.max_column == 77
        rows = list(sheet.iter_rows(min_row=2, max_row=4, values_only=True))
        assert [row[0] for row in rows] == [
            "CSPL-THK-1-Staple-DN-CH-FMTM-7R",
            "CSPL-THK-1-Staple-DN-CH-FMTM-7EC",
            "CSPL-THK-1-Staple-DN-CH-FMTM-6SB",
        ]
        assert [row[1] for row in rows] == [
            "Fried Mee Tai Mak (7R - Regular)",
            "Fried Mee Tai Mak (7EC - Easy)",
            "Fried Mee Tai Mak (6SB - Chopped)",
        ]
    finally:
        workbook.close()


def _real_fixture_path() -> Path | None:
    candidates = (
        Path("UAT-EMOS-Dish_Template_10092026.xlsx"),
        Path("tests/fixtures/UAT-EMOS-Dish_Template_10092026.xlsx"),
    )
    return next((candidate for candidate in candidates if candidate.is_file()), None)


REAL_UAT_FIXTURE = _real_fixture_path()


@pytest.mark.skipif(REAL_UAT_FIXTURE is None, reason="real UAT fixture was not attached")
def test_real_uat_fixture_expands_1251_rows_to_3136():
    assert REAL_UAT_FIXTURE is not None
    source = REAL_UAT_FIXTURE.read_bytes()
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        source, REAL_UAT_FIXTURE.name, EXPANSION_PROMPT
    )
    workbook = load_workbook(BytesIO(generated), read_only=True, data_only=False)
    try:
        assert workbook["Sheet1"].max_row - 1 == 3136
        assert workbook["Sheet1"].max_column == 77
    finally:
        workbook.close()
