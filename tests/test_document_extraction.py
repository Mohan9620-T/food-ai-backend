from io import BytesIO

import pytesseract
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.errors import DependencyError
from pypdfium2 import PdfBitmap
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from app.services.chat_document_service import (
    ChatDocumentService,
    DocumentProcessingUnavailableError,
    InvalidDocumentError,
)


def text_pdf(text: str = "Calories: 1800; Allergy: peanuts") -> bytes:
    buffer = BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(72, 720, text)
    canvas.save()
    return buffer.getvalue()


def scanned_pdf() -> bytes:
    buffer = BytesIO()
    with Image.new("RGB", (300, 120), "white") as image:
        ImageDraw.Draw(image).text((20, 30), "Calories: 1800", fill="black")
        canvas = Canvas(buffer, pagesize=(300, 120))
        canvas.drawImage(ImageReader(image), 0, 0, width=300, height=120)
        canvas.save()
    return buffer.getvalue()


def test_text_pdf_extracts_without_running_ocr(monkeypatch):
    def unexpected_ocr(*args, **kwargs):
        pytest.fail("A searchable PDF must not need an OCR installation")

    monkeypatch.setattr(pytesseract, "image_to_string", unexpected_ocr)
    assert "Allergy: peanuts" in ChatDocumentService().extract(text_pdf(), "report.PDF")


def test_scanned_pdf_renders_to_pillow_without_numpy(monkeypatch):
    images = []

    def unexpected_numpy(*args, **kwargs):
        pytest.fail("Rendering PDF pages must not require NumPy")

    def read_scan(image, *, timeout):
        assert isinstance(image, Image.Image)
        assert image.size == (600, 240)
        assert image.getpixel((0, 0)) == (255, 255, 255)
        assert timeout == ChatDocumentService.OCR_PAGE_TIMEOUT_SECONDS
        images.append(image)
        return "Calories: 1800\n"

    monkeypatch.setattr(PdfBitmap, "to_numpy", unexpected_numpy)
    monkeypatch.setattr(pytesseract, "image_to_string", read_scan)
    assert ChatDocumentService().extract(scanned_pdf(), "scan.pdf") == "Calories: 1800"
    assert len(images) == 1
    with pytest.raises(ValueError, match="closed"):
        images[0].getpixel((0, 0))


def test_mixed_pdf_keeps_page_order_and_only_runs_ocr_for_scanned_pages(monkeypatch):
    writer = PdfWriter()
    for source in (
        text_pdf("First section"),
        scanned_pdf(),
        text_pdf("Third section"),
        scanned_pdf(),
    ):
        writer.add_page(PdfReader(BytesIO(source)).pages[0])
    buffer = BytesIO()
    writer.write(buffer)
    ocr_calls = []

    def read_scan(image, *, timeout):
        assert image.size == (600, 240), "Searchable pages must not run OCR"
        ocr_calls.append(image.size)
        return "Second section" if len(ocr_calls) == 1 else "Fourth section"

    monkeypatch.setattr(pytesseract, "image_to_string", read_scan)
    extracted = ChatDocumentService().extract(buffer.getvalue(), "mixed.pdf")
    assert [line for line in extracted.splitlines() if line] == [
        "First section",
        "Second section",
        "Third section",
        "Fourth section",
    ]
    assert len(ocr_calls) == 2


def test_mixed_pdf_does_not_silently_drop_scanned_pages_when_ocr_is_missing(monkeypatch):
    writer = PdfWriter(clone_from=BytesIO(text_pdf()))
    writer.add_page(PdfReader(BytesIO(scanned_pdf())).pages[0])
    buffer = BytesIO()
    writer.write(buffer)

    def fail_ocr(*args, **kwargs):
        raise pytesseract.TesseractNotFoundError()

    monkeypatch.setattr(pytesseract, "image_to_string", fail_ocr)
    with pytest.raises(DocumentProcessingUnavailableError, match="OCR is unavailable"):
        ChatDocumentService().extract(buffer.getvalue(), "mixed.pdf")


def test_searchable_pdf_with_physically_empty_page_does_not_need_ocr(monkeypatch):
    writer = PdfWriter(clone_from=BytesIO(text_pdf()))
    writer.add_blank_page(width=300, height=120)
    buffer = BytesIO()
    writer.write(buffer)

    def unexpected_ocr(*args, **kwargs):
        pytest.fail("A physically empty page must not need OCR")

    monkeypatch.setattr(pytesseract, "image_to_string", unexpected_ocr)
    assert "Allergy: peanuts" in ChatDocumentService().extract(buffer.getvalue(), "report.pdf")


def test_searchable_pdf_with_vector_only_page_does_not_need_ocr(monkeypatch):
    vector_buffer = BytesIO()
    canvas = Canvas(vector_buffer, pagesize=(300, 120))
    canvas.rect(20, 20, 100, 50)
    canvas.save()
    writer = PdfWriter(clone_from=BytesIO(text_pdf()))
    writer.add_page(PdfReader(BytesIO(vector_buffer.getvalue())).pages[0])
    buffer = BytesIO()
    writer.write(buffer)

    def unexpected_ocr(*args, **kwargs):
        pytest.fail("A vector-only page must not need OCR")

    monkeypatch.setattr(pytesseract, "image_to_string", unexpected_ocr)
    assert "Allergy: peanuts" in ChatDocumentService().extract(buffer.getvalue(), "report.pdf")


@pytest.mark.parametrize(
    ("failure", "message"),
    [
        (pytesseract.TesseractNotFoundError(), "OCR is unavailable"),
        (pytesseract.TesseractError(1, "Language data unavailable"), "using OCR"),
        (RuntimeError("Tesseract process timeout"), "took too long"),
    ],
)
def test_scanned_pdf_reports_ocr_failures_without_calling_it_corrupt(monkeypatch, failure, message):
    def fail_ocr(*args, **kwargs):
        raise failure

    monkeypatch.setattr(pytesseract, "image_to_string", fail_ocr)
    with pytest.raises(DocumentProcessingUnavailableError, match=message):
        ChatDocumentService().extract(scanned_pdf(), "scan.pdf")


def test_scanned_pdf_without_readable_text_is_rejected(monkeypatch):
    monkeypatch.setattr(pytesseract, "image_to_string", lambda *args, **kwargs: " \n")
    with pytest.raises(InvalidDocumentError, match="does not contain readable text"):
        ChatDocumentService().extract(scanned_pdf(), "blank.pdf")


def test_password_protected_pdf_has_actionable_error():
    writer = PdfWriter(clone_from=BytesIO(text_pdf()))
    writer.encrypt("test-password")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(InvalidDocumentError, match="password-protected"):
        ChatDocumentService().extract(buffer.getvalue(), "protected.pdf")


def test_pdf_with_empty_user_password_can_be_read():
    writer = PdfWriter(clone_from=BytesIO(text_pdf()))
    writer.encrypt("", "owner-password")
    buffer = BytesIO()
    writer.write(buffer)
    assert "Calories: 1800" in ChatDocumentService().extract(buffer.getvalue(), "report.pdf")


def test_missing_pdf_encryption_dependency_is_not_reported_as_corruption(monkeypatch):
    def fail_reader(*args, **kwargs):
        raise DependencyError("Cryptography is required")

    monkeypatch.setattr("app.services.chat_document_service.PdfReader", fail_reader)
    with pytest.raises(DocumentProcessingUnavailableError, match="encryption"):
        ChatDocumentService().extract(text_pdf(), "report.pdf")


def test_corrupt_pdf_is_still_rejected():
    with pytest.raises(InvalidDocumentError, match="corrupt"):
        ChatDocumentService().extract(b"%PDF-1.7\nnot a real PDF", "bad.pdf")


def test_generated_pdf_treats_model_output_as_text():
    content = "# Summary\nReference: <link> and <b> unclosed\nCalories <1800 & protein >90"
    data, filename, content_type = ChatDocumentService().render(content, "pdf")
    assert filename.endswith(".pdf")
    assert content_type == "application/pdf"
    rendered_text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "Reference: <link> and <b> unclosed" in rendered_text
    assert "Calories <1800 & protein >90" in rendered_text


def test_docx_includes_paragraphs_and_table_cells():
    document = Document()
    document.add_heading("Project overview", level=1)
    document.add_paragraph("Milestone: launch")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Alex"
    table.cell(1, 1).text = "In progress"
    buffer = BytesIO()
    document.save(buffer)
    text = ChatDocumentService().extract(buffer.getvalue(), "project.docx")
    assert "Project overview\nMilestone: launch" in text
    assert "Owner\tStatus\nAlex\tIn progress" in text


def test_xlsx_includes_multiple_sheets_values_and_closes_archive(monkeypatch):
    workbook = Workbook()
    workbook.active.title = "Projects"
    workbook.active.append(["Project", "Budget", "Approved"])
    workbook.active.append(["Launch", 1250, True])
    workbook.create_sheet("Owners").append(["Alex", None, "Team A"])
    buffer = BytesIO()
    workbook.save(buffer)
    opened = []

    def track_workbook(*args, **kwargs):
        reader = load_workbook(*args, **kwargs)
        opened.append(reader)
        return reader

    monkeypatch.setattr("app.services.chat_document_service.load_workbook", track_workbook)
    text = ChatDocumentService().extract(buffer.getvalue(), "projects.xlsx")
    assert "[Projects]\nProject\tBudget\tApproved\nLaunch\t1250\tTrue" in text
    assert "[Owners]\nAlex\t\tTeam A" in text
    # One view preserves formulas and the second exposes cached values.
    assert len(opened) == 2
    assert opened[0]._archive.fp is None


@pytest.mark.parametrize("filename", ["empty.docx", "empty.xlsx"])
def test_empty_office_documents_are_rejected(filename):
    document = Document() if filename.endswith(".docx") else Workbook()
    buffer = BytesIO()
    document.save(buffer)
    with pytest.raises(InvalidDocumentError, match="does not contain readable text"):
        ChatDocumentService().extract(buffer.getvalue(), filename)


@pytest.mark.parametrize("filename", ["corrupt.docx", "corrupt.xlsx"])
def test_corrupt_office_documents_are_rejected(filename):
    with pytest.raises(InvalidDocumentError, match="corrupt"):
        ChatDocumentService().extract(b"not an Office document", filename)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
@pytest.mark.parametrize("extension", [".txt", ".csv"])
def test_text_and_csv_handle_unicode_bom_and_windows_exports(encoding, extension):
    text = 'Name,Note\r\nAm\u00e9lie,"Project, phase 2"'
    extracted = ChatDocumentService().extract(text.encode(encoding), f"notes{extension}")
    assert "Am\u00e9lie" in extracted
    assert "Project, phase 2" in extracted
    assert "\ufeff" not in extracted
    if extension == ".csv":
        assert "Name\tNote\nAm\u00e9lie\tProject, phase 2" == extracted


@pytest.mark.parametrize("extension", [".txt", ".csv"])
def test_binary_data_cannot_be_imported_as_text(extension):
    with pytest.raises(InvalidDocumentError, match="binary data"):
        ChatDocumentService().extract(b"MZ\x00\x01\x02", f"invalid{extension}")


@pytest.mark.parametrize("extension", [".txt", ".csv"])
def test_windows_1252_text_encoding_is_detected(extension):
    extracted = ChatDocumentService().extract(b"Name: Am\xe9lie", f"notes{extension}")
    assert "Am\xe9lie" in extracted


def test_unsupported_document_extension_is_rejected():
    with pytest.raises(InvalidDocumentError, match="Unsupported document type"):
        ChatDocumentService().extract(b"plain text", "notes.exe")
