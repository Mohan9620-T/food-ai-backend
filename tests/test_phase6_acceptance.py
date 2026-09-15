from io import BytesIO

from openpyxl import Workbook, load_workbook
from reportlab.pdfgen.canvas import Canvas

from app.services.chat_document_service import ChatDocumentService


def _login(client, email: str) -> dict[str, str]:
    client.post(
        "/users/",
        json={"fullname": "Acceptance User", "email": email, "password": "secret123"},
    )
    token = client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


def _pdf_bytes() -> bytes:
    output = BytesIO()
    canvas = Canvas(output)
    canvas.drawString(72, 760, "Customer: Atlas Foods")
    canvas.drawString(72, 740, "Status: Active")
    canvas.save()
    return output.getvalue()


def _records_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Code", "Status", "Customer"])
    sheet.append(["C003", "Active", "Gamma"])
    sheet.append(["A001", "Inactive", "Alpha"])
    sheet.append(["B002", "Active", "Beta"])
    notes = workbook.create_sheet("Notes")
    notes.append(["Keep", "Unchanged"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_acceptance_01_upload_pdf_read_and_summarize(client, monkeypatch):
    headers = _login(client, "acceptance-01@example.com")
    source = _pdf_bytes()

    async def summarize(self, extracted, instruction):
        del self
        assert "Atlas Foods" in extracted.text
        assert "summarize" in instruction.casefold()
        return "Atlas Foods is an active customer."

    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)
    response = client.post(
        "/chat/documents",
        headers=headers,
        data={"message": "Read this document and summarize it."},
        files={"file": ("customer.pdf", source, "application/pdf")},
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "complete"
    assert body["response"] == "Atlas Foods is an active customer."
    assert body["fidelity"] == "HIGH"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    assert download.content == source


def test_acceptance_07_filter_active_records_and_sort_by_code(client):
    headers = _login(client, "acceptance-07@example.com")
    source = _records_workbook()
    response = client.post(
        "/chat/documents",
        headers=headers,
        data={"message": "Filter the active records and sort them by Code"},
        files={
            "file": (
                "records.xlsx",
                source,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["attachment"]["kind"] == "generated"
    assert body["fidelity"] == "FULL"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    workbook = load_workbook(BytesIO(download.content), data_only=False)
    try:
        assert list(workbook["Records"].values) == [
            ("Code", "Status", "Customer"),
            ("B002", "Active", "Beta"),
            ("C003", "Active", "Gamma"),
        ]
        assert list(workbook["Notes"].values) == [("Keep", "Unchanged")]
    finally:
        workbook.close()

    history = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()
    original_id = history["messages"][0]["document_attachment"]["id"]
    original = client.get(f"/chat/documents/{original_id}/download", headers=headers)
    assert original.content == source
