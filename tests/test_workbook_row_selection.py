import hashlib
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.api import chat, chat_documents
from app.services.document.document_automation_service import (
    AvailableDocument,
    DocumentAutomationService,
)
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.exceptions import DocumentLocatorError
from app.services.spreadsheet.row_selection import RowSelection, WorkbookRowSelector

LOOKUP_REQUEST = "read the excel document end to end and give me the data's on the particular \r\n24908\r\n24697\r\n24702\r\n19534"
FOLLOWUP_REQUEST = "give me the this 24908\n24697\n24702\n19534 id's row data's only"

REQUEST = """I will upload an Excel file.
From the uploaded Excel file, find the rows containing these three IDs:
- 24138
- 24101
- 24102
Return only the data related to these three IDs.
Requirements:
1. Search all sheets in the uploaded Excel file.
2. Identify the column that contains the ID values.
3. Find all rows where the ID exactly matches 24138, 24101, or 24102.
4. Keep the complete row data for those matching IDs, including all columns.
5. Do not include rows belonging to any other IDs.
6. Preserve the original column names and column order.
7. Preserve the original row data without modifying or formatting the values unnecessarily.
8. If an ID appears multiple times, include all matching rows.
9. Do not modify the original uploaded Excel file.
10. Create a new Excel file containing only the matching records.
11. If one or more requested IDs are not found, clearly mention which IDs were not found.
12. Return the generated Excel file as an attachment.
Expected result:
The output Excel should contain only records for:
24138
24101
24102"""


def workbook_bytes():
    book = Workbook()
    first = book.active
    first.title = "Items"
    first.append(["Id", "Name", "Empty", "Value", "Code"])
    first.append(["24138", "One", None, 0, "001"])
    first.append(["24101", "Two", None, False, "002"])
    first.append(["24138", "Duplicate", None, 3.5, "003"])
    first.append(["241380", "Prefix does not match", None, 8, "004"])
    first.append(["024138", "Leading zero is a different ID", None, 8, "005"])
    first.append(["999", "24102 in a non-ID column", None, 24102, "006"])
    first["A1"].font = Font(bold=True)
    first["E2"].number_format = "@"
    first.column_dimensions["B"].width = 35
    second = book.create_sheet("More")
    second.append(["Metadata"])
    second.append(["Other", "Dish ID", "Blank"])
    second.append(["Numeric ID", 24102, None])
    second.append(["Repeated numeric ID", 24138, None])
    book.create_sheet("Notes").append(["Instructions only"])
    buf = BytesIO()
    book.save(buf)
    book.close()
    return buf.getvalue()


def test_numbered_prompt_parses_only_the_requested_ids_and_bypasses_ai():
    assert RowSelection.from_instruction(REQUEST).ids == ["24138", "24101", "24102"]
    model = MagicMock()
    plan = DocumentAutomationService(chat_service=model).plan(
        REQUEST, (AvailableDocument(1, "source.xlsx", DocumentType.XLSX, True),)
    )
    assert plan.steps[0].intent.operation == DocumentOperation.EXTRACT_MATCHING_ROWS
    assert plan.steps[0].intent.parameters == {"ids": ["24138", "24101", "24102"]}
    assert plan.steps[0].source_filenames == ("source.xlsx",)
    model.chat.assert_not_called()


@pytest.mark.parametrize(
    "instruction",
    [
        LOOKUP_REQUEST,
        FOLLOWUP_REQUEST,
        "Get data for 24908, 24697, 24702 and 19534",
        "Show IDs 24908, 24697, 24702, 19534",
    ],
)
def test_natural_lookup_wording_and_numbers_before_ids(instruction):
    selection = RowSelection.from_instruction(instruction, workbook_context=True)
    assert selection.ids == ["24908", "24697", "24702", "19534"]


@pytest.mark.parametrize(
    "instruction",
    [
        "Show the data for 2020 to 2026",
        "Give me a recipe for 4 people",
        "Read this Excel end to end and give me the data",
        "How many rows are in this Excel?",
        "Show rows except IDs 24908, 24697",
        "Give me a summary of this file in 3 paragraphs",
    ],
)
def test_lookup_does_not_guess_ids_from_counts_dates_or_unspecified_data(instruction):
    assert RowSelection.from_instruction(instruction, workbook_context=True) is None


def test_short_new_ids_override_old_creation_prompt_and_select_original():
    from app.schemas.chat import ChatHistoryMessage

    model = MagicMock()
    plan = DocumentAutomationService(chat_service=model).plan(
        "Create Excel data for 24908",
        (
            AvailableDocument(1, "original.xlsx", DocumentType.XLSX),
            AvailableDocument(2, "filtered.xlsx", DocumentType.XLSX, True, "generated"),
        ),
        conversation_history=[ChatHistoryMessage(role="user", content=REQUEST)],
    )
    assert plan.steps[0].intent.parameters == {"ids": ["24908"]}
    assert plan.steps[0].source_filenames == ("original.xlsx",)
    model.chat.assert_not_called()


def test_upload_lookup_followup_chat_and_stream_without_ai(client, monkeypatch):
    def no_model(*args, **kwargs):
        raise AssertionError("Workbook lookup must not use AI")

    monkeypatch.setattr(type(chat_documents.service), "summarize", no_model)
    monkeypatch.setattr(type(chat.service), "chat", no_model)
    monkeypatch.setattr(type(chat.service), "stream_chat", no_model)
    credentials = {"email": "lookup@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Lookup test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    book = Workbook()
    book.active.append(["Id", "Name", "Empty", "Value"])
    for identifier in [24908, 24697, 24702, 19534, 12345]:
        book.active.append([identifier, f"Item {identifier}", None, 0])
    raw = BytesIO()
    book.save(raw)
    book.close()
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"message": LOOKUP_REQUEST},
        files={"file": ("lookup.xlsx", raw.getvalue())},
    )
    assert upload.status_code == 200, upload.text
    data = upload.json()
    assert data["analysis_status"] == "complete"
    assert "Matched 4 rows" in data["response"]
    for identifier in [24908, 24697, 24702, 19534]:
        assert f"Item {identifier}" in data["response"]
    assert "Item 12345" not in data["response"]
    session = data["session_id"]
    output = client.get(f"/chat/documents/{data['attachment']['id']}/download", headers=headers)
    reopened = load_workbook(BytesIO(output.content))
    assert reopened.active.max_row == 5
    assert reopened.active["C2"].value is None and reopened.active["D2"].value == 0
    reopened.close()
    # This ID exists only in the upload, not the latest generated selection.
    followup = client.post(
        "/chat/documents/spreadsheet",
        headers=headers,
        json={"session_id": session, "instruction": "give me this 12345 id's row data's only"},
    )
    assert followup.status_code == 200, followup.text
    assert "Item 12345" in followup.json()["response"]
    assert (
        followup.json()["attachment"]["source_document_ids"]
        == data["attachment"]["source_document_ids"]
    )
    for endpoint in ["/chat/", "/chat/stream"]:
        response = client.post(
            endpoint,
            params={"session_id": session},
            headers=headers,
            json={"message": FOLLOWUP_REQUEST},
        )
        assert response.status_code == 200, response.text
        for identifier in [24908, 24697, 24702, 19534]:
            assert f"Item {identifier}" in response.text
        assert "unavailable" not in response.text and "interrupted" not in response.text
    missing = client.post(
        "/chat/documents/spreadsheet",
        headers=headers,
        json={"session_id": session, "instruction": "Get data for 99999"},
    )
    assert "IDs not found: 99999" in missing.json()["response"]
    original_id = data["attachment"]["source_document_ids"][0]
    assert (
        client.get(f"/chat/documents/{original_id}/download", headers=headers).content
        == raw.getvalue()
    )
    saved = client.get(f"/chat/documents/{data['attachment']['id']}/data", headers=headers)
    assert saved.json()["extracted_data"]["tables"]


@pytest.mark.parametrize(
    "instruction",
    [
        "Create a new Excel budget with 12 rows and 3 columns",
        "Create an Excel with records about ideas",
        "Filter Excel rows to exclude IDs 24138, 24101",
        "Find Excel rows with IDs 24138, 24101 and create a Word report",
        "Filter Excel rows with IDs 24138, 24101 and sort by Price",
    ],
)
def test_unrelated_or_exclusion_requests_are_not_inclusion_filters(instruction):
    assert RowSelection.from_instruction(instruction) is None


def test_all_sheets_duplicate_rows_values_types_and_formatting_are_preserved():
    source = workbook_bytes()
    digest = hashlib.sha256(source).hexdigest()
    result = WorkbookRowSelector().select(
        source, RowSelection(ids=["24138", "24101", "24102", "missing"])
    )
    assert result.counts == {"Items": 3, "More": 2}
    assert result.missing_ids == ("missing",)
    assert "Matched 5 rows" in result.summary and "IDs not found: missing" in result.summary
    original = load_workbook(BytesIO(source))
    selected = load_workbook(BytesIO(result.file_data))
    try:
        assert selected.sheetnames == ["Items", "More"]
        assert list(selected["Items"].values) == list(original["Items"].values)[:4]
        assert list(selected["More"].values) == list(original["More"].values)[1:]
        assert selected["Items"]["A1"].font.bold
        assert selected["Items"]["E2"].number_format == "@"
        assert selected["Items"].column_dimensions["B"].width == 35
        assert selected["Items"]["D3"].data_type == "b"
        assert selected["Items"]["C2"].value is None
        assert original["Items"].max_row == 7
    finally:
        original.close()
        selected.close()
    assert hashlib.sha256(source).hexdigest() == digest


def test_no_match_reports_missing_ids_and_writes_only_headers():
    result = WorkbookRowSelector().select(workbook_bytes(), RowSelection(ids=["99999"]))
    assert result.missing_ids == ("99999",)
    assert "Matched 0 rows" in result.summary
    selected = load_workbook(BytesIO(result.file_data))
    try:
        assert all(sheet.max_row == 1 for sheet in selected)
    finally:
        selected.close()


def test_ambiguous_columns_require_an_explicit_column():
    book = Workbook()
    book.active.append(["Item ID", "Order ID"])
    book.active.append(["24138", "24138"])
    buf = BytesIO()
    book.save(buf)
    with pytest.raises(DocumentLocatorError, match="multiple ID columns"):
        WorkbookRowSelector().select(buf.getvalue(), RowSelection(ids=["24138"]))
    result = WorkbookRowSelector().select(
        buf.getvalue(), RowSelection(ids=["24138"], column="Item ID")
    )
    assert sum(result.counts.values()) == 1


def test_filter_refuses_to_silently_drop_uncalculated_formula_data():
    book = Workbook()
    book.active.append(["Id", "Amount"])
    book.active.append(["24138", "=1+2"])
    buf = BytesIO()
    book.save(buf)
    with pytest.raises(DocumentLocatorError, match="no saved formula result"):
        WorkbookRowSelector().select(buf.getvalue(), RowSelection(ids=["24138"]))


def test_description_mode_cannot_invent_rows_from_an_unselected_source():
    plan = DocumentAutomationService().plan(
        REQUEST,
        (AvailableDocument(1, "source.xlsx", DocumentType.XLSX, True),),
        source_mode="description",
    )
    assert plan.status == "clarification_required"
    assert "upload" in plan.clarifying_question.lower()


def test_direct_creation_download_and_history_use_real_source_cells(client, monkeypatch):
    def no_model(*args, **kwargs):
        raise AssertionError("Exact row selection must not call AI")

    monkeypatch.setattr(chat_documents.automation_service, "chat_service", MagicMock(chat=no_model))
    monkeypatch.setattr(
        type(chat_documents.pipeline_service.document_service), "generate_content", no_model
    )
    credentials = {"email": "exact-rows@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Row selection test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    raw = workbook_bytes()
    upload = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("source.xlsx", raw)},
    ).json()
    payload = {
        "session_id": upload["session_id"],
        "source_document_id": upload["attachment"]["id"],
        "instruction": REQUEST,
    }
    response = client.post("/chat/documents/automate", headers=headers, json=payload)
    assert response.status_code == 200
    built = response.json()
    assert built["status"] == "done", built
    assert "Matched 5 rows" in built["response"]
    document = built["attachments"][0]
    assert document["source_document_ids"] == [upload["attachment"]["id"]]
    output = client.get(f"/chat/documents/{document['id']}/download", headers=headers)
    book = load_workbook(BytesIO(output.content))
    assert book["Items"].max_row == 4 and book["More"].max_row == 3
    book.close()
    saved = client.get(f"/chat/documents/{document['id']}/data", headers=headers).json()
    assert saved["extracted_data"]["tables"]
    assert (
        client.get(
            f"/chat/documents/{upload['attachment']['id']}/download", headers=headers
        ).content
        == raw
    )
