import hashlib
import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from app.api import chat_documents
from app.models.chat import ChatDocumentAttachment, ChatMessageRecord
from app.services.document.extraction_models import StructuredDocumentContent


def ready():
    return {
        "status": "ready",
        "steps": [{"operation": "create_document", "output_type": "xlsx", "parameters": {}}],
    }


def question(text="Which region?", options=True):
    payload = {"status": "clarification_required", "clarifying_question": text, "steps": []}
    if options:
        payload["clarification_options"] = [
            {"id": str(i), "label": label, "recommended": i == 1}
            for i, label in enumerate(["South India", "North India", "All India"], 1)
        ]
    return payload


@pytest.fixture
def build_setup(client, monkeypatch):
    credentials = {"email": "stateless-wizard@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Review test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    headers = {"Authorization": "Bearer " + token}
    sid = client.post("/chat/sessions", json={}, headers=headers).json()["id"]
    model = MagicMock(return_value=json.dumps(ready()))
    # Replace this dependency, avoiding methods shadowed on shared service
    # instances by older execution tests.
    monkeypatch.setattr(chat_documents.automation_service, "chat_service", MagicMock(chat=model))
    content = StructuredDocumentContent.model_validate(
        {
            "title": "Dishes",
            "assumptions": ["Prices are approximate INR amounts."],
            "tables": [{"headers": ["Dish", "Price"], "rows": [["Idli", "30"], ["Dosa", "50"]]}],
        }
    )
    generate = AsyncMock(return_value=content)
    monkeypatch.setattr(
        type(chat_documents.pipeline_service.document_service), "generate_content", generate
    )
    return headers, sid, generate, model


def call(client, headers, sid, instruction="Create an Excel file with dishes", **kwargs):
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={"session_id": sid, "instruction": instruction, **kwargs},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_question_answer_generates_once_and_saves_choices(client, db_session, build_setup):
    headers, sid, generate, model = build_setup
    model.side_effect = [
        json.dumps(question()),
        json.dumps(ready()),
    ]
    first = call(client, headers, sid)
    assert first["clarification"]["options"][0]["recommended"] is True
    result = call(client, headers, sid, "South India")
    assert result["status"] == "done"
    assert result["plan_summary"] is None
    db_session.expire_all()
    saved_review = client.get(f"/chat/sessions/{sid}", headers=headers).json()["messages"][-1][
        "automation"
    ]
    assert saved_review["response"] == result
    assert saved_review["request_instruction"] == "Create an Excel file with dishes"
    assert saved_review["instruction"] == "South India"
    assert saved_review["choices"] == [{"question": "Which region?", "answer": "South India"}]
    assert result["status"] == "done"
    assert len(result["attachments"]) == 1
    assert db_session.query(ChatDocumentAttachment).filter_by(kind="generated").count() == 1
    generate.assert_awaited_once()
    assert "South India" in model.call_args.args[0]
    attachment = result["attachments"][0]
    raw = client.get(f"/chat/documents/{attachment['id']}/download", headers=headers).content
    workbook = load_workbook(BytesIO(raw))
    for sheet in workbook:
        page = client.get(
            f"/chat/documents/{attachment['id']}/preview",
            params={"sheet": sheet.title},
            headers=headers,
        ).json()
        assert page["rows"] == [list(row) for row in sheet.values]
        assert page["sha256"] == hashlib.sha256(raw).hexdigest()
    assert attachment["provenance"] == "general_knowledge"


def test_new_request_does_not_inherit_an_unfinished_wizard(client, build_setup):
    headers, sid, generate, model = build_setup
    model.side_effect = [json.dumps(question()), json.dumps(ready())]
    call(client, headers, sid)
    instruction = "Create a new Excel with an inventory table"
    assert call(client, headers, sid, instruction)["status"] == "done"
    saved = client.get(f"/chat/sessions/{sid}", headers=headers).json()["messages"][-1][
        "automation"
    ]
    assert saved["request_instruction"] == instruction
    assert saved["choices"] == []
    generate.assert_awaited_once()


def test_skip_upload_generates_directly_and_saves_source_mode(client, build_setup):
    headers, sid, generate, model = build_setup
    instruction = "Create a new Excel from this image with Item and Quantity columns"
    assert call(client, headers, sid, instruction)["status"] == "clarification_required"
    result = call(client, headers, sid, instruction, source_mode="description")
    assert result["status"] == "done"
    model.assert_not_called()
    saved = client.get(f"/chat/sessions/{sid}", headers=headers).json()["messages"][-1][
        "automation"
    ]
    assert saved["source_mode"] == "description"
    assert saved["source_document_id"] is None
    assert saved["choices"][-1]["answer"] == "Create from written requirements without uploading."
    assert result["status"] == "done"
    assert result["attachments"][0]["source_document_ids"] == []
    generate.assert_awaited_once()
    assert generate.call_args.kwargs["documents"] == []


def test_legacy_confirm_cannot_skip_remaining_questions(client, build_setup):
    headers, sid, generate, model = build_setup
    model.side_effect = [
        json.dumps(question()),
        json.dumps(question("Which currency?")),
        json.dumps(ready()),
    ]
    assert call(client, headers, sid)["status"] == "clarification_required"
    assert (
        call(client, headers, sid, "South India", confirm=True)["status"]
        == "clarification_required"
    )
    generate.assert_not_awaited()
    assert call(client, headers, sid, "INR")["status"] == "done"
    generate.assert_awaited_once()


def test_plain_question_stays_compatible(client, build_setup):
    headers, sid, generate, model = build_setup
    model.return_value = json.dumps(question("What should I title this file?", False))
    result = call(client, headers, sid)
    assert result["response"] == "What should I title this file?"
    assert result["clarification"] == {
        "question": result["response"],
        "options": [],
        "allow_other": True,
    }
    generate.assert_not_awaited()


def test_history_endpoint_restores_choices_for_a_legacy_layout_question(
    client, db_session, build_setup
):
    headers, sid, _, _ = build_setup
    text = 'What text should appear under the column "Install"?'
    record = ChatMessageRecord(
        session_id=sid,
        sender="bot",
        content=text,
        automation={
            "instruction": "yes",
            "request_instruction": "Header - Uninstall\nColumn - Install",
            "choices": [{"question": "Which Excel columns?", "answer": "yes"}],
            "response": {
                "session_id": sid,
                "status": "clarification_required",
                "response": text,
                "clarification": None,
            },
        },
    )
    db_session.add(record)
    db_session.commit()
    response = client.get(f"/chat/sessions/{sid}", headers=headers)
    assert response.status_code == 200
    saved = response.json()["messages"][-1]
    assert saved["content"] == text
    assert len(saved["automation"]["response"]["clarification"]["options"]) == 3
    db_session.refresh(record)
    assert record.automation["response"]["clarification"] is None


@pytest.mark.parametrize("compatibility", [{}, {"confirm": False}, {"confirm": True}])
def test_direct_creation_and_explicit_source_use_existing_pipeline(
    client, build_setup, compatibility
):
    headers, sid, generate, _ = build_setup
    upload = client.post(
        "/chat/documents",
        data={"session_id": sid, "analyze": "false"},
        files={"file": ("source.txt", b"Dish: Idli. Price: 30 INR.", "text/plain")},
        headers=headers,
    ).json()
    source = upload["attachment"]["id"]
    result = call(client, headers, sid, source_document_id=source, **compatibility)
    assert result["status"] == "done"
    generate.assert_awaited_once()
    assert result["attachments"][0]["source_document_ids"] == [source]
    assert result["attachments"][0]["provenance"] == "uploaded_source"


def test_generation_failure_never_returns_attachment(client, build_setup):
    headers, sid, generate, _ = build_setup
    generate.side_effect = ValueError("generation failed")
    result = call(client, headers, sid, confirm=True)
    assert result["status"] == "failed"
    assert not result["attachments"]


def test_preview_preserves_blanks_zero_formula_and_sheet_structure(client, build_setup):
    headers, sid, _, _ = build_setup
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Name", "Blank", "Zero", "Formula"])
    ws.append(["Item", None, 0, "=1+1"])
    wb.create_sheet("Dishes").append(["Idli", 30])
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    uploaded = client.post(
        "/chat/documents",
        data={"session_id": sid, "analyze": "false"},
        files={
            "file": (
                "sample.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text
    document_id = uploaded.json()["attachment"]["id"]
    path = f"/chat/documents/{document_id}/preview"
    preview = client.get(
        path,
        params={
            "sheet": "Summary",
            "row_offset": 1,
            "column_offset": 1,
            "row_limit": 1,
            "column_limit": 3,
        },
        headers=headers,
    ).json()
    assert preview["rows"] == [[None, 0, "=1+1"]]
    assert preview["formulas"] == [[None, None, "=1+1"]]
    assert preview["sheets"] == [
        {"name": "Summary", "rows": 2, "columns": 4},
        {"name": "Dishes", "rows": 1, "columns": 2},
    ]
    assert client.get(path, params={"row_limit": 201}, headers=headers).status_code == 422
    assert client.get(path, params={"sheet": "missing"}, headers=headers).status_code == 404
    assert client.get(path).status_code == 401
