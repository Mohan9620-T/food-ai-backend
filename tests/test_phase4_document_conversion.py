from contextlib import closing
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.models.chat import ChatDocumentAttachment
from app.services.conversion.document_conversion_service import (
    DocumentConversionService,
    LibreOfficeConverter,
    UnsupportedDocumentConversionError,
)
from app.services.document.document_operation_registry import DocumentType, Fidelity
from app.services.document.exceptions import DocumentProcessingUnavailableError
from app.services.document.extraction_models import (
    GeneratedTableContent,
    StructuredDocumentContent,
)
from app.services.pdf.pdf_generator import PdfGenerator
from app.services.powerpoint.pptx_generator import PptxGenerator
from app.services.spreadsheet.excel_generator import ExcelGenerator
from app.services.word.docx_generator import DocxGenerator


def _headers(client, email: str) -> dict[str, str]:
    registered = client.post(
        "/users/",
        json={"fullname": "Phase Four", "email": email, "password": "secret123"},
    )
    assert registered.status_code in (200, 201)
    token = client.post("/users/login", json={"email": email, "password": "secret123"}).json()[
        "access_token"
    ]
    return {"Authorization": f"Bearer {token}"}


def _table_pdf() -> bytes:
    return PdfGenerator().generate_structured(
        StructuredDocumentContent(
            title="Quarterly Sales",
            paragraphs=["Approved sales information."],
            tables=[
                GeneratedTableContent(
                    title="Sales Totals",
                    headers=["Item", "Total"],
                    rows=[["Apple", "12"], ["Pear", "8"]],
                )
            ],
        )
    )


def test_capability_matrix_contains_the_declared_reliable_paths():
    assert set(DocumentConversionService.CAPABILITY_MATRIX) == {
        (DocumentType.PDF, DocumentType.DOCX),
        (DocumentType.PDF, DocumentType.XLSX),
        (DocumentType.PDF, DocumentType.TXT),
        (DocumentType.PDF, DocumentType.MARKDOWN),
        (DocumentType.DOCX, DocumentType.PDF),
        (DocumentType.DOCX, DocumentType.TXT),
        (DocumentType.DOCX, DocumentType.MARKDOWN),
        (DocumentType.PPTX, DocumentType.PDF),
        (DocumentType.XLSX, DocumentType.PDF),
        (DocumentType.XLSX, DocumentType.CSV),
        (DocumentType.CSV, DocumentType.XLSX),
        (DocumentType.TXT, DocumentType.MARKDOWN),
        (DocumentType.MARKDOWN, DocumentType.TXT),
    }


def test_runtime_matrix_disables_office_pdf_paths_when_libreoffice_is_unavailable():
    service = DocumentConversionService(libreoffice=LibreOfficeConverter(binary=""))
    available_pairs = {
        (capability.source_type, capability.output_type)
        for capability in service.available_capabilities()
    }

    assert (DocumentType.PDF, DocumentType.DOCX) in available_pairs
    assert (DocumentType.PDF, DocumentType.XLSX) in available_pairs
    assert (DocumentType.DOCX, DocumentType.PDF) not in available_pairs
    assert (DocumentType.PPTX, DocumentType.PDF) not in available_pairs
    assert (DocumentType.XLSX, DocumentType.PDF) not in available_pairs
    with pytest.raises(DocumentProcessingUnavailableError, match="not available"):
        service.ensure_supported(DocumentType.DOCX, DocumentType.PDF)


def test_pdf_to_docx_reconstructs_text_and_tables_and_discloses_fidelity():
    source = _table_pdf()
    original = bytes(source)
    generated = DocumentConversionService().convert(
        source, "../Quarterly:Report.pdf", DocumentType.DOCX
    )

    assert source == original
    assert generated.filename == "Quarterly-Report-converted.docx"
    assert generated.fidelity == Fidelity.BEST_EFFORT
    assert generated.fidelity_note is not None
    assert "not pixel-perfect" in generated.fidelity_note
    document = Document(BytesIO(generated.file_data))
    assert document.paragraphs[0].text == "Quarterly:Report - Reconstructed"
    assert document.tables[0].cell(1, 0).text == "Apple"


def test_pdf_to_xlsx_requires_and_reopens_real_extracted_tables():
    source = _table_pdf()
    generated = DocumentConversionService().convert(source, "quarterly.pdf", DocumentType.XLSX)

    assert generated.fidelity == Fidelity.PARTIAL
    assert generated.fidelity_note is not None
    with closing(
        load_workbook(BytesIO(generated.file_data), read_only=True, data_only=True)
    ) as workbook:
        assert workbook.sheetnames == ["Page 1 table 1"]
        assert list(workbook.active.values) == [
            ("Item", "Total"),
            ("Apple", "12"),
            ("Pear", "8"),
        ]


def test_pdf_to_xlsx_refuses_to_claim_success_when_no_table_exists():
    source = PdfGenerator().generate("A paragraph without tabular data.")

    with pytest.raises(UnsupportedDocumentConversionError, match="No extractable tables"):
        DocumentConversionService().convert(source, "notes.pdf", DocumentType.XLSX)


@pytest.mark.parametrize(
    ("filename", "file_data"),
    [
        ("report.docx", DocxGenerator().generate("# Report\n\nApproved content")),
        ("slides.pptx", PptxGenerator().generate("# Slides\n\nApproved content")),
        ("figures.xlsx", ExcelGenerator().generate("Item,Total\nApple,12")),
    ],
    ids=("docx", "pptx", "xlsx"),
)
def test_real_office_to_pdf_conversion_when_libreoffice_is_installed(filename, file_data):
    converter = LibreOfficeConverter()
    if not converter.available:
        pytest.skip("LibreOffice is not installed; real Office-to-PDF conversion is unavailable")

    generated = DocumentConversionService(libreoffice=converter).convert(
        file_data, filename, DocumentType.PDF
    )

    assert generated.file_data.startswith(b"%PDF")
    assert generated.fidelity == Fidelity.HIGH


def test_conversion_fidelity_is_assigned_for_existing_deterministic_paths():
    xlsx = DocumentConversionService().convert(
        b"Item,Total\nApple,12\n", "sales.csv", DocumentType.XLSX
    )
    csv_document = DocumentConversionService().convert(
        xlsx.file_data, xlsx.filename, DocumentType.CSV
    )
    markdown = DocumentConversionService().convert(
        b"Approved notes", "notes.txt", DocumentType.MARKDOWN
    )

    assert xlsx.fidelity == Fidelity.HIGH
    assert csv_document.fidelity == Fidelity.PARTIAL
    assert "formatting is not retained" in (csv_document.fidelity_note or "")
    assert markdown.fidelity == Fidelity.FULL


def test_csv_to_xlsx_preserves_multiline_cells_and_ignores_empty_records():
    source = b'Heading\r\n"Line one\nLine two"\r\n\r\nItem,Total\r\nApple,12\r\n'

    converted = DocumentConversionService().convert(source, "multiline.csv", DocumentType.XLSX)

    assert converted.fidelity == Fidelity.HIGH
    with closing(
        load_workbook(BytesIO(converted.file_data), read_only=True, data_only=True)
    ) as workbook:
        assert list(workbook.active.values) == [
            ("Heading", None),
            ("Line one\nLine two", None),
            ("Item", "Total"),
            ("Apple", "12"),
        ]


@pytest.mark.parametrize(
    ("filename", "file_data", "output_type", "expected_fidelity", "expected_text"),
    [
        (
            "notes.pdf",
            PdfGenerator().generate("Approved source content"),
            DocumentType.TXT,
            Fidelity.HIGH,
            "Approved source content",
        ),
        (
            "notes.pdf",
            PdfGenerator().generate("Approved source content"),
            DocumentType.MARKDOWN,
            Fidelity.PARTIAL,
            "Approved source content",
        ),
        (
            "notes.docx",
            DocxGenerator().generate("# Approved\n\nSource content"),
            DocumentType.TXT,
            Fidelity.HIGH,
            "Source content",
        ),
        (
            "notes.docx",
            DocxGenerator().generate("# Approved\n\nSource content"),
            DocumentType.MARKDOWN,
            Fidelity.HIGH,
            "Source content",
        ),
    ],
    ids=("pdf-txt", "pdf-markdown", "docx-txt", "docx-markdown"),
)
def test_pdf_and_docx_text_conversions_are_real_and_reopenable(
    filename, file_data, output_type, expected_fidelity, expected_text
):
    generated = DocumentConversionService().convert(file_data, filename, output_type)

    assert generated.fidelity == expected_fidelity
    assert generated.filename.endswith(".md" if output_type == DocumentType.MARKDOWN else ".txt")
    assert expected_text in generated.file_data.decode("utf-8")


def test_pdf_to_word_acceptance_request_persists_and_downloads_reconstructed_file(
    client, db_session
):
    headers = _headers(client, "phase4-pdf-word@example.com")
    original = _table_pdf()
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("quarterly.pdf", original, "application/pdf")},
        data={"message": "Import quarterly report", "analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text

    response = client.post(
        "/chat/documents/pipeline",
        json={
            "session_id": uploaded.json()["session_id"],
            "instruction": ("Take all the information from this PDF and create a Word document."),
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["steps"][0]["operation"] == "convert_document"
    assert body["steps"][0]["fidelity"] == "BEST_EFFORT"
    assert "not pixel-perfect" in body["response"]
    assert body["attachment"]["filename"] == "quarterly-converted.docx"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    assert Document(BytesIO(download.content)).tables[0].cell(2, 1).text == "8"
    attachments = db_session.query(ChatDocumentAttachment).all()
    assert len(attachments) == 2
    assert attachments[0].kind == "uploaded"
    assert attachments[0].file_data == original


def test_pdf_table_to_excel_acceptance_request_persists_a_nonempty_workbook(client):
    headers = _headers(client, "phase4-pdf-excel@example.com")
    uploaded = client.post(
        "/chat/documents",
        files={"file": ("quarterly.pdf", _table_pdf(), "application/pdf")},
        data={"message": "Import quarterly report", "analyze": "false"},
        headers=headers,
    )
    assert uploaded.status_code == 200, uploaded.text

    response = client.post(
        "/chat/documents/pipeline",
        json={
            "session_id": uploaded.json()["session_id"],
            "instruction": "Extract the table and create an Excel.",
        },
        headers=headers,
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["steps"][0]["operation"] == "extract_table_to_excel"
    download = client.get(f"/chat/documents/{body['attachment']['id']}/download", headers=headers)
    assert download.status_code == 200
    with closing(load_workbook(BytesIO(download.content), read_only=True)) as workbook:
        assert list(workbook.active.values)[1] == ("Apple", "12")
