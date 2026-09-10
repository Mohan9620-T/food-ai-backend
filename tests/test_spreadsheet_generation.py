from datetime import date, datetime
from io import BytesIO
from typing import Sequence

import pytest
from openpyxl import Workbook, load_workbook

from app.services.chat_document_service import ChatDocumentService, InvalidDocumentError

SPLIT_PROMPT = (
    "I have many categories. Please split the item list category-wise and create a "
    "separate sheet for each category, with the corresponding items added to each sheet."
)
FILTER_PROMPT = (
    "The Excel sheet contains different types of Category values. "
    "Please add a filter only to the Category* column."
)


@pytest.mark.parametrize(
    "instruction",
    [
        "Split this Excel file category-wise.",
        "Create separate sheets for each category.",
        "Group the items by category and create an Excel file.",
        "Arrange the item list category-wise in separate sheets.",
        "Create a new Excel workbook with one sheet per category.",
    ],
)
def test_common_category_split_instructions_are_detected(instruction):
    assert ChatDocumentService.is_spreadsheet_update_request(instruction) is True


@pytest.mark.parametrize(
    "instruction",
    [
        "Please add a filter only to the Category* column.",
        "Add a filter to Category*.",
        "Filter the Category column.",
        "Please enable filtering for the Category* column.",
        "Add an Excel filter only to the Category column.",
        "Add a dropdown filter to the Category* column.",
        "Create a filter only for the Category column.",
        "Add a filter to the Category column and generate the Excel file.",
    ],
)
def test_filter_column_instructions_are_detected(instruction):
    assert ChatDocumentService.is_spreadsheet_update_request(instruction) is True
    assert ChatDocumentService.is_filter_column_request(instruction) is True


def _workbook_bytes(
    header: str = "Category", rows: Sequence[tuple[object, ...]] | None = None
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Items"
    sheet.append(["Item", header, "Quantity", "Created"])
    workbook_rows = (
        rows
        if rows is not None
        else [
            ("Apple", "Fruit", 2, date(2026, 9, 1)),
            ("Carrot", "Vegetable", 3, date(2026, 9, 2)),
            ("Banana", "Fruit", 4, date(2026, 9, 3)),
        ]
    )
    for row in workbook_rows:
        sheet.append(row)
    sheet["D2"].number_format = "yyyy-mm-dd"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.parametrize(
    "header",
    ["Category", "category", "CATEGORY", "Category*", "Item Category", "Category Name"],
)
def test_category_header_variations_create_valid_category_workbook(header):
    generated, filename, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes(header), "products.xlsx", SPLIT_PROMPT
    )

    assert filename == "products_category_wise.xlsx"
    workbook = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert workbook.sheetnames == ["Fruit", "Vegetable"]
        assert list(workbook["Fruit"].values) == [
            ("Item", header, "Quantity", "Created"),
            ("Apple", "Fruit", 2, datetime(2026, 9, 1)),
            ("Banana", "Fruit", 4, datetime(2026, 9, 3)),
        ]
        assert list(workbook["Vegetable"].values)[1][0] == "Carrot"
        assert workbook["Fruit"]["C2"].data_type == "n"
        assert workbook["Fruit"]["D2"].number_format == "yyyy-mm-dd"
    finally:
        workbook.close()


def test_explicit_alternate_grouping_column_is_supported():
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes("Department"),
        "departments.xlsx",
        "Split the item list by Department into a separate sheet for each value.",
    )

    workbook = load_workbook(BytesIO(generated), read_only=True)
    try:
        assert workbook.sheetnames == ["Fruit", "Vegetable"]
    finally:
        workbook.close()


def test_blank_invalid_duplicate_and_long_categories_get_safe_unique_sheet_names():
    rows = [
        ("One", "Dry/Goods", 1, None),
        ("Two", "Dry:Goods", 2, None),
        ("Three", "A category name that is much longer than thirty one characters", 3, None),
        ("Four", None, 4, None),
        ("Five", "   ", 5, None),
    ]
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes(rows=rows), "unsafe.xlsx", SPLIT_PROMPT
    )

    workbook = load_workbook(BytesIO(generated), read_only=True)
    try:
        assert workbook.sheetnames[:2] == ["Dry - Goods", "Dry - Goods (2)"]
        assert workbook.sheetnames[-1] == "Uncategorized"
        assert workbook["Uncategorized"].max_row == 3
        assert all(len(name) <= 31 for name in workbook.sheetnames)
        assert not any(
            any(character in name for character in r":\/?*[]") for name in workbook.sheetnames
        )
    finally:
        workbook.close()


def test_generated_workbook_is_formatted_and_original_upload_is_unchanged():
    source = _workbook_bytes()
    original_snapshot = bytes(source)

    generated, _, _ = ChatDocumentService().format_spreadsheet(
        source, "inventory.xlsx", SPLIT_PROMPT
    )

    assert source == original_snapshot
    original = load_workbook(BytesIO(source), read_only=True)
    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert original.sheetnames == ["Items"]
        assert revised.sheetnames == ["Fruit", "Vegetable"]
        fruit = revised["Fruit"]
        assert fruit.freeze_panes == "A2"
        assert fruit.auto_filter.ref == fruit.dimensions
        assert fruit["A1"].font.bold is True
        assert fruit["A1"].fill.fgColor.rgb == "001F4E78"
        assert fruit["A2"].alignment.horizontal == "left"
        assert fruit["C2"].alignment.horizontal == "right"
        assert fruit["A2"].alignment.wrap_text is True
        assert 10 <= fruit.column_dimensions["A"].width <= 60
    finally:
        original.close()
        revised.close()


def test_filter_finds_category_header_and_creates_single_column_autofilter():
    source = _workbook_bytes("Category*")
    original_snapshot = bytes(source)

    generated, filename, actions = ChatDocumentService().format_spreadsheet(
        source, "KSCB ItemMaster.xlsx", FILTER_PROMPT
    )

    assert source == original_snapshot
    assert filename == "KSCB_ItemMaster_filter_updated.xlsx"
    assert "Category* column" in " ".join(actions)
    original = load_workbook(BytesIO(source), data_only=False)
    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        original_sheet = original["Items"]
        revised_sheet = revised["Items"]
        category_header = ChatDocumentService._find_requested_header(
            tuple(revised_sheet[1]), "Category"
        )
        assert category_header is not None
        assert category_header.column == 2
        assert revised_sheet.auto_filter.ref == "B1:B4"
        assert original_sheet.auto_filter.ref is None
        assert list(revised_sheet.values) == list(original_sheet.values)
        assert revised_sheet.max_row == original_sheet.max_row
        assert revised_sheet.max_column == original_sheet.max_column
        assert revised_sheet["B1"].value == "Category*"
        assert revised_sheet["D2"].number_format == original_sheet["D2"].number_format
    finally:
        original.close()
        revised.close()


def test_filter_uses_actual_category_column_position_and_preserves_formulas():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ItemCode", "ItemName", "Category*", "Uom", "Total"])
    sheet.append(["BNS001", "Coffee", "Beverages", "CTN", "=1+1"])
    sheet.append(["BNS002", "Tea", "Beverages", "PKT", "=2+2"])
    sheet.append(["CHC001", "Chicken", "Meat", "KG", "=3+3"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    generated, _, _ = ChatDocumentService().format_spreadsheet(
        source.getvalue(), "items.xlsx", FILTER_PROMPT
    )

    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised.active.auto_filter.ref == "C1:C4"
        assert revised.active["E2"].value == "=1+1"
        assert revised.active["E2"].data_type == "f"
        assert list(revised.active.values)[0] == (
            "ItemCode",
            "ItemName",
            "Category*",
            "Uom",
            "Total",
        )
    finally:
        revised.close()


def test_filter_converts_csv_source_to_valid_xlsx_without_changing_values():
    source = b"ItemCode,ItemName,Category*,Uom\r\nBNS001,Coffee,Beverages,CTN\r\nCHC001,Chicken,Meat,KG\r\n"
    original_snapshot = bytes(source)

    generated, filename, _ = ChatDocumentService().format_spreadsheet(
        source, "items.csv", FILTER_PROMPT
    )

    assert source == original_snapshot
    assert filename == "items_filter_updated.xlsx"
    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised.active.auto_filter.ref == "C1:C3"
        assert list(revised.active.values) == [
            ("ItemCode", "ItemName", "Category*", "Uom"),
            ("BNS001", "Coffee", "Beverages", "CTN"),
            ("CHC001", "Chicken", "Meat", "KG"),
        ]
    finally:
        revised.close()


def test_kscb_sized_filter_workbook_preserves_807_items_and_41_suppliers():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Item Master"
    headers = ["ItemCode", "ItemName", "Category*", "Supplier", "Uom", "UnitPrice"]
    sheet.append(headers)
    for item_number in range(1, 808):
        sheet.append(
            [
                f"ITEM{item_number:04d}",
                f"Item {item_number}",
                f"Category {(item_number - 1) % 12 + 1}",
                f"Supplier {(item_number - 1) % 41 + 1}",
                "EA",
                item_number / 10,
            ]
        )
    source_buffer = BytesIO()
    workbook.save(source_buffer)
    workbook.close()
    source = source_buffer.getvalue()
    original_snapshot = bytes(source)

    generated, filename, _ = ChatDocumentService().format_spreadsheet(
        source,
        "KSCB_ItemMaster_ERP_CSSB_807Items_41Suppliers.xlsx",
        FILTER_PROMPT,
    )

    assert source == original_snapshot
    assert filename == "KSCB_ItemMaster_ERP_CSSB_807Items_41Suppliers_filter_updated.xlsx"
    original = load_workbook(BytesIO(source), data_only=False)
    revised = load_workbook(BytesIO(generated), data_only=False)
    try:
        assert revised["Item Master"].auto_filter.ref == "C1:C808"
        assert revised["Item Master"].max_row == 808
        assert revised["Item Master"].max_column == len(headers)
        assert list(revised["Item Master"].values) == list(original["Item Master"].values)
        assert (
            len({row[3] for row in revised["Item Master"].iter_rows(min_row=2, values_only=True)})
            == 41
        )
    finally:
        original.close()
        revised.close()


@pytest.mark.parametrize(
    "header",
    [
        "Category*",
        "Category",
        "category*",
        "category",
        "CATEGORY*",
        "CATEGORY",
        "Item Category",
        "Category Name",
    ],
)
def test_filter_matches_supported_category_header_variations_without_renaming(header):
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes(header), "items.xlsx", "Add a filter to the Category column."
    )

    revised = load_workbook(BytesIO(generated), read_only=False)
    try:
        assert revised["Items"].auto_filter.ref == "B1:B4"
        assert revised["Items"]["B1"].value == header
    finally:
        revised.close()


def test_filter_replaces_category_split_full_range_without_losing_category_sheets():
    service = ChatDocumentService()
    category_workbook, _, _ = service.format_spreadsheet(
        _workbook_bytes("Category*"), "items.xlsx", SPLIT_PROMPT
    )

    filtered, _, _ = service.format_spreadsheet(
        category_workbook, "items_category_wise.xlsx", FILTER_PROMPT
    )

    revised = load_workbook(BytesIO(filtered), data_only=False)
    try:
        assert revised.sheetnames == ["Fruit", "Vegetable"]
        assert revised["Fruit"].auto_filter.ref == "B1:B3"
        assert revised["Vegetable"].auto_filter.ref == "B1:B2"
        assert list(revised["Fruit"].values)[1][0] == "Apple"
    finally:
        revised.close()


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        (b"not an xlsx", "valid XLSX"),
        (_workbook_bytes("Department"), "couldn't find a Category\\* column"),
        (_workbook_bytes("Category*", rows=[]), "no data rows"),
    ],
)
def test_filter_invalid_workbooks_return_clean_errors(source, expected):
    with pytest.raises(InvalidDocumentError, match=expected):
        ChatDocumentService().format_spreadsheet(source, "items.xlsx", FILTER_PROMPT)


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        (b"not an xlsx", "valid XLSX"),
        (_workbook_bytes("Department"), "Category column"),
        (_workbook_bytes(rows=[]), "no item rows"),
    ],
)
def test_invalid_workbooks_return_clean_actionable_errors(source, expected):
    with pytest.raises(InvalidDocumentError, match=expected):
        ChatDocumentService().format_spreadsheet(source, "items.xlsx", SPLIT_PROMPT)


def test_empty_workbook_returns_clean_error():
    workbook = Workbook()
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(InvalidDocumentError, match="spreadsheet is empty"):
        ChatDocumentService().format_spreadsheet(buffer.getvalue(), "empty.xlsx", SPLIT_PROMPT)


def test_multi_sheet_selection_uses_largest_item_table_and_allows_explicit_sheet():
    workbook = Workbook()
    small = workbook.active
    small.title = "Small"
    small.append(["Item", "Category"])
    small.append(["Apple", "Fruit"])
    large = workbook.create_sheet("Main Items")
    large.append(["Item", "Category"])
    large.append(["Rice", "Grocery"])
    large.append(["Beans", "Grocery"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    service = ChatDocumentService()
    generated, _, _ = service.format_spreadsheet(buffer.getvalue(), "items.xlsx", SPLIT_PROMPT)
    largest = load_workbook(BytesIO(generated), read_only=True)
    try:
        assert largest.sheetnames == ["Grocery"]
        assert largest["Grocery"].max_row == 3
    finally:
        largest.close()

    selected, _, _ = service.format_spreadsheet(
        buffer.getvalue(), "items.xlsx", f"{SPLIT_PROMPT} Use worksheet Small."
    )
    explicit = load_workbook(BytesIO(selected), read_only=True)
    try:
        assert explicit.sheetnames == ["Fruit"]
    finally:
        explicit.close()


def test_tied_multi_sheet_item_tables_require_an_explicit_sheet_name():
    workbook = Workbook()
    for index, title in enumerate(("First", "Second")):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(["Item", "Category"])
        sheet.append([title, "Group"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    with pytest.raises(InvalidDocumentError, match="More than one worksheet"):
        ChatDocumentService().format_spreadsheet(buffer.getvalue(), "items.xlsx", SPLIT_PROMPT)


def test_large_category_workbook_is_complete_and_reopenable():
    rows = [(f"Item {index}", f"Category {index % 20}", index, None) for index in range(5_000)]
    generated, _, _ = ChatDocumentService().format_spreadsheet(
        _workbook_bytes(rows=rows), "large.xlsx", SPLIT_PROMPT
    )

    workbook = load_workbook(BytesIO(generated), read_only=True)
    try:
        assert len(workbook.sheetnames) == 20
        assert sum(sheet.max_row - 1 for sheet in workbook.worksheets) == 5_000
    finally:
        workbook.close()
