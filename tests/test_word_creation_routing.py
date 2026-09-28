import json
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from app.api import chat_documents
from app.services.chat_document_service import ChatDocumentService
from app.services.document.document_automation_service import DocumentAutomationService
from app.services.document.document_intent_service import DocumentIntentService
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.web_lookup import automatic_web_question
from app.utils.document_output import requests_word_output


@pytest.mark.parametrize(
    "prompt",
    [
        "Can you give me a Word document about workplace organization?",
        "Can you crate a Word document with a weekly schedule?",
        "Give me this content as a Word document?",
        "word document ah kudu",
        "word document create pannu",
        "docx file venum",
    ],
)
def test_word_delivery_is_not_an_automatic_web_question(prompt):
    assert requests_word_output(prompt)
    assert DocumentIntentService.is_creation_request(prompt)
    assert not automatic_web_question(prompt)


@pytest.mark.parametrize(
    "prompt",
    [
        "Read this Word document and give me a summary",
        'What does "create a Word document" mean?',
        "Create a PDF from this Word document",
        "Create an Excel file. Do not create a Word document.",
        "What is the latest Microsoft Word version?",
    ],
)
def test_input_format_and_quoted_content_do_not_select_word_output(prompt):
    assert not requests_word_output(prompt)


@pytest.mark.parametrize(
    "prompt",
    [
        "Create a Word document about workplace organization",
        "Can you give me a Word document with a weekly schedule?",
        "Word document create pannu about gardening",
    ],
)
def test_complete_word_brief_does_not_depend_on_a_model_planner(prompt):
    model = MagicMock()
    model.chat.side_effect = AssertionError("An explicit brief does not need an AI planner")
    plan = DocumentAutomationService(chat_service=model).plan(prompt, ())
    assert plan.status == "ready"
    assert len(plan.steps) == 1
    assert plan.steps[0].intent.operation == DocumentOperation.CREATE_DOCUMENT
    assert plan.steps[0].intent.output_type == DocumentType.DOCX
    assert plan.steps[0].source_filenames == ()
    model.chat.assert_not_called()


def test_word_export_keeps_the_conversion_route():
    assert not DocumentIntentService.is_creation_request("Export this PDF to Word")
    assert DocumentIntentService.is_conversion_request("Export this PDF to Word")


@pytest.mark.parametrize(
    "prompt", ["Create a Word document", "Give me a Word document for me please"]
)
def test_missing_word_content_still_asks_for_a_brief(prompt):
    model = MagicMock()
    model.chat.return_value = json.dumps(
        {
            "status": "clarification_required",
            "steps": [],
            "clarifying_question": "What should the Word document contain?",
        }
    )
    plan = DocumentAutomationService(chat_service=model).plan(prompt, ())
    assert plan.status == "clarification_required"
    assert plan.steps == ()
    assert "contain" in plan.clarifying_question


def test_word_request_creates_downloadable_file_and_persists_it(client, monkeypatch):
    credentials = {"email": "word-delivery@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Word delivery test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    session = client.post("/chat/sessions", headers=headers, json={"title": "Word test"}).json()
    planner = MagicMock(side_effect=AssertionError("No separate planner needed"))
    monkeypatch.setattr(chat_documents.automation_service.chat_service, "chat", planner)
    completion = AsyncMock(
        return_value=json.dumps(
            {
                "title": "Workplace Organization",
                "paragraphs": ["Keep shared work areas ready for use."],
                "tables": [
                    {
                        "title": "Weekly Checklist",
                        "headers": ["Task", "Frequency"],
                        "rows": [["Clear the desk", "Daily"], ["Review supplies", "Weekly"]],
                    }
                ],
            }
        )
    )
    monkeypatch.setattr(ChatDocumentService, "_complete", completion)
    response = client.post(
        "/chat/documents/automate",
        headers=headers,
        json={
            "session_id": session["id"],
            "instruction": "Can you give me a Word document about workplace organization with a checklist?",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "done", body
    assert len(body["attachments"]) == 1
    attachment = body["attachments"][0]
    assert attachment["filename"].endswith(".docx")
    download = client.get(f"/chat/documents/{attachment['id']}/download", headers=headers)
    assert download.status_code == 200
    assert "wordprocessingml.document" in download.headers["content-type"]
    document = Document(BytesIO(download.content))
    assert "Workplace Organization" in [paragraph.text for paragraph in document.paragraphs]
    assert document.tables[0].cell(1, 0).text == "Clear the desk"
    saved = client.get(f"/chat/sessions/{session['id']}", headers=headers).json()
    assert saved["messages"][-1]["attachments"][0]["id"] == attachment["id"]
    planner.assert_not_called()
