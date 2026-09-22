import threading
from io import BytesIO
from time import monotonic
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.api import chat, chat_documents
from app.models.chat import ChatDocumentAttachment, ChatMessageRecord
from app.schemas.vision_result import VisionItem, VisionResult
from app.services.chat_document_service import ChatDocumentService
from app.services.document.document_automation_service import (
    AvailableDocument,
    DocumentAutomationService,
)
from app.services.document.document_operation_registry import DocumentOperation, DocumentType
from app.services.document.exceptions import DocumentProcessingUnavailableError
from app.services.document.extraction_models import StructuredDocumentContent
from app.services.document.image_document_reader import ImageDocumentReader
from app.services.image_parser_service import VisionModelUnavailableError
from app.services.vision_providers.failover_provider import FailoverVisionProvider
from app.services.vision_runtime import vision_inference_slot

TRANSCRIPTION = "| Item name | Country name |\n| --- | --- |\n| Curry | India |\n| Sushi | Japan |"


@pytest.fixture
def headers(client):
    credentials = {"email": "image-file@example.com", "password": "test-password"}
    client.post("/users/", json={**credentials, "fullname": "Image file test"})
    token = client.post("/users/login", json=credentials).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def upload(client, headers, image, **data):
    return client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false", **data},
        files={"file": ("food.png", image, "image/png")},
    )


def test_image_save_only_preserves_original_without_running_models(
    client, headers, monkeypatch, valid_png_bytes, db_session
):
    model = MagicMock(side_effect=AssertionError("Save-only must not call a model"))
    monkeypatch.setattr(type(chat_documents.service), "extract_document", model)
    response = upload(client, headers, valid_png_bytes)
    assert response.status_code == 200
    result = response.json()
    assert result["analysis_status"] == "skipped"
    document = db_session.get(ChatDocumentAttachment, result["attachment"]["id"])
    assert document.raw_text == ""
    assert document.file_data == valid_png_bytes
    assert document.message.image_data == valid_png_bytes
    assert (
        client.get(f"/chat/documents/{document.id}/download", headers=headers).content
        == valid_png_bytes
    )
    model.assert_not_called()


@pytest.mark.parametrize("mime,data", [("image/jpeg", b"not an image"), ("image/png", b"%PDF-")])
def test_invalid_image_cannot_be_saved(client, headers, mime, data):
    response = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": ("food.png", data, mime)},
    )
    assert response.status_code in {415, 422}


@pytest.mark.parametrize("format", ["xlsx", "docx", "pdf"])
def test_image_direct_creation_download_reopens_with_matching_parser(
    client, headers, monkeypatch, valid_png_bytes, format, db_session
):
    provider = MagicMock()
    provider.infer.return_value = VisionResult(image_type="text", answer=TRANSCRIPTION)
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    content = StructuredDocumentContent.model_validate(
        {
            "title": "Food labels",
            "tables": [
                {
                    "title": "Items",
                    "headers": ["Item name", "Country name"],
                    "rows": [["Curry", "India"], ["Sushi", "Japan"]],
                }
            ],
        }
    )
    generate = AsyncMock(return_value=content)
    monkeypatch.setattr(ChatDocumentService, "generate_content", generate)
    source = upload(client, headers, valid_png_bytes).json()
    payload = {
        "session_id": source["session_id"],
        "source_document_id": source["attachment"]["id"],
        "instruction": f"Read this image and create a {format} document with item name and country name",
    }
    result = client.post("/chat/documents/automate", headers=headers, json=payload).json()
    assert result["status"] == "done", result
    attachment = result["attachments"][0]
    assert attachment["source_document_ids"] == [source["attachment"]["id"]]
    assert attachment["provenance"] == "uploaded_source"
    assert result["steps"][0]["fidelity"] == "BEST_EFFORT"
    data = client.get(f"/chat/documents/{attachment['id']}/download", headers=headers).content
    if format == "xlsx":
        workbook = load_workbook(BytesIO(data))
        assert list(workbook["Items"].values) == [
            ("Item name", "Country name"),
            ("Curry", "India"),
            ("Sushi", "Japan"),
        ]
        assert workbook["Items"]["A1"].font.bold
        workbook.close()
    elif format == "docx":
        assert Document(BytesIO(data)).tables[0].cell(1, 0).text == "Curry"
    else:
        assert "Curry" in "".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    extracted = generate.call_args.kwargs["documents"][0]
    assert extracted.text == TRANSCRIPTION
    assert extracted.tables[0].rows[-1] == ("Sushi", "Japan")
    assert provider.infer.call_args.kwargs["max_tokens"] == 4096
    db_session.expire_all()
    stored_source = db_session.get(ChatDocumentAttachment, source["attachment"]["id"])
    assert stored_source.raw_text == TRANSCRIPTION
    assert stored_source.extracted_data["tables"][0]["rows"][-1] == ["Sushi", "Japan"]
    stored_output = db_session.get(ChatDocumentAttachment, attachment["id"])
    assert "Curry" in stored_output.raw_text
    assert stored_output.extracted_data["document_type"] == format
    assert stored_output.generation_metadata["content"]["tables"][0]["rows"][0] == [
        "Curry",
        "India",
    ]
    assert (
        stored_output.generation_metadata["source_extractions"][0]["data"]["text"] == TRANSCRIPTION
    )
    history = client.get(f"/chat/sessions/{source['session_id']}", headers=headers).json()[
        "messages"
    ]
    assert history[-1]["automation"]["response"] == result
    assert history[-1]["attachments"][0]["id"] == attachment["id"]
    saved_data = client.get(f"/chat/documents/{attachment['id']}/data", headers=headers)
    assert saved_data.status_code == 200
    assert saved_data.json()["extracted_data"] == stored_output.extracted_data


@pytest.mark.parametrize(
    "instruction",
    [
        "Create Excel from this image with item name and country name",
        "Create a Word document on this content",
    ],
)
def test_old_image_turn_becomes_owned_source_without_duplicate_attachments(
    client, headers, monkeypatch, valid_png_bytes, db_session, instruction
):
    monkeypatch.setattr(type(chat.vision_service), "describe", lambda *args: "Food labels")
    monkeypatch.setattr(
        ImageDocumentReader,
        "read",
        lambda *args, **kwargs: ImageDocumentReader._document(TRANSCRIPTION, ()),
    )
    monkeypatch.setattr(
        ChatDocumentService,
        "generate_content",
        AsyncMock(
            return_value=StructuredDocumentContent(title="Labels", paragraphs=["Curry India"])
        ),
    )
    source = client.post(
        "/chat/vision", headers=headers, files={"image": ("food.png", valid_png_bytes, "image/png")}
    ).json()
    payload = {"session_id": source["session_id"], "instruction": instruction}
    for _ in range(2):
        result = client.post("/chat/documents/automate", headers=headers, json=payload)
        assert result.status_code == 200
        assert result.json()["status"] == "done", result.json()
    images = db_session.query(ChatDocumentAttachment).filter_by(kind="uploaded").all()
    assert len(images) == 1
    assert images[0].file_data == valid_png_bytes
    assert (
        db_session.query(ChatMessageRecord).filter(ChatMessageRecord.image_data.isnot(None)).count()
        == 1
    )
    other = client.post(
        "/users/",
        json={"email": "other-image@example.com", "password": "test-password", "fullname": "Other"},
    )
    assert other.status_code in {200, 201}
    token = client.post(
        "/users/login", json={"email": "other-image@example.com", "password": "test-password"}
    ).json()["access_token"]
    assert (
        client.post(
            "/chat/documents/automate", headers={"Authorization": f"Bearer {token}"}, json=payload
        ).status_code
        == 404
    )


@pytest.mark.parametrize(
    "instruction,filename",
    [
        ("Create Excel from this image", "second.png"),
        ("Convert the first image to Word", "first.png"),
        ("Create a PDF from the previous image", "first.png"),
        ("Extract the table from this image into Excel", "second.png"),
    ],
)
def test_image_source_selection_ignores_old_workbook(instruction, filename):
    model = MagicMock(side_effect=AssertionError("Planning should be deterministic"))
    documents = (
        AvailableDocument(1, "first.png", DocumentType.IMAGE),
        AvailableDocument(2, "second.png", DocumentType.IMAGE),
        AvailableDocument(3, "old.xlsx", DocumentType.XLSX, is_latest=True),
    )
    plan = DocumentAutomationService(chat_service=model).plan(instruction, documents)
    assert plan.status == "ready"
    assert plan.steps[0].intent.operation == DocumentOperation.CREATE_DOCUMENT
    assert plan.steps[0].source_filenames == (filename,)


def test_missing_format_asks_before_building():
    plan = DocumentAutomationService(chat_service=MagicMock()).plan(
        "Create a document from this image", (AvailableDocument(1, "food.png", DocumentType.IMAGE),)
    )
    assert plan.status == "clarification_required"
    assert "Excel, Word, or PDF" in plan.clarifying_question


def test_unavailable_vision_and_ocr_do_not_become_document_data(monkeypatch, valid_png_bytes):
    provider = MagicMock()
    provider.infer.side_effect = VisionModelUnavailableError("offline")
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    monkeypatch.setattr(
        "pytesseract.image_to_string", MagicMock(side_effect=OSError("missing OCR"))
    )
    with pytest.raises(DocumentProcessingUnavailableError, match="Image extraction is unavailable"):
        ImageDocumentReader().read(valid_png_bytes)


def test_local_ocr_fallback_is_marked_for_review(monkeypatch, valid_png_bytes):
    provider = MagicMock()
    provider.infer.side_effect = ValueError("invalid JSON")
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    monkeypatch.setattr("pytesseract.image_to_string", lambda *args, **kwargs: "Curry India")
    result = ImageDocumentReader().read(valid_png_bytes)
    assert result.text == "Curry India"
    assert any("Local OCR" in warning for warning in result.warnings)


def test_confident_local_ocr_keeps_positions_without_cloud_call(monkeypatch, valid_png_bytes):
    words = ["Curry", "India", "Sushi", "Japan", "Pizza", "Italy", "Paella", "Spain"]
    data = {
        "text": words,
        "conf": [95] * 8,
        "block_num": list(range(8)),
        "par_num": [1] * 8,
        "line_num": [1] * 8,
        "left": [10, 10, 100, 100, 200, 200, 300, 300],
        "top": [10, 30] * 4,
    }
    monkeypatch.setattr("pytesseract.image_to_data", lambda *args, **kwargs: data)
    provider = MagicMock(side_effect=AssertionError("Readable labels should use local OCR"))
    monkeypatch.setattr("app.services.document.image_document_reader.get_vision_provider", provider)
    result = ImageDocumentReader().read(valid_png_bytes)
    assert '"left": 10, "top": 30, "text": "India"' in result.text
    assert all(word in result.text for word in words)
    assert result.used_ocr
    provider.assert_not_called()


def test_scene_request_uses_vision_even_if_ocr_finds_readable_labels(monkeypatch, valid_png_bytes):
    monkeypatch.setattr(
        ImageDocumentReader, "_local_text", lambda *args: "A sign with readable words"
    )
    provider = MagicMock()
    provider.infer.return_value = VisionResult(
        image_type="other", answer="There are three visible chairs."
    )
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    document = ChatDocumentService().extract_document(
        valid_png_bytes,
        "room.png",
        image_instruction="Count the objects in this photo and create Excel",
    )
    assert "three visible chairs" in document.text
    assert document.source.filename == "room.png"
    assert "Count the objects" in provider.infer.call_args.kwargs["user_prompt"]


def test_visual_failure_cannot_be_substituted_with_unrelated_ocr(monkeypatch, valid_png_bytes):
    monkeypatch.setattr(
        ImageDocumentReader, "_local_text", lambda *args: "A sign with readable words"
    )
    provider = MagicMock()
    provider.infer.side_effect = VisionModelUnavailableError("offline")
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    with pytest.raises(
        DocumentProcessingUnavailableError, match="Text extraction alone cannot verify"
    ):
        ImageDocumentReader().read(valid_png_bytes, instruction="Count the people in the photo")


def test_scene_items_are_valid_evidence_when_provider_omits_answer(monkeypatch, valid_png_bytes):
    provider = MagicMock()
    provider.infer.return_value = VisionResult(
        image_type="other",
        items=[
            VisionItem(
                name="red square", confidence="high", visual_evidence="One red square on the left"
            )
        ],
    )
    monkeypatch.setattr(
        "app.services.document.image_document_reader.get_vision_provider", lambda: provider
    )
    document = ImageDocumentReader().read(
        valid_png_bytes, instruction="Count the objects and create Excel"
    )
    assert "red square" in document.text
    assert "One red square on the left" in document.text


@pytest.mark.parametrize(
    "format,mime,extension",
    [
        ("JPEG", "image/jpeg", "jpg"),
        ("PNG", "image/png", "png"),
        ("WEBP", "image/webp", "webp"),
        ("GIF", "image/gif", "gif"),
    ],
)
def test_common_image_formats_remain_downloadable_without_conversion(
    client, headers, format, mime, extension
):
    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (80, 60), "white").save(buffer, format=format)
    original = buffer.getvalue()
    response = client.post(
        "/chat/documents",
        headers=headers,
        data={"analyze": "false"},
        files={"file": (f"source.{extension}", original, mime)},
    )
    assert response.status_code == 200
    attachment = response.json()["attachment"]
    assert attachment["content_type"] == mime
    assert (
        client.get(f"/chat/documents/{attachment['id']}/download", headers=headers).content
        == original
    )


def test_image_failover_preserves_token_budget_and_remaining_deadline():
    nvidia, ollama = MagicMock(), MagicMock()
    nvidia.infer.side_effect = VisionModelUnavailableError("offline")
    ollama.infer.return_value = VisionResult(image_type="text", answer=TRANSCRIPTION)
    FailoverVisionProvider(nvidia, ollama).infer(
        "system", "prompt", "image", max_tokens=4096, timeout_seconds=20
    )
    assert ollama.infer.call_args.kwargs["max_tokens"] == 4096
    assert 0 < ollama.infer.call_args.kwargs["timeout_seconds"] <= 20


def test_image_queue_wait_is_bounded():
    with vision_inference_slot():
        with pytest.raises(VisionModelUnavailableError, match="busy"):
            with vision_inference_slot(timeout_seconds=0.01):
                pytest.fail("The occupied slot must not be entered")


def test_vision_request_does_not_block_health_endpoint(
    client, headers, monkeypatch, valid_png_bytes
):
    entered, release = threading.Event(), threading.Event()

    def slow(*args):
        entered.set()
        release.wait(3)
        return "Food labels"

    monkeypatch.setattr(type(chat.vision_service), "describe", slow)
    worker = threading.Thread(
        target=lambda: client.post(
            "/chat/vision",
            headers=headers,
            files={"image": ("food.png", valid_png_bytes, "image/png")},
        )
    )
    worker.start()
    try:
        assert entered.wait(2)
        started = monotonic()
        assert client.get("/health").status_code == 200
        assert monotonic() - started < 1
        assert not release.is_set()
    finally:
        release.set()
        worker.join(4)
