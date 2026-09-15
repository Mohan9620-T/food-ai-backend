import asyncio
from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

from app.models.chat import ChatDocumentAttachment, ChatMessageRecord, ChatSession
from app.repositories.chat_repository import ChatRepository
from app.services.chat_document_service import ChatDocumentService
from app.services.chat_service import ChatModelUnavailableError


def _headers(client, email="resilient-upload@example.com"):
    response = client.post(
        "/users/", json={"fullname": "Upload User", "email": email, "password": "secret123"}
    )
    assert response.status_code in (200, 201)
    login = client.post("/users/login", json={"email": email, "password": "secret123"})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def _file_bytes(extension):
    text = "Company: Catering Solutions; Status: Draft"
    buffer = BytesIO()
    if extension == "pdf":
        canvas = Canvas(buffer)
        canvas.drawString(72, 750, text)
        canvas.save()
    elif extension == "docx":
        document = Document()
        document.add_paragraph(text)
        document.save(buffer)
    elif extension == "xlsx":
        workbook = Workbook()
        workbook.active.append([text])
        workbook.save(buffer)
        workbook.close()
    else:
        return text.encode("utf-8")
    return buffer.getvalue()


def _upload_without_ai(client, headers, text="Company: Catering Solutions"):
    result = client.post(
        "/chat/documents",
        files={"file": ("source.txt", text.encode("utf-8"), "text/plain")},
        data={"analyze": "false"},
        headers=headers,
    )
    assert result.status_code == 200, result.text
    return result.json()


@pytest.mark.parametrize("extension", ["pdf", "docx", "txt", "csv", "xlsx"])
def test_readable_upload_is_saved_and_downloadable_when_ai_is_unavailable(
    client, monkeypatch, db_session, extension
):
    headers = _headers(client)

    async def unavailable(*args):
        raise ChatModelUnavailableError("Document AI generation timed out")

    monkeypatch.setattr(ChatDocumentService, "summarize", unavailable)
    original = _file_bytes(extension)
    result = client.post(
        "/chat/documents",
        files={"file": (f"source.{extension}", original, "application/octet-stream")},
        data={"message": "Read the company details"},
        headers=headers,
    )
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["analysis_status"] == "unavailable"
    assert "extracted" in body["response"].lower()
    assert "ai" in body["response"].lower()
    assert "Catering Solutions" in body["response"]
    saved = db_session.query(ChatDocumentAttachment).one()
    assert saved.kind == "uploaded"
    expected_text = "Company: Catering Solutions; Status: Draft"
    if extension == "xlsx":
        expected_text = f"[Sheet]\n{expected_text}"
    assert saved.raw_text == expected_text
    assert saved.structured_summary is None
    assert saved.file_data == original
    download = client.get(f"/chat/documents/{saved.id}/download", headers=headers)
    assert download.status_code == 200
    assert download.content == original
    messages = client.get(f"/chat/sessions/{body['session_id']}", headers=headers).json()[
        "messages"
    ]
    assert len(messages) == 2
    assert messages[0]["document_attachment"]["id"] == saved.id
    assert messages[1]["content"] == body["response"]
    assert db_session.query(ChatSession).count() == 1


def test_upload_is_persisted_before_waiting_for_ai(client, monkeypatch, db_session):
    headers = _headers(client)
    observations = []

    async def summarize(self, raw_text, instruction):
        attachment = db_session.query(ChatDocumentAttachment).one()
        observations.append((attachment.raw_text, attachment.file_data))
        assert attachment.structured_summary is None
        assert db_session.query(ChatMessageRecord).count() >= 1
        assert db_session.query(ChatSession).count() == 1
        await asyncio.sleep(0)
        raise ChatModelUnavailableError("AI timed out after upload was saved")

    monkeypatch.setattr(ChatDocumentService, "summarize", summarize)
    response = client.post(
        "/chat/documents",
        files={"file": ("source.txt", b"Save this before AI starts", "text/plain")},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert observations == [("Save this before AI starts", b"Save this before AI starts")]
    assert response.json()["analysis_status"] == "unavailable"


def test_import_only_skips_ai_and_keeps_raw_text_for_later_context(client, monkeypatch, db_session):
    headers = _headers(client)

    async def unexpected(*args):
        raise AssertionError("Import only must not invoke any AI analysis")

    monkeypatch.setattr(ChatDocumentService, "summarize", unexpected)
    body = _upload_without_ai(client, headers)
    assert body["analysis_status"] == "skipped"
    saved = db_session.query(ChatDocumentAttachment).one()
    assert saved.structured_summary is None
    assert saved.raw_text == "Company: Catering Solutions"
    sources = ChatRepository().get_document_summaries(db_session, body["session_id"])
    assert any("Company: Catering Solutions" in source for source in sources)


def test_successful_analysis_replaces_pending_reply_and_saves_summary(
    client, monkeypatch, db_session
):
    headers = _headers(client)

    async def completed(*args):
        return "# Summary\nCatering Solutions has draft status."

    monkeypatch.setattr(ChatDocumentService, "summarize", completed)
    response = client.post(
        "/chat/documents",
        files={"file": ("source.txt", b"Company: Catering Solutions", "text/plain")},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_status"] == "complete"
    assert body["response"] == "# Summary\nCatering Solutions has draft status."
    saved = db_session.query(ChatDocumentAttachment).one()
    assert saved.structured_summary == body["response"]
    assert saved.raw_text == "Company: Catering Solutions"
    assert db_session.query(ChatMessageRecord).count() == 2
    sources = ChatRepository().get_document_summaries(db_session, body["session_id"])
    assert sources == [body["response"]]


@pytest.mark.parametrize("output_format", ["pdf", "docx"])
def test_export_uploaded_source_uses_entire_raw_text_without_ai(
    client, monkeypatch, db_session, output_format
):
    headers = _headers(client)
    source_text = "\n".join(
        ["FIRST SOURCE LINE"]
        + [f"Row {index}: Catering Solutions report content." for index in range(150)]
        + ["LAST SOURCE LINE MUST NOT BE TRUNCATED"]
    )
    assert len(source_text) > 2000

    async def unavailable(*args):
        raise ChatModelUnavailableError("Document AI generation timed out")

    monkeypatch.setattr(ChatDocumentService, "summarize", unavailable)
    upload = client.post(
        "/chat/documents",
        files={"file": ("source.txt", source_text.encode("utf-8"), "text/plain")},
        headers=headers,
    )
    assert upload.status_code == 200, upload.text
    body = upload.json()
    assert body["analysis_status"] == "unavailable"
    assert "LAST SOURCE LINE MUST NOT BE TRUNCATED" not in body["response"]

    async def unexpected(*args):
        raise AssertionError("Source export must not invoke AI")

    monkeypatch.setattr(ChatDocumentService, "generate_content", unexpected)
    result = client.post(
        "/chat/documents/generate",
        json={
            "mode": "export",
            "source_document_id": body["attachment"]["id"],
            "output_format": output_format,
        },
        headers=headers,
    )
    assert result.status_code == 200, result.text
    generated = result.json()
    assert generated["session_id"] == body["session_id"]
    assert generated["response"] == source_text
    assert db_session.query(ChatSession).count() == 1
    saved = db_session.query(ChatDocumentAttachment).filter_by(kind="generated").one()
    assert saved.structured_summary == source_text
    download = client.get(f"/chat/documents/{saved.id}/download", headers=headers)
    assert download.status_code == 200
    if output_format == "pdf":
        text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(download.content)).pages)
    else:
        text = "\n".join(p.text for p in Document(BytesIO(download.content)).paragraphs)
    assert "FIRST SOURCE LINE" in text
    assert "Row 149: Catering Solutions report content." in text
    assert "LAST SOURCE LINE MUST NOT BE TRUNCATED" in text


def test_source_export_rejects_other_users_documents(client, monkeypatch, db_session):
    owner = _headers(client, "source-owner@example.com")
    uploaded = _upload_without_ai(client, owner)
    stranger = _headers(client, "source-stranger@example.com")

    def unexpected(*args, **kwargs):
        raise AssertionError("Unauthorized source export must not render content")

    monkeypatch.setattr(ChatDocumentService, "render", unexpected)
    response = client.post(
        "/chat/documents/generate",
        json={"mode": "export", "source_document_id": uploaded["attachment"]["id"]},
        headers=stranger,
    )
    assert response.status_code == 404
    assert db_session.query(ChatDocumentAttachment).count() == 1
    assert db_session.query(ChatSession).count() == 1


def test_source_export_rejects_mismatched_session(client, monkeypatch, db_session):
    headers = _headers(client)
    uploaded = _upload_without_ai(client, headers)
    other_session_id = client.post(
        "/chat/sessions", json={"title": "Other chat"}, headers=headers
    ).json()["id"]

    def unexpected(*args, **kwargs):
        raise AssertionError("Source export must not cross conversation boundaries")

    monkeypatch.setattr(ChatDocumentService, "render", unexpected)
    response = client.post(
        "/chat/documents/generate",
        json={
            "mode": "export",
            "session_id": other_session_id,
            "source_document_id": uploaded["attachment"]["id"],
        },
        headers=headers,
    )
    assert response.status_code == 404
    assert db_session.query(ChatDocumentAttachment).count() == 1


@pytest.mark.parametrize("invalid_source", ["generated", "empty"])
def test_source_export_requires_an_uploaded_document_with_readable_text(
    client, db_session, invalid_source
):
    headers = _headers(client)
    uploaded = _upload_without_ai(client, headers)
    source = db_session.query(ChatDocumentAttachment).one()
    if invalid_source == "generated":
        source.kind = "generated"
    else:
        source.raw_text = " \n\t "
    db_session.commit()
    response = client.post(
        "/chat/documents/generate",
        json={"mode": "export", "source_document_id": uploaded["attachment"]["id"]},
        headers=headers,
    )
    assert response.status_code == 422, response.text
    assert db_session.query(ChatDocumentAttachment).count() == 1
    assert db_session.query(ChatMessageRecord).count() == 2


@pytest.mark.parametrize(
    ("payload", "expected_status"),
    [
        ({}, 422),
        ({"mode": "export"}, 422),
        ({"mode": "ai", "source_document_id": 1, "instruction": "Create a report"}, 404),
        ({"source_document_id": 1, "instruction": "Create a report"}, 404),
        ({"mode": "export", "source_document_id": 0}, 422),
    ],
)
def test_generation_request_validation_and_missing_sources(
    client, db_session, payload, expected_status
):
    headers = _headers(client)
    response = client.post("/chat/documents/generate", json=payload, headers=headers)
    assert response.status_code == expected_status, response.text
    assert db_session.query(ChatSession).count() == 0
    assert db_session.query(ChatDocumentAttachment).count() == 0


@pytest.mark.parametrize(
    ("filename", "content", "expected_status"),
    [("bad.exe", b"unsupported", 415), ("bad.pdf", b"not-a-pdf", 422), ("empty.txt", b"", 422)],
)
def test_import_only_does_not_bypass_file_validation(
    client, monkeypatch, db_session, filename, content, expected_status
):
    headers = _headers(client)

    async def unexpected(*args):
        raise AssertionError("Invalid files must not invoke AI")

    monkeypatch.setattr(ChatDocumentService, "summarize", unexpected)
    response = client.post(
        "/chat/documents",
        files={"file": (filename, content, "application/octet-stream")},
        data={"analyze": "false"},
        headers=headers,
    )
    assert response.status_code == expected_status
    assert db_session.query(ChatSession).count() == 0
    assert db_session.query(ChatDocumentAttachment).count() == 0
