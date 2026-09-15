"""Run persistent, real-file acceptance checks for the document agent.

The script treats every uploaded document as untrusted data. It exercises the
same deterministic readers, generators, validators, and conversion registry as
the application, writes generated artifacts outside production storage, and
emits machine-readable evidence for the Markdown acceptance report.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import platform
import shutil
import subprocess
import sys
import time
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.services.chat_document_service import ChatDocumentService
from app.services.conversion.document_conversion_service import (
    DocumentConversionService,
    UnsupportedDocumentConversionError,
)
from app.services.document.document_generation_service import DocumentGenerationService
from app.services.document.document_operation_registry import DocumentType
from app.services.document.document_reading_service import DocumentReadingService
from app.services.document.document_validation_service import DocumentValidationService
from app.services.document.exceptions import DocumentProcessingUnavailableError
from app.services.document.extraction_models import (
    GeneratedTableContent,
    StructuredDocumentContent,
)

SAMPLE_NAMES = (
    "UAT-EMOS-Dish Template 10092026.xlsx",
    "sample-1.pdf",
    "nutrition-document.docx",
    "nutrition-document.pdf",
    "sample-1mb.pdf",
)
CREATION_TYPES = (
    DocumentType.DOCX,
    DocumentType.PDF,
    DocumentType.XLSX,
    DocumentType.CSV,
    DocumentType.PPTX,
    DocumentType.TXT,
    DocumentType.MARKDOWN,
)
UAT_EXPANSION_PROMPT = (
    "Create separate rows based on Regular, Easy to Chew, Soft & Bite, Minced & Moist and Pureed."
)


@dataclass
class Result:
    category: str
    action: str
    source: str | None
    output: str | None
    status: str
    seconds: float
    fidelity: str | None = None
    bytes: int | None = None
    parser: str | None = None
    content: str | None = None
    limitation: str | None = None


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _parser_name(document_type: DocumentType) -> str:
    return {
        DocumentType.PDF: "pypdf/pdfplumber",
        DocumentType.DOCX: "python-docx",
        DocumentType.XLSX: "openpyxl",
        DocumentType.CSV: "csv",
        DocumentType.PPTX: "python-pptx",
        DocumentType.TXT: "UTF-8 decoder",
        DocumentType.MARKDOWN: "Markdown text reader",
        DocumentType.IMAGE: "Pillow/Tesseract",
    }[document_type]


def _profile(extracted, title: str) -> StructuredDocumentContent:
    paragraphs = [
        f"Source validation profile generated from {title}.",
        f"Extracted coverage: {extracted.coverage or 'available content'}.",
    ]
    visible_lines = [line.strip() for line in extracted.text.splitlines() if line.strip()]
    paragraphs.extend(line[:800] for line in visible_lines[:8])
    tables: list[GeneratedTableContent] = []
    if extracted.tables:
        source_table = extracted.tables[0]
        rows = [list(row[:8]) for row in source_table.rows[:9]]
        headers = rows.pop(0) if source_table.header_detected and rows else []
        width = len(headers) or max((len(row) for row in rows), default=0)
        if width:
            headers = (headers + [f"Column {index}" for index in range(1, width + 1)])[:width]
            normalized_rows = [(row + [""] * width)[:width] for row in rows]
            tables.append(
                GeneratedTableContent(
                    title=f"Sample from {source_table.name}",
                    headers=headers,
                    rows=normalized_rows,
                )
            )
    if not paragraphs[2:] and not tables:
        paragraphs.append("The source contained no additional preview text.")
    return StructuredDocumentContent(title=title, paragraphs=paragraphs, tables=tables)


def _save_and_reopen(
    *,
    data: bytes,
    filename: str,
    output_dir: Path,
    reader: DocumentReadingService,
    validator: DocumentValidationService,
) -> tuple[Path, str]:
    document_type = (
        DocumentType.MARKDOWN
        if filename.endswith(".md")
        else DocumentType(Path(filename).suffix.lstrip(".").casefold())
    )
    validator.validate(data, document_type)
    path = output_dir / filename
    path.write_bytes(data)
    reopened = reader.read(path.read_bytes(), path.name)
    if not reopened.text.strip() and not reopened.tables:
        raise AssertionError("Reopened output did not contain readable content.")
    content = (
        f"chars={len(reopened.text)}, tables={len(reopened.tables)}, "
        f"coverage={reopened.coverage or 'verified'}"
    )
    return path, content


def _record_generation(
    *,
    profile: StructuredDocumentContent,
    source_name: str,
    round_number: int,
    output_type: DocumentType,
    output_dir: Path,
    generator: DocumentGenerationService,
    reader: DocumentReadingService,
    validator: DocumentValidationService,
) -> Result:
    started = time.perf_counter()
    requested = f"{Path(source_name).stem}-generated-{round_number}.{output_type.value}"
    try:
        generated = generator.generate(
            profile,
            output_type.value,
            requested_filename=requested,
        )
        path, content = _save_and_reopen(
            data=generated.file_data,
            filename=generated.filename,
            output_dir=output_dir,
            reader=reader,
            validator=validator,
        )
        return Result(
            category="creation",
            action=f"create_{output_type.value}",
            source=source_name,
            output=str(path.resolve()),
            status="PASS",
            seconds=round(time.perf_counter() - started, 3),
            fidelity=generated.fidelity.value,
            bytes=len(generated.file_data),
            parser=_parser_name(output_type),
            content=content,
        )
    except Exception as error:  # evidence runner must report every failure and continue
        return Result(
            category="creation",
            action=f"create_{output_type.value}",
            source=source_name,
            output=None,
            status="FAIL",
            seconds=round(time.perf_counter() - started, 3),
            limitation=f"{type(error).__name__}: {error}",
        )


def _record_conversion(
    *,
    source_path: Path,
    output_type: DocumentType,
    output_dir: Path,
    converter: DocumentConversionService,
    reader: DocumentReadingService,
    validator: DocumentValidationService,
) -> Result:
    started = time.perf_counter()
    source_data = source_path.read_bytes()
    source_hash = _sha256(source_data)
    action = f"{source_path.suffix.lstrip('.').casefold()}_to_{output_type.value}"
    try:
        generated = converter.convert(source_data, source_path.name, output_type)
        output_name = (
            f"{Path(source_path).stem}-to-{output_type.value}"
            f"{DocumentGenerationService.EXTENSIONS[output_type]}"
        )
        path, content = _save_and_reopen(
            data=generated.file_data,
            filename=output_name,
            output_dir=output_dir,
            reader=reader,
            validator=validator,
        )
        if _sha256(source_path.read_bytes()) != source_hash:
            raise AssertionError("Source file changed during conversion.")
        return Result(
            category="conversion",
            action=action,
            source=str(source_path.resolve()),
            output=str(path.resolve()),
            status="PASS",
            seconds=round(time.perf_counter() - started, 3),
            fidelity=generated.fidelity.value,
            bytes=len(generated.file_data),
            parser=_parser_name(output_type),
            content=content,
            limitation=generated.fidelity_note,
        )
    except UnsupportedDocumentConversionError as error:
        return Result(
            category="conversion",
            action=action,
            source=str(source_path.resolve()),
            output=None,
            status="UNSUPPORTED",
            seconds=round(time.perf_counter() - started, 3),
            limitation=str(error),
        )
    except DocumentProcessingUnavailableError as error:
        return Result(
            category="conversion",
            action=action,
            source=str(source_path.resolve()),
            output=None,
            status="SKIPPED",
            seconds=round(time.perf_counter() - started, 3),
            limitation=str(error),
        )
    except Exception as error:
        return Result(
            category="conversion",
            action=action,
            source=str(source_path.resolve()),
            output=None,
            status="FAIL",
            seconds=round(time.perf_counter() - started, 3),
            limitation=f"{type(error).__name__}: {error}",
        )


def _record_uat_expansion(source_path: Path, output_dir: Path) -> Result:
    """Run the production dietary-category expansion against the real UAT file."""
    started = time.perf_counter()
    source_data = source_path.read_bytes()
    source_hash = _sha256(source_data)
    try:
        generated, filename, actions = ChatDocumentService().format_spreadsheet(
            source_data,
            source_path.name,
            UAT_EXPANSION_PROMPT,
        )
        output_path = output_dir / filename
        output_path.write_bytes(generated)
        workbook = load_workbook(BytesIO(generated), read_only=True, data_only=False)
        try:
            sheet = workbook["Sheet1"]
            data_rows = sheet.max_row - 1
            columns = sheet.max_column
        finally:
            workbook.close()
        if data_rows != 3_136 or columns != 77:
            raise AssertionError(
                f"Expected 3,136 data rows and 77 columns; got {data_rows} and {columns}."
            )
        if _sha256(source_path.read_bytes()) != source_hash:
            raise AssertionError("Source file changed during category expansion.")
        return Result(
            category="modification",
            action="expand_dish_by_dietary_category",
            source=str(source_path.resolve()),
            output=str(output_path.resolve()),
            status="PASS",
            seconds=round(time.perf_counter() - started, 3),
            fidelity="high",
            bytes=len(generated),
            parser="openpyxl",
            content=(f"rows={data_rows}, columns={columns}, actions={' | '.join(actions)}"),
        )
    except Exception as error:
        return Result(
            category="modification",
            action="expand_dish_by_dietary_category",
            source=str(source_path.resolve()),
            output=None,
            status="FAIL",
            seconds=round(time.perf_counter() - started, 3),
            limitation=f"{type(error).__name__}: {error}",
        )


def _environment() -> dict[str, object]:
    libraries = {}
    for name in (
        "pypdf",
        "pdfplumber",
        "python-docx",
        "openpyxl",
        "python-pptx",
        "reportlab",
        "pytesseract",
    ):
        libraries[name] = importlib.metadata.version(name)
    soffice = DocumentConversionService().libreoffice.binary
    soffice_version = None
    if soffice:
        version_binary = Path(soffice)
        if version_binary.suffix.casefold() == ".exe":
            console_binary = version_binary.with_suffix(".com")
            if console_binary.is_file():
                version_binary = console_binary
        completed = subprocess.run(
            [str(version_binary), "--headless", "--version"],
            capture_output=True,
            check=False,
            text=True,
            timeout=15,
        )
        soffice_version = (completed.stdout or completed.stderr).strip() or None
    return {
        "timestamp_utc": datetime.now(UTC).isoformat(),
        "os": platform.platform(),
        "python": sys.version,
        "libraries": libraries,
        "libreoffice_path": soffice,
        "libreoffice_version": soffice_version,
        "tesseract_path": shutil.which("tesseract"),
    }


def run(sample_dir: Path, output_dir: Path) -> dict[str, object]:
    output_dir.mkdir(parents=True, exist_ok=False)
    reader = DocumentReadingService()
    validator = DocumentValidationService()
    generator = DocumentGenerationService(validator=validator)
    converter = DocumentConversionService(
        reader=reader,
        generator=generator,
        validator=validator,
    )
    samples = {name: sample_dir / name for name in SAMPLE_NAMES}
    missing = [str(path) for path in samples.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing sample files: " + ", ".join(missing))

    results: list[Result] = []
    extracted_by_name = {}
    source_hashes = {name: _sha256(path.read_bytes()) for name, path in samples.items()}
    for name, path in samples.items():
        started = time.perf_counter()
        try:
            extracted = reader.read(path.read_bytes(), name)
            extracted_by_name[name] = extracted
            results.append(
                Result(
                    category="reading",
                    action=f"read_{extracted.document_type.value}",
                    source=str(path.resolve()),
                    output=None,
                    status="PASS",
                    seconds=round(time.perf_counter() - started, 3),
                    fidelity=validator.fidelity_for_extraction(extracted).fidelity.value,
                    bytes=path.stat().st_size,
                    parser=_parser_name(extracted.document_type),
                    content=(
                        f"chars={len(extracted.text)}, blocks={len(extracted.blocks)}, "
                        f"tables={len(extracted.tables)}, coverage={extracted.coverage}"
                    ),
                )
            )
        except Exception as error:
            results.append(
                Result(
                    category="reading",
                    action=f"read_{path.suffix.lstrip('.')}",
                    source=str(path.resolve()),
                    output=None,
                    status="FAIL",
                    seconds=round(time.perf_counter() - started, 3),
                    limitation=f"{type(error).__name__}: {error}",
                )
            )

    generation_sources = (
        "UAT-EMOS-Dish Template 10092026.xlsx",
        "sample-1.pdf",
        "nutrition-document.docx",
    )
    generated_paths: dict[tuple[str, DocumentType], Path] = {}
    for round_number, source_name in enumerate(generation_sources, start=1):
        extracted = extracted_by_name[source_name]
        profile = _profile(extracted, f"Validation from {Path(source_name).stem}")
        for output_type in CREATION_TYPES:
            result = _record_generation(
                profile=profile,
                source_name=source_name,
                round_number=round_number,
                output_type=output_type,
                output_dir=output_dir,
                generator=generator,
                reader=reader,
                validator=validator,
            )
            results.append(result)
            if result.status == "PASS" and result.output:
                generated_paths[(source_name, output_type)] = Path(result.output)

    results.append(
        _record_uat_expansion(samples["UAT-EMOS-Dish Template 10092026.xlsx"], output_dir)
    )

    for source_name in (
        "sample-1.pdf",
        "nutrition-document.pdf",
        "sample-1mb.pdf",
    ):
        for output_type in (DocumentType.DOCX, DocumentType.TXT, DocumentType.MARKDOWN):
            results.append(
                _record_conversion(
                    source_path=samples[source_name],
                    output_type=output_type,
                    output_dir=output_dir,
                    converter=converter,
                    reader=reader,
                    validator=validator,
                )
            )
        results.append(
            _record_conversion(
                source_path=samples[source_name],
                output_type=DocumentType.XLSX,
                output_dir=output_dir,
                converter=converter,
                reader=reader,
                validator=validator,
            )
        )

    for output_type in (DocumentType.PDF, DocumentType.TXT, DocumentType.MARKDOWN):
        results.append(
            _record_conversion(
                source_path=samples["nutrition-document.docx"],
                output_type=output_type,
                output_dir=output_dir,
                converter=converter,
                reader=reader,
                validator=validator,
            )
        )
    for output_type in (DocumentType.CSV, DocumentType.PDF):
        results.append(
            _record_conversion(
                source_path=samples["UAT-EMOS-Dish Template 10092026.xlsx"],
                output_type=output_type,
                output_dir=output_dir,
                converter=converter,
                reader=reader,
                validator=validator,
            )
        )

    conversion_sources = (
        (("sample-1.pdf", DocumentType.PPTX), DocumentType.PDF),
        (("UAT-EMOS-Dish Template 10092026.xlsx", DocumentType.CSV), DocumentType.XLSX),
        (("sample-1.pdf", DocumentType.TXT), DocumentType.MARKDOWN),
        (("sample-1.pdf", DocumentType.MARKDOWN), DocumentType.TXT),
    )
    for source_key, output_type in conversion_sources:
        source_path = generated_paths.get(source_key)
        if source_path is None:
            results.append(
                Result(
                    category="conversion",
                    action=f"generated_{source_key[1].value}_to_{output_type.value}",
                    source=source_key[0],
                    output=None,
                    status="SKIPPED",
                    seconds=0,
                    limitation="The prerequisite generated source failed validation.",
                )
            )
            continue
        results.append(
            _record_conversion(
                source_path=source_path,
                output_type=output_type,
                output_dir=output_dir,
                converter=converter,
                reader=reader,
                validator=validator,
            )
        )

    source_integrity = {
        name: _sha256(path.read_bytes()) == source_hashes[name] for name, path in samples.items()
    }
    payload = {
        "environment": _environment(),
        "sample_dir": str(sample_dir.resolve()),
        "output_dir": str(output_dir.resolve()),
        "source_integrity": source_integrity,
        "results": [asdict(result) for result in results],
        "summary": {
            status: sum(result.status == status for result in results)
            for status in ("PASS", "FAIL", "SKIPPED", "UNSUPPORTED")
        },
    }
    (output_dir / "results.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sample-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    payload = run(args.sample_dir, args.output_dir)
    print(json.dumps(payload["summary"], sort_keys=True))
    print(payload["output_dir"])
    return 1 if payload["summary"]["FAIL"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
